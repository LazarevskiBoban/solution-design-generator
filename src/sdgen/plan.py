from __future__ import annotations

import re
from collections import Counter
from pathlib import Path
from typing import Literal

import yaml
from pydantic import BaseModel, Field

from sdgen.analyze import slugify
from sdgen.blueprint import TITLE_MAX, Blueprint, Section, cap_title, composite_fields
from sdgen.brief import Brief, dump_brief, looks_joined, split_joined, table_rows
from sdgen.content import detail_key
from sdgen.flow import FlowSpec
from sdgen.grounding import GROUNDING_RULE
from sdgen.llm import LLMClient
from sdgen.manifest import FieldSpec, Manifest
from sdgen.material import material_text
from sdgen.render import ExtraSlide

Source = Literal["draft", "keep", "blank", "diagram", "mechanical"]
SOURCES = ("draft", "keep", "blank", "diagram", "mechanical")
EFFORT_RE = re.compile(r"effort", re.IGNORECASE)
VERSION_TITLE_RE = re.compile(r"version|author|contributor", re.IGNORECASE)
GIST_CHARS = 200
WALKTHROUGH_SUFFIX = "_walkthrough"
OPEN_QUESTIONS_KEY = "extra_open_questions"
OPEN_QUESTIONS_COLUMNS = ["Ref", "Question", "Ask", "Status"]
OPEN_QUESTIONS_RE = re.compile(r"open questions?", re.IGNORECASE)

EXTRA_DEFAULTS = {
    "acceptance_criteria": ("Acceptance Criteria and Test Scenarios", "table", ["Ref", "Scenario", "Expected result", "Evidence"]),
    "operations": ("Error Handling, Monitoring and Operations", "text", []),
    "decisions_log": ("Decisions", "table", ["Ref", "Decision", "Owner", "Status"]),
    "build_notes": ("Build Notes", "text", []),
}
# The slides a developer builds from, offered for every design as ticked extras: key, title, columns, what the rows hold.
DEVELOPER_EXTRAS: list[tuple[str, str, list[str], str]] = [
    ("extra_interface_inventory", "Interface inventory", ["#", "Party", "Flow", "Direction", "Source", "Target", "Encryption", "Cut-off"], "One row per party and flow: what moves where, encrypted or not, by when."),
    ("extra_configuration", "Configuration", ["Component", "Parameter", "Value", "Note"], "Adapter, folder, polling, data store, partner directory, keystore and alert parameters a developer sets."),
    ("extra_file_naming", "File naming", ["Flow", "Pattern", "Temp name", "Example"], "Names on both sides, including the temporary name while a file is written."),
    ("extra_connectivity", "Connectivity and environments", ["Environment", "Endpoint", "Host", "Port", "Account", "Key or cert", "Network path"], "Every endpoint per environment; [TBC] where the brief is silent."),
    ("extra_security_access", "Security and access", ["Endpoint", "Folder or resource", "Account", "Read", "Write", "Move or delete"], "Least privilege per folder or resource and account."),
    ("extra_cutover", "Cutover and dual run", ["Step", "When", "Owner", "Check"], "The steps to go live, in order, with who checks what."),
    ("extra_raci", "RACI", ["Activity", "R", "A", "C", "I"], "Responsible, accountable, consulted and informed for keys, endpoints, folders and runs."),
    ("extra_build_checklist", "Build checklist", ["#", "Step", "Depends on", "Owner"], "Build steps in dependency order: endpoints and access first, keys second, build last."),
    ("extra_assumptions", "Assumptions and constraints", ["#", "Statement", "Type", "Owner", "Status"], "Assumptions and constraints from the investigation, each with an owner."),
    ("extra_nfr", "Non-functional requirements", ["Aspect", "Requirement", "Source"], "Volumes, frequency, sizes, security, retention, environments, availability."),
]
DEVELOPER_KEYS = {key for key, _, _, _ in DEVELOPER_EXTRAS}
OPERATIONS_RE = re.compile(r"operation|error handling|runbook", re.IGNORECASE)
SEQUENCE_RE = re.compile(r"handshake|sequence|negotiat|\blog-?in\b|authenticat|steps between", re.IGNORECASE)

PLAN_SCHEMA = {
    "type": "object",
    "properties": {
        "decisions": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "key": {"type": "string"},
                    "use": {"type": "boolean"},
                    "title": {"type": "string"},
                    "source": {"type": "string", "enum": list(SOURCES)},
                    "reason": {"type": "string"},
                },
                "required": ["key", "use", "source"],
            },
        },
        "extras": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "key": {"type": "string"},
                    "title": {"type": "string"},
                    "kind": {"type": "string", "enum": ["table", "text"]},
                    "columns": {"type": "array", "items": {"type": "string"}},
                    "before": {"type": "string"},
                    "reason": {"type": "string"},
                },
                "required": ["key", "title", "kind"],
            },
        },
        "flows": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {"section": {"type": "string"}, "title": {"type": "string"}, "purpose": {"type": "string"}, "material_id": {"type": "string"}},
                "required": ["section", "title", "purpose"],
            },
        },
        "clear": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {"slide": {"type": "integer"}, "shape": {"type": "integer"}, "reason": {"type": "string"}},
                "required": ["slide", "shape"],
            },
        },
    },
    "required": ["decisions"],
}
LEFTOVER_MAX_CHARS = 200

SYSTEM_PROMPT = """You plan which slides of a solution-design template apply to one specific integration.
You get the template outline (one line per section with its key, kind, what it asks for and a
gist of the text an earlier project wrote there) and the brief of the new integration.
For every section decide:
- use: false when the section only makes sense for the earlier project (its flows, systems or
  API calls have no counterpart in the brief); true otherwise. The cover, document version
  control, contents and guiding principles always apply. A diagram section is a reusable slot:
  when the brief describes a flow, pattern, exception path or handshake it can show, keep it
  and give it that flow's title; mark it not used only when no flow of the brief is left for it.
- title: a new title only when the template title names the earlier project's systems or
  flows and the slide's purpose still fits the new integration (for example a flow slide that
  becomes the new integration's equivalent flow); otherwise an empty string keeps the title.
  A title has at most 60 characters, no arrows, and never the integration's name: the subject
  stays on the cover only.
- source: "draft" when the model should write it, "keep" for template text that applies as it
  is (guiding principles, contents), "blank" when the slide should stay empty, "diagram" for
  diagram slides to draw from the brief, "mechanical" for cover, version control and references.
- reason: one short sentence.
When several sections share a title, the template repeats a slide type; keep only as many
of them as the brief needs and mark the rest not used, so the document has no duplicates.
Propose extra slides only when the brief has content for them: an acceptance criteria table
(columns Ref, Scenario, Expected result, Evidence), an operations text slide, a decisions table
(columns Ref, Decision, Owner, Status), and a build notes text slide (key build_notes) for
developer detail such as file naming, cut-off times, reprocessing steps and configuration keys,
when the brief carries operations, investigation or non-functional detail. Place extras before
the effort estimation. Open questions get their own slide from the brief; never propose one.
The developer slides (interface inventory, configuration, file naming, connectivity, security and
access, cutover, RACI, build checklist, assumptions, non-functional requirements) are added
automatically: do not propose them.
For every diagram section in use without an uploaded image, add a flow entry with the title
and purpose of the diagram to draw from the brief. When a reference picture listed below is the
author's own drawing of what that section asks for, give the entry its material_id: the picture
is then placed on the slide instead of a drawing.
The outline may end with template texts that belong to no field. List in "clear" the ones
that describe the earlier project (its systems, plans, dates, names, examples) so they are
emptied; leave generic wording alone (labels, legends, headings, instructions, "see ...").
Titles, extra slides and diagram purposes name only systems and flows from the brief.
""" + GROUNDING_RULE + """
Return only JSON matching the schema."""


class SectionDecision(BaseModel):
    key: str
    use: bool = True
    title: str = ""
    source: Source = "draft"
    reason: str = ""


class ExtraSection(BaseModel):
    key: str
    title: str
    kind: Literal["table", "text"] = "text"
    columns: list[str] = Field(default_factory=list)
    prototype: str = ""
    before: str = ""
    reason: str = ""
    include: bool = True
    generated: bool = False  # built from a drawing, never written by the model


class FlowRequest(BaseModel):
    width_in: float = 0.0  # drawing area of the slot, in inches; zero when unknown
    height_in: float = 0.0
    section: str
    title: str = ""
    purpose: str = ""
    material_id: str = ""  # a reference picture that takes the slot instead of a drawing
    kind: Literal["flow", "sequence"] = "flow"  # a sequence draws participants with numbered arrows


def flow_kind(*texts: str) -> Literal["flow", "sequence"]:
    return "sequence" if SEQUENCE_RE.search(" ".join(texts)) else "flow"


class Leftover(BaseModel):
    """Template text that is bound to no field; the plan decides whether it is project-specific."""

    slide: int
    shape: int
    text: str
    reason: str = ""
    include: bool = True


class SectionPlan(BaseModel):
    decisions: list[SectionDecision] = Field(default_factory=list)
    extras: list[ExtraSection] = Field(default_factory=list)
    flows: list[FlowRequest] = Field(default_factory=list)
    clear: list[Leftover] = Field(default_factory=list)
    model: str = ""

    def decision(self, key: str) -> SectionDecision | None:
        return next((d for d in self.decisions if d.key == key), None)

    def save(self, path: str | Path) -> None:
        Path(path).write_text(yaml.safe_dump(self.model_dump(mode="json"), sort_keys=False, allow_unicode=True, width=100), encoding="utf-8")

    @classmethod
    def load(cls, path: str | Path) -> SectionPlan:
        return cls.model_validate(yaml.safe_load(Path(path).read_text(encoding="utf-8")) or {})


def developer_extras(blueprint: Blueprint, manifest: Manifest | None) -> list[ExtraSection]:
    """The catalogue of developer slides, as ticked extras, for a template that has a plain table slide to clone."""
    extras = []
    for key, title, columns, reason in DEVELOPER_EXTRAS:
        prototype = _prototype(blueprint, "table", list(columns), manifest, allow_composite=False)
        if prototype:
            extras.append(ExtraSection(key=key, title=title, kind="table", columns=list(columns), prototype=prototype, before=_before(blueprint), reason=reason))
    return extras


def default_plan(blueprint: Blueprint, images: set[str] | None = None, manifest: Manifest | None = None) -> SectionPlan:
    """Every section applies; sources follow the section kind; with a manifest the developer slides come along."""
    images = images or set()
    decisions = []
    flows = []
    for section in blueprint.sections:
        source: Source = "draft"
        if section.kind in ("static", "divider"):
            source = "keep"
        elif section.kind == "cover":
            source = "mechanical"
        elif section.kind == "diagram":
            has_image = any(k in images for k in section.fields)
            source = "draft" if has_image else "diagram"
            if not has_image:
                flows.append(FlowRequest(section=section.key, title=section.title, purpose=section.ask, kind=flow_kind(section.title, section.ask)))
        decisions.append(SectionDecision(key=section.key, use=True, source=source))
    return SectionPlan(decisions=decisions, flows=flows, extras=developer_extras(blueprint, manifest) if manifest is not None else [])


def plan_sections(
    brief: Brief,
    blueprint: Blueprint,
    manifest: Manifest,
    llm: LLMClient,
    images: set[str] | None = None,
    leftovers: list[Leftover] | None = None,
) -> SectionPlan:
    base = default_plan(blueprint, images, manifest)
    pictures = picture_ids(brief)
    data = llm.complete_json(SYSTEM_PROMPT, _prompt(brief, blueprint, manifest, images or set(), leftovers or []), _schema_with_pictures(pictures) if pictures else PLAN_SCHEMA, name="section_plan")
    plan = merge_plan(base, data, blueprint, images or set(), leftovers or [], manifest, brief)
    plan.model = str(getattr(llm, "label", llm.name))
    return plan


def picture_ids(brief: Brief) -> list[str]:
    """The ids of the reference pictures a planner may place on a diagram slide."""
    return [m.id for m in brief.material if m.kind == "image" and m.file]


def _schema_with_pictures(ids: list[str]) -> dict:
    import copy

    schema = copy.deepcopy(PLAN_SCHEMA)
    schema["properties"]["flows"]["items"]["properties"]["material_id"] = {"type": "string", "enum": list(ids) + [""]}
    return schema


def leftover_texts(template_path: str | Path, manifest: Manifest, blueprint: Blueprint, max_chars: int = LEFTOVER_MAX_CHARS) -> list[Leftover]:
    """Short texts on content slides that no field replaces; candidates for clearing per design."""
    from pptx import Presentation

    from sdgen.inventory import walk_shapes

    bound = {(b.slide, b.shape.id) for f in manifest.fields for b in f.bindings}
    sections = {s.slide: s for s in blueprint.sections}
    result: list[Leftover] = []
    for number, slide in enumerate(Presentation(str(template_path)).slides, 1):
        section = sections.get(number)
        if section is None or section.kind in ("divider", "static"):
            continue
        title_id = slide.shapes.title.shape_id if slide.shapes.title is not None else None
        for shape in walk_shapes(slide.shapes):
            if shape.shape_id == title_id or (number, shape.shape_id) in bound or not getattr(shape, "has_text_frame", False):
                continue
            text = " ".join(shape.text_frame.text.split())
            if not text or "{{" in text or len(text) > max_chars:
                continue
            result.append(Leftover(slide=number, shape=shape.shape_id, text=text))
    return result


def merge_plan(
    base: SectionPlan,
    data: dict,
    blueprint: Blueprint,
    images: set[str],
    leftovers: list[Leftover] | None = None,
    manifest: Manifest | None = None,
    brief: Brief | None = None,
) -> SectionPlan:
    sections = {s.key: s for s in blueprint.sections}
    plan = base.model_copy(deep=True)
    for item in data.get("decisions") or []:
        decision = plan.decision(str(item.get("key", "")))
        if decision is None:
            continue
        section = sections[decision.key]
        if section.kind == "cover":
            continue
        decision.use = bool(item.get("use", decision.use))
        title = cap_title(str(item.get("title") or ""))
        decision.title = "" if title == section.title else title
        source = str(item.get("source") or decision.source)
        if source in SOURCES:
            decision.source = source  # type: ignore[assignment]
        decision.reason = str(item.get("reason") or "").strip()
    extras = []
    for item in data.get("extras") or []:
        title = str(item.get("title") or "").strip()
        if not title:
            continue
        key = slugify(str(item.get("key") or title))
        if not key.startswith("extra_"):
            key = f"extra_{key}"
        kind = "table" if str(item.get("kind") or "text") == "table" else "text"
        columns = [str(c).strip() for c in (item.get("columns") or []) if str(c).strip()] if kind == "table" else []
        if kind == "table" and not columns:
            columns = next((c for k, (t, kd, c) in EXTRA_DEFAULTS.items() if k.split("_")[0] in key or k.split("_")[0] in title.lower()), ["Ref", "Item", "Notes"])
        pasted = table_rows(brief.operations) if brief is not None and OPERATIONS_RE.search(f"{key} {title}") else None
        if pasted is not None:
            kind, columns = "table", pasted[0]  # the brief holds the operations as a table: the slide keeps its columns
        if key in DEVELOPER_KEYS or slugify(title) in {slugify(t) for _, t, _, _ in DEVELOPER_EXTRAS}:
            continue  # the catalogue version of this slide comes along anyway
        if any(e.key == key for e in extras):
            continue
        before = str(item.get("before") or "").strip()
        extras.append(
            ExtraSection(
                key=key,
                title=title,
                kind=kind,
                columns=columns,
                prototype=_prototype(blueprint, kind, columns, manifest),
                before=before if before in sections else _before(blueprint),
                reason=str(item.get("reason") or "").strip(),
            )
        )
    plan.extras = extras + [e for e in base.extras if e.key in DEVELOPER_KEYS]
    flows = []
    for item in data.get("flows") or []:
        key = str(item.get("section") or "")
        section = sections.get(key)
        decision = plan.decision(key)
        if section is None or section.kind != "diagram" or decision is None or not decision.use:
            continue
        if any(k in images for k in section.fields) or any(f.section == key for f in flows):
            continue  # a model answer may name the same diagram twice; the first entry wins
        material = str(item.get("material_id") or "").strip()
        title, purpose = str(item.get("title") or section.title).strip(), str(item.get("purpose") or "").strip()
        flows.append(FlowRequest(section=key, title=title, purpose=purpose, material_id=material if brief is not None and material in picture_ids(brief) else "", kind=flow_kind(section.title, section.ask, title, purpose)))
    if flows:
        plan.flows = flows
    else:
        plan.flows = [f for f in plan.flows if (plan.decision(f.section) or SectionDecision(key=f.section)).use]
    known = {(item.slide, item.shape): item for item in (leftovers or [])}
    clear = []
    for item in data.get("clear") or []:
        try:
            key = (int(item.get("slide")), int(item.get("shape")))
        except (TypeError, ValueError):
            continue
        leftover = known.get(key)
        if leftover is not None:
            clear.append(leftover.model_copy(update={"reason": str(item.get("reason") or "").strip(), "include": True}))
    plan.clear = clear
    return plan


def apply_plan(plan: SectionPlan, design, blueprint: Blueprint) -> None:
    """Writes the plan into the design's hidden sections, modes and titles."""
    sections = {s.key: s for s in blueprint.sections}
    design.hidden = [d.key for d in plan.decisions if not d.use and d.key in sections and sections[d.key].kind != "cover"]
    design.modes = {d.key: d.source for d in plan.decisions if d.use and d.source in ("keep", "blank") and d.key in sections and sections[d.key].fields}
    design.titles = {d.key: d.title.strip() for d in plan.decisions if d.use and d.title.strip() and d.key in sections}
    design.plan = plan


def active_extras(plan: SectionPlan | None) -> list[ExtraSection]:
    return [e for e in (plan.extras if plan else []) if e.include]


def extended_manifest(manifest: Manifest, blueprint: Blueprint, extras: list[ExtraSection]) -> Manifest:
    """The manifest plus one field per extra slide, bound to its prototype's shape."""
    fields = list(manifest.fields)
    known = {f.key for f in fields}
    for extra in extras:
        spec = _prototype_spec(manifest, blueprint, extra)
        if spec is None or extra.key in known:
            continue
        binding = spec.bindings[0].model_copy(update={"keep_last_row_if": None})
        fields.append(
            FieldSpec(
                key=extra.key,
                label=extra.title,
                kind="table" if extra.kind == "table" else "bullets",
                columns=list(extra.columns),
                guidance=extra.reason,
                bindings=[binding],
            )
        )
        known.add(extra.key)
    return manifest.model_copy(update={"fields": fields})


def extended_blueprint(blueprint: Blueprint, extras: list[ExtraSection]) -> Blueprint:
    """The outline plus one section per extra slide, placed before its target section."""
    sections = list(blueprint.sections)
    for extra in extras:
        proto = _prototype_section(blueprint, extra)
        if proto is None or any(s.key == extra.key for s in sections):
            continue
        ask = ("table: " + ", ".join(extra.columns)) if extra.columns else "text"
        section = Section(key=extra.key, title=extra.title, kind="table" if extra.kind == "table" else "text", slide=proto.slide, fields=[extra.key], ask=ask, generated=extra.generated)
        position = next((i for i, s in enumerate(sections) if s.key == extra.before), len(sections))
        sections.insert(position, section)
    return blueprint.model_copy(update={"sections": sections})


def extra_slides(plan: SectionPlan | None, manifest: Manifest, blueprint: Blueprint, fields: dict, hidden: list[str] | None = None) -> list[ExtraSlide]:
    """Render input for the included extras; `manifest` is the extended manifest."""
    slides = []
    sections = {s.key: s for s in blueprint.sections}
    for extra in active_extras(plan):
        spec = manifest.field(extra.key)
        if spec is None or extra.key in (hidden or []):
            continue
        before = sections[extra.before].slide if extra.before in sections else 0
        slides.append(ExtraSlide(key=extra.key, title=extra.title, spec=spec, value=fields.get(extra.key), before=before))
    return slides


def walkthrough_extras(design, blueprint: Blueprint, manifest: Manifest, flows: dict[str, FlowSpec], order: list[str] | None = None) -> list[ExtraSection]:
    """A "how it works" text slide after every drawn diagram whose slide has no written text or table of its own."""
    prototype = _prototype(blueprint, "text", manifest=manifest)
    if not prototype:
        return []
    template_keys = [s.key for s in blueprint.sections if not s.generated]
    keys = [k for k in (order or template_keys) if k in template_keys]
    extras = []
    for section in blueprint.sections:
        flow = flows.get(section.key)
        if flow is None or not flow.nodes or section.generated or section.key in design.hidden:
            continue
        if any(design.images.get(k) for k in section.fields):
            continue
        written = design.modes.get(section.key, "text") == "text"
        if written and any(f is not None and f.kind != "image" for f in (manifest.field(k) for k in section.fields)):
            continue  # the slide's own text explains the drawing
        position = keys.index(section.key) if section.key in keys else -1
        following = keys[position + 1] if 0 <= position < len(keys) - 1 else ""
        suffix = ": how it works"
        title = cap_title(design.titles.get(section.key, section.title), TITLE_MAX - len(suffix))
        extras.append(ExtraSection(key=f"{section.key}{WALKTHROUGH_SUFFIX}", title=f"{title}{suffix}", kind="text", prototype=prototype, before=following, reason="numbered steps taken from the drawing", generated=True))
    return extras


def open_questions_extras(design, blueprint: Blueprint, manifest: Manifest) -> list[ExtraSection]:
    """One generated slide listing the brief's open questions, unless the template already has such a slide."""
    text = design.brief.open_questions.strip()
    if not text or OPEN_QUESTIONS_KEY in design.hidden:
        return []
    if any(OPEN_QUESTIONS_RE.search(s.title) and not s.generated and s.key not in design.hidden for s in blueprint.sections):
        return []
    extra = ExtraSection(key=OPEN_QUESTIONS_KEY, title="Open Questions", kind="table", columns=list(OPEN_QUESTIONS_COLUMNS), prototype=_prototype(blueprint, "table", OPEN_QUESTIONS_COLUMNS, manifest), before=_before(blueprint), reason="from the brief's open questions", generated=True)
    spec = _prototype_spec(manifest, blueprint, extra) if extra.prototype else None
    if spec is not None and spec.kind == "table" and len(spec.columns) >= 3:
        # A cloned table keeps its physical width, so the columns follow the prototype.
        return [extra.model_copy(update={"columns": OPEN_QUESTIONS_COLUMNS[: len(spec.columns)]})]
    prototype = _prototype(blueprint, "text", manifest=manifest)
    return [extra.model_copy(update={"kind": "text", "columns": [], "prototype": prototype})] if prototype else []


def open_question_lines(text: str) -> list[tuple[str, str]]:
    """(question, who to ask) per non-empty line of the brief field."""
    if looks_joined(text, 2):
        text = split_joined(text, 2)
    found = []
    for raw in text.splitlines():
        line = raw.strip().lstrip("-*• ").strip()
        if line:
            question, _, who = line.partition("|")
            found.append((question.strip(), who.strip()))
    return found


def open_question_rows(text: str, columns: list[str] | None = None) -> list[dict]:
    columns = columns or OPEN_QUESTIONS_COLUMNS
    rows = []
    for number, (question, who) in enumerate(open_question_lines(text), 1):
        values = [str(number), question, who]
        rows.append({column: (values[index] if index < len(values) else "") for index, column in enumerate(columns)})
    return rows


def open_questions_text(text: str) -> str:
    return "\n".join(f"{number}. {question}" + (f" (ask: {who})" if who else "") for number, (question, who) in enumerate(open_question_lines(text), 1))


def _prototype_section(blueprint: Blueprint, extra: ExtraSection) -> Section | None:
    if extra.prototype:
        section = blueprint.section(extra.prototype)
        if section is not None:
            return section
    wanted = "table" if extra.kind == "table" else "text"
    return next((s for s in blueprint.sections if s.kind == wanted and s.fields), None)


def _prototype_spec(manifest: Manifest, blueprint: Blueprint, extra: ExtraSection) -> FieldSpec | None:
    proto = _prototype_section(blueprint, extra)
    if proto is None:
        return None
    specs = [f for f in (manifest.field(k) for k in proto.fields) if f is not None and f.kind != "image" and f.bindings]
    wanted = "table" if extra.kind == "table" else "text"
    match = next((f for f in specs if (f.kind == "table") == (wanted == "table")), None)
    return match or (specs[0] if specs else None)


def _prompt(brief: Brief, blueprint: Blueprint, manifest: Manifest, images: set[str], leftovers: list[Leftover] | None = None) -> str:
    lines = ["# Template outline"]
    counts = Counter(s.title for s in blueprint.sections)
    seen: Counter[str] = Counter()
    for section in blueprint.sections:
        gist = " ".join(section.example.split())
        gist = gist if len(gist) <= GIST_CHARS else gist[:GIST_CHARS].rsplit(" ", 1)[0] + " …"
        has_image = any(k in images for k in section.fields)
        seen[section.title] += 1
        title = section.title if counts[section.title] == 1 else f"{section.title} (slide {seen[section.title]} of {counts[section.title]} with this title)"
        parts = [section.key, f"slide {section.slide}", section.kind, title]
        if section.ask:
            parts.append(f"asks: {section.ask}")
        if gist:
            parts.append(f"earlier document: {gist}")
        if section.kind == "diagram":
            parts.append("image uploaded: " + ("yes" if has_image else "no"))
        lines.append("- " + " | ".join(parts))
    if leftovers:
        lines += ["", "# Template text not bound to any field (slide | shape id | text)"]
        lines += [f"- {item.slide} | {item.shape} | {item.text}" for item in leftovers]
    block = material_text(brief.material, total=8000)
    if block:
        lines += ["", "# Reference material", block]
    pictures = [m for m in brief.material if m.kind == "image" and m.file]
    if pictures:
        lines += ["", "# Reference pictures (material_id | title | what it shows)"]
        lines += [f"- {m.id} | {m.label} | {' '.join(m.text.split())[:160]}" for m in pictures]
    lines += ["", "# Brief", dump_brief(brief)]
    return "\n".join(lines)


def detail_prototypes(blueprint: Blueprint, manifest: Manifest) -> dict[str, FieldSpec]:
    """Per composite field, the field to fill on a cloned plain slide with the complete content."""
    found: dict[str, FieldSpec] = {}
    for key in composite_fields(blueprint, manifest):
        spec = manifest.field(key)
        if spec is None:
            continue
        table = spec.kind == "table"
        prototype = _prototype(blueprint, "table" if table else "text", spec.columns if table else None, manifest, allow_composite=False)
        if not prototype:
            continue
        proto_spec = _prototype_spec(manifest, blueprint, ExtraSection(key=key, title=spec.label, kind="table" if table else "text", prototype=prototype))
        if proto_spec is None or not proto_spec.bindings or (proto_spec.kind == "table") != table:
            continue
        binding = proto_spec.bindings[0].model_copy(update={"keep_prefix": None, "keep_last_row_if": None})
        found[key] = FieldSpec(key=detail_key(key), label=spec.label, kind="table" if table else "bullets", columns=list(spec.columns), bindings=[binding])
    return found


def _prototype(blueprint: Blueprint, kind: str, columns: list[str] | None = None, manifest: Manifest | None = None, allow_composite: bool = True) -> str:
    """The template slide to clone for an extra: a plain slide of the same kind, tables by closest column count."""
    wanted = "table" if kind == "table" else "text"
    best, best_score = "", None
    for section in blueprint.sections:
        if section.kind != wanted or not section.fields or section.generated or VERSION_TITLE_RE.search(section.title):
            continue
        score = 0
        if wanted == "table" and manifest is not None:
            spec = next((manifest.field(k) for k in section.fields if manifest.field(k) and manifest.field(k).kind == "table"), None)
            if spec is None or VERSION_TITLE_RE.search(" ".join(spec.columns)):
                continue
            score = abs(len(spec.columns) - len(columns or [])) if columns else 0
        elif manifest is not None:
            # The roomiest text box makes the best prototype for a text slide.
            spec = next((manifest.field(k) for k in section.fields if manifest.field(k) and manifest.field(k).kind in ("text", "bullets")), None)
            score = -max((b.max_chars or 0) for b in spec.bindings) if spec is not None and spec.bindings else 0
        if best_score is None or score < best_score:
            best, best_score = section.key, score
    if best or not allow_composite:
        return best
    return next((s.key for s in blueprint.sections if s.kind == "composite"), "")


def _before(blueprint: Blueprint) -> str:
    for section in blueprint.sections:
        if EFFORT_RE.search(section.title):
            return section.key
    return ""
