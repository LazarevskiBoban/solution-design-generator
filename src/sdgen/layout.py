from __future__ import annotations

from dataclasses import dataclass, field
from functools import reduce

from pptx.util import Inches

from sdgen.inventory import find_shape

BAND_PAD = Inches(0.05)
HEADER_MAX_HEIGHT = Inches(0.45)
HEADER_GAP = Inches(0.15)
GAP_TIE = Inches(0.05)
CONTAIN_TOL = Inches(0.05)
BOTTOM_BAND = Inches(0.6)
BOTTOM_MARGIN = Inches(0.35)
FULL_WIDTH = 0.8


@dataclass(frozen=True)
class Box:
    left: int
    top: int
    right: int
    bottom: int

    @property
    def width(self) -> int:
        return self.right - self.left

    @property
    def height(self) -> int:
        return self.bottom - self.top

    def union(self, other: Box) -> Box:
        return Box(min(self.left, other.left), min(self.top, other.top), max(self.right, other.right), max(self.bottom, other.bottom))

    def contains(self, other: Box, tol: int = CONTAIN_TOL) -> bool:
        return self.left - tol <= other.left and self.top - tol <= other.top and other.right <= self.right + tol and other.bottom <= self.bottom + tol

    def overlap_x(self, other: Box) -> int:
        return max(0, min(self.right, other.right) - max(self.left, other.left))

    def overlaps_y(self, other: Box) -> bool:
        return self.top < other.bottom and other.top < self.bottom


@dataclass
class Block:
    """A header bar, a container and what sits inside it: the unit that moves or grows together."""

    members: list[int]
    box: Box
    fields: set[str] = field(default_factory=set)
    header: int | None = None


@dataclass
class SlideLayout:
    title: int | None
    top_band: list[int]
    bottom_band: list[int]
    blocks: list[Block]
    content: Box | None
    bound: dict[int, str] = field(default_factory=dict)
    floor: int = 0  # how far down a block may grow: above the footer band, else near the slide bottom

    def block_of(self, shape_id: int) -> Block | None:
        return next((b for b in self.blocks if shape_id in b.members), None)

    def top_block(self) -> Block | None:
        return min(self.blocks, key=lambda b: (b.box.top, b.box.left)) if self.blocks else None

    def is_full_width(self, block: Block) -> bool:
        if self.content is None or block.box.width < FULL_WIDTH * self.content.width:
            return False
        return not any(other is not block and other.box.overlaps_y(block.box) for other in self.blocks)


def analyse_slide(slide, bound: dict[int, str], slide_height: int) -> SlideLayout:
    """Groups the top-level shapes of a slide into a title band, a footer band and content blocks.

    `bound` maps shape ids to the field keys filled on this slide; bound shapes are always content.
    """
    shapes = {s.shape_id: s for s in slide.shapes if None not in (s.left, s.top, s.width, s.height)}
    boxes = {i: Box(s.left, s.top, s.left + s.width, s.top + s.height) for i, s in shapes.items()}
    title = slide.shapes.title
    title_id = title.shape_id if title is not None and title.shape_id in shapes else None
    if title_id is not None:
        title_box = boxes[title_id]
        top_band = [i for i, b in boxes.items() if i not in bound and (i == title_id or (b.bottom <= title_box.bottom + BAND_PAD and b.top <= title_box.top + HEADER_GAP))]
    else:
        limit = min((boxes[i].top for i in bound if i in boxes), default=0)
        top_band = [i for i, b in boxes.items() if i not in bound and b.bottom <= limit]
    bottom_band = [i for i, b in boxes.items() if i not in bound and i not in top_band and b.top >= slide_height - BOTTOM_BAND]
    content = [i for i in boxes if i not in top_band and i not in bottom_band]

    parent: dict[int, int | None] = {}
    for i in sorted(content, key=lambda i: -(boxes[i].width * boxes[i].height)):
        parent[i] = next((p for p in parent if parent[p] is None and _can_contain(shapes[p], p in bound) and boxes[p].contains(boxes[i])), None)
    parentless = [i for i in content if parent[i] is None]

    attached: dict[int, int] = {}
    for h in parentless:
        if h in bound or boxes[h].height > HEADER_MAX_HEIGHT or not _has_text(shapes[h]):
            continue
        best = None
        for s in parentless:
            if s == h or s in attached:
                continue
            gap = boxes[s].top - boxes[h].bottom
            if not -CONTAIN_TOL <= gap <= HEADER_GAP or boxes[h].overlap_x(boxes[s]) < 0.5 * min(boxes[h].width, boxes[s].width):
                continue
            key = (gap // GAP_TIE, boxes[s].left)
            if best is None or key < best[0]:
                best = (key, s)
        if best is not None:
            attached[h] = best[1]

    blocks: list[Block] = []
    for s in parentless:
        if s in attached:
            continue
        members = [s] + [c for c in content if parent[c] == s] + [h for h, t in attached.items() if t == s]
        box = reduce(Box.union, (boxes[m] for m in members))
        header = next((h for h, t in attached.items() if t == s), None)
        blocks.append(Block(members=members, box=box, fields={bound[m] for m in members if m in bound}, header=header))
    content_box = reduce(Box.union, (b.box for b in blocks)) if blocks else None
    if bottom_band:
        floor = min(boxes[i].top for i in bottom_band) - BAND_PAD
    else:
        floor = max(content_box.bottom if content_box else 0, slide_height - BOTTOM_MARGIN)
    return SlideLayout(title=title_id, top_band=top_band, bottom_band=bottom_band, blocks=blocks, content=content_box, bound=dict(bound), floor=floor)


def shift_shapes(slide, shape_ids: list[int], dy: int) -> None:
    for shape_id in shape_ids:
        shape = find_shape(slide, shape_id)
        if shape is not None:
            shape.top = shape.top + dy


def _can_contain(shape, is_bound: bool) -> bool:
    return bool(getattr(shape, "has_text_frame", False)) and not is_bound


def _has_text(shape) -> bool:
    return bool(getattr(shape, "has_text_frame", False)) and bool(shape.text_frame.text.strip())
