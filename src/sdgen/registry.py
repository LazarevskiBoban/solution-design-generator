from __future__ import annotations

import re
import shutil
from pathlib import Path

from pptx import Presentation
from pydantic import BaseModel

from sdgen.analyze import analyze_deck
from sdgen.blueprint import Blueprint, derive_blueprint, mark_static_fields
from sdgen.content import Content, dump_markdown, load_markdown_file
from sdgen.diagrams import add_image_slots
from sdgen.inventory import inspect_deck
from sdgen.manifest import Manifest
from sdgen.tokenize import capture_content, tokenize_deck

MANIFEST_FILE = "manifest.yaml"
BLUEPRINT_FILE = "blueprint.yaml"
TEMPLATE_FILE = "template.pptx"
SOURCE_FILE = "source.pptx"
ORIGINAL_FILE = "original.md"
NAME_RE = re.compile(r"[^a-z0-9_-]+")


def safe_name(name: str) -> str:
    cleaned = NAME_RE.sub("-", name.strip().lower()).strip("-")
    if not cleaned:
        raise ValueError("template name must contain letters or digits")
    return cleaned


class TemplateEntry(BaseModel):
    name: str
    directory: Path
    manifest: Manifest
    blueprint: Blueprint | None = None
    original: Content | None = None

    @property
    def template_path(self) -> Path:
        return self.directory / self.manifest.source

    @property
    def source_path(self) -> Path:
        return self.directory / SOURCE_FILE


class Registry:
    def __init__(self, root: str | Path = "templates") -> None:
        self.root = Path(root)

    def names(self) -> list[str]:
        if not self.root.is_dir():
            return []
        return sorted(p.name for p in self.root.iterdir() if (p / MANIFEST_FILE).is_file())

    def remove(self, name: str) -> None:
        directory = self.root / name
        if not (directory / MANIFEST_FILE).is_file():
            raise FileNotFoundError(f"template '{name}' not found under {self.root}")
        shutil.rmtree(directory)

    def load(self, name: str) -> TemplateEntry:
        directory = self.root / name
        manifest_path = directory / MANIFEST_FILE
        if not manifest_path.is_file():
            raise FileNotFoundError(f"template '{name}' not found under {self.root}")
        blueprint_path = directory / BLUEPRINT_FILE
        blueprint = Blueprint.load(blueprint_path) if blueprint_path.is_file() else None
        manifest = Manifest.load(manifest_path)
        original_path = directory / ORIGINAL_FILE
        original = load_markdown_file(original_path, manifest) if original_path.is_file() else None
        return TemplateEntry(name=name, directory=directory, manifest=manifest, blueprint=blueprint, original=original)

    def add(
        self,
        name: str,
        deck_path: str | Path,
        manifest: Manifest,
        tokenize: bool = True,
        blueprint: Blueprint | None = None,
    ) -> TemplateEntry:
        name = safe_name(name)
        directory = self.root / name
        directory.mkdir(parents=True, exist_ok=True)
        source = directory / SOURCE_FILE
        if Path(deck_path).resolve() != source.resolve():
            shutil.copyfile(deck_path, source)
        manifest = manifest.model_copy(update={"name": name, "source": TEMPLATE_FILE})
        if blueprint is not None:
            manifest = mark_static_fields(manifest, blueprint)
        original = None
        if tokenize:
            prs = Presentation(str(source))
            original = capture_content(prs, manifest)
            manifest = tokenize_deck(prs, manifest)
            if blueprint is not None:
                manifest, blueprint = add_image_slots(prs, manifest, blueprint)
            prs.save(str(directory / TEMPLATE_FILE))
            (directory / ORIGINAL_FILE).write_text(dump_markdown(original, manifest), encoding="utf-8")
        else:
            shutil.copyfile(source, directory / TEMPLATE_FILE)
        manifest.save(directory / MANIFEST_FILE)
        if blueprint is not None:
            blueprint = blueprint.model_copy(update={"name": name})
            blueprint.save(directory / BLUEPRINT_FILE)
        return TemplateEntry(name=name, directory=directory, manifest=manifest, blueprint=blueprint, original=original)

    def reanalyze(self, name: str) -> TemplateEntry:
        """Analyses the stored original again, keeping the section and field names chosen before."""
        current = self.load(name)
        if not current.source_path.is_file():
            raise FileNotFoundError(f"template '{name}' has no stored original; upload the deck again")
        deck = inspect_deck(str(current.source_path))
        analysis = analyze_deck(deck)
        manifest = _carry_fields(analysis.to_manifest(name), current.manifest)
        blueprint = derive_blueprint(deck, analysis, manifest, name)
        if current.blueprint is not None:
            blueprint = _carry_sections(blueprint, current.blueprint)
        return self.add(name, current.source_path, manifest, blueprint=blueprint)

    def save_blueprint(self, name: str, blueprint: Blueprint) -> None:
        directory = self.root / name
        if not directory.is_dir():
            raise FileNotFoundError(f"template '{name}' not found under {self.root}")
        blueprint.save(directory / BLUEPRINT_FILE)

    def save(self, name: str, manifest: Manifest) -> None:
        directory = self.root / name
        if not directory.is_dir():
            raise FileNotFoundError(f"template '{name}' not found under {self.root}")
        manifest.save(directory / MANIFEST_FILE)


def _carry_fields(fresh: Manifest, stored: Manifest) -> Manifest:
    """Fresh analysis with the keys, labels and prefixes the user gave the same shapes before."""
    by_place = {_place(b): f for f in stored.fields for b in f.bindings}
    used: set[str] = set()
    fields = []
    for spec in fresh.fields:
        match = next((by_place[_place(b)] for b in spec.bindings if _place(b) in by_place), None)
        if match is not None and match.key not in used:
            prefixes = {_place(b): b.keep_prefix for b in match.bindings}
            bindings = [b.model_copy(update={"keep_prefix": prefixes.get(_place(b), b.keep_prefix)}) for b in spec.bindings]
            spec = spec.model_copy(update={"key": match.key, "label": match.label, "bindings": bindings})
        elif spec.key in used:
            spec = spec.model_copy(update={"key": _free_key(spec.key, used)})
        used.add(spec.key)
        fields.append(spec)
    slides = fresh.slides.model_copy(
        update={"exclude": sorted(set(fresh.slides.exclude) | set(stored.slides.exclude)), "prototypes": {**fresh.slides.prototypes, **stored.slides.prototypes}}
    )
    return fresh.model_copy(update={"fields": fields, "slides": slides})


def _carry_sections(fresh: Blueprint, stored: Blueprint) -> Blueprint:
    by_slide = {s.slide: s for s in stored.sections}
    used: set[str] = set()
    sections = []
    for section in fresh.sections:
        old = by_slide.get(section.slide)
        if old is not None:
            update = {"key": old.key, "title": old.title, "kind": old.kind, "ask": old.ask, "optional": old.optional}
            if old.kind == "diagram" and section.images == 0:
                update["images"] = 1
            section = section.model_copy(update=update)
        if section.key in used:
            section = section.model_copy(update={"key": _free_key(section.key, used)})
        used.add(section.key)
        sections.append(section)
    return fresh.model_copy(update={"sections": sections})


def _place(binding) -> tuple:
    """Where a binding lives; token bindings share a shape, so the token tells them apart."""
    return (binding.slide, binding.shape.id, binding.token if binding.mode == "token" else None)


def _free_key(key: str, used: set[str]) -> str:
    candidate, n = key, 2
    while candidate in used:
        candidate = f"{key}_{n}"
        n += 1
    return candidate
