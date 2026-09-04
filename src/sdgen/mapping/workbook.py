from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sdgen.mapping.model import MappingSet

COLUMNS = ["Target field", "Type", "Required", "Source field", "Transformation rule", "Example", "Notes"]
WIDTHS = [42, 16, 10, 42, 40, 28, 40]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SHEET_NAME_RE = re.compile(r"[\[\]:*?/\\]")


def write_workbook(mapping: MappingSet, path: str | Path) -> Path:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_summary(summary, mapping)
    used = {"Summary"}
    for source in mapping.sources:
        sheet = workbook.create_sheet(_sheet_name(source.name, used))
        _write_source_sheet(sheet, mapping, source.name)
    path = Path(path)
    workbook.save(str(path))
    return path


def _write_summary(sheet, mapping: MappingSet) -> None:
    sheet.append(["Mapping", mapping.name])
    sheet.append(["Target", mapping.target.name, mapping.target.file])
    sheet.append(["Target fields", len(mapping.target.fields)])
    sheet.append([])
    sheet.append(["Source", "File", "Mapped", "Target fields", "Required fields still unmapped"])
    for cell in sheet[5]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for item in mapping.summary():
        source = mapping.source(item.source)
        sheet.append([item.source, source.file if source else "", item.mapped, item.total, ", ".join(item.unmapped_required)])
    for index, width in enumerate([28, 40, 10, 14, 80], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet["A1"].font = Font(bold=True)


def _write_source_sheet(sheet, mapping: MappingSet, source_name: str) -> None:
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    entries = mapping.entries_for(source_name)
    for field in mapping.target.fields:
        entry = entries.get(field.path)
        sheet.append(
            [
                field.path,
                field.type,
                "yes" if field.required else ("no" if field.required is False else ""),
                entry.source_path if entry else "",
                entry.rule if entry else "",
                entry.example if entry else field.example,
                entry.note if entry else field.description,
            ]
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(sheet.max_row, 1)}"


def _sheet_name(name: str, used: set[str]) -> str:
    base = SHEET_NAME_RE.sub("-", name).strip()[:28] or "Source"
    candidate, n = base, 2
    while candidate in used:
        candidate = f"{base[:25]} {n}"
        n += 1
    used.add(candidate)
    return candidate
