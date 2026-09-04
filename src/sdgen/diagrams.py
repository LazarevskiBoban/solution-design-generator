from __future__ import annotations

import io

from PIL import Image, ImageDraw, ImageFont
from pptx.shapes.picture import Picture
from pptx.util import Emu, Inches

from sdgen.blueprint import Blueprint
from sdgen.inventory import walk_shapes
from sdgen.manifest import Binding, FieldSpec, Manifest, ShapeRef

RELATIONSHIP_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
SLOT_PREFIX = "Diagram slot: "
MIN_SLOT_INCHES = 2.0
MARGIN = Inches(0.42)
CONTENT_TOP = Inches(1.0)
BOTTOM_MARGIN = Inches(0.45)
PIXELS_PER_INCH = 96


def add_image_slots(prs, manifest: Manifest, blueprint: Blueprint) -> tuple[Manifest, Blueprint]:
    slides = list(prs.slides)
    fields = list(manifest.fields)
    sections = []
    for section in blueprint.sections:
        if section.kind != "diagram" or not 1 <= section.slide <= len(slides):
            sections.append(section)
            continue
        slide = slides[section.slide - 1]
        bound = {b.shape.id for f in fields for b in f.bindings if b.slide == section.slide}
        removable = [s for s in slide.shapes if not _keep(s, bound)]
        drawn = [s for s in removable if not isinstance(s, Picture)] or removable
        box = _bounding_box(drawn, prs, slide.shapes.title)
        _remove(slide, removable)
        picture = slide.shapes.add_picture(_placeholder_png(section.title, box), *box)
        picture.name = f"{SLOT_PREFIX}{section.title}"

        key = _unique_key(f"{section.key}_diagram", {f.key for f in fields})
        fields.append(
            FieldSpec(
                key=key,
                label=f"{section.title} diagram",
                kind="image",
                guidance="PNG or JPG. Several images produce one slide each.",
                bindings=[Binding(slide=section.slide, shape=ShapeRef(id=picture.shape_id, name=picture.name), fit="contain")],
            )
        )
        sections.append(section.model_copy(update={"fields": section.fields + [key], "images": max(section.images, 1)}))
    return manifest.model_copy(update={"fields": fields}), blueprint.model_copy(update={"sections": sections})


def _keep(shape, bound: set[int]) -> bool:
    if shape.is_placeholder or shape.shape_id in bound:
        return True
    return any(child.shape_id in bound for child in walk_shapes([shape]))


def _bounding_box(shapes, prs, title=None) -> tuple[int, int, int, int]:
    boxes = [(s.left, s.top, s.left + s.width, s.top + s.height) for s in shapes if None not in (s.left, s.top, s.width, s.height)]
    default = (MARGIN, CONTENT_TOP, prs.slide_width - 2 * MARGIN, prs.slide_height - CONTENT_TOP - BOTTOM_MARGIN)
    if not boxes:
        return default
    left = max(min(b[0] for b in boxes), 0)
    top = max(min(b[1] for b in boxes), 0)
    right = min(max(b[2] for b in boxes), prs.slide_width)
    bottom = min(max(b[3] for b in boxes), prs.slide_height)
    if title is not None and None not in (title.top, title.height):
        top = max(top, title.top + title.height + Inches(0.1))
    if right - left < Inches(MIN_SLOT_INCHES) or bottom - top < Inches(MIN_SLOT_INCHES):
        return default
    return left, top, right - left, bottom - top


def _remove(slide, shapes) -> None:
    rids: set[str] = set()
    for shape in shapes:
        element = shape._element
        for node in element.iter():
            for attr, value in node.attrib.items():
                if attr.startswith(RELATIONSHIP_NS):
                    rids.add(value)
        element.getparent().remove(element)
    for rid in rids:
        if rid in slide.part.rels:
            slide.part.drop_rel(rid)


def _placeholder_png(title: str, box: tuple[int, int, int, int]) -> io.BytesIO:
    width = max(int(Emu(box[2]).inches * PIXELS_PER_INCH), 200)
    height = max(int(Emu(box[3]).inches * PIXELS_PER_INCH), 100)
    image = Image.new("RGB", (width, height), (236, 240, 245))
    draw = ImageDraw.Draw(image)
    draw.rectangle([2, 2, width - 3, height - 3], outline=(150, 160, 175), width=2)
    try:
        font = ImageFont.load_default(size=max(14, height // 18))
        small = ImageFont.load_default(size=max(11, height // 28))
    except TypeError:
        font = small = ImageFont.load_default()
    lines = [(f"Diagram: {title}", font), ("Upload an image for this section", small)]
    y = height / 2 - sum(_text_height(draw, t, f) for t, f in lines) / 2
    for text, used in lines:
        w = _text_width(draw, text, used)
        draw.text(((width - w) / 2, y), text, fill=(70, 80, 95), font=used)
        y += _text_height(draw, text, used) + 6
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def _text_width(draw, text, font) -> float:
    left, _, right, _ = draw.textbbox((0, 0), text, font=font)
    return right - left


def _text_height(draw, text, font) -> float:
    _, top, _, bottom = draw.textbbox((0, 0), text, font=font)
    return bottom - top


def _unique_key(key: str, used: set[str]) -> str:
    candidate, n = key, 2
    while candidate in used:
        candidate = f"{key}_{n}"
        n += 1
    return candidate
