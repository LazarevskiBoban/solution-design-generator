from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Literal

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches
from pydantic import BaseModel, Field

from sdgen.blueprint import cap_title, clean_title
from sdgen.check import FOOTER_TYPES, check_presentation
from sdgen.content import Content, ImageValue, detail_key, images_of, parse_pipe_table
from sdgen.diagrams import remove_shapes
from sdgen.fill.image import replace_picture
from sdgen.flow import FlowSpec, draw_flow
from sdgen.fill.slides import clone_slide, move_slide, remove_slide
from sdgen.fill.table import append_rows, clear_table_body, fill_table, fit_columns, has_footer, header_texts, resize_columns, row_heights
from sdgen.fill.text import Block, Span, capacity_chars_of, capacity_lines_of, fit_text_shape, line_chars_of, overflow_ratio, parse_blocks, replace_literal_everywhere, replace_token, set_rich_text, strip_leading_label, theme_fonts
from sdgen.inventory import find_shape, walk_shapes
from sdgen.layout import BLOCK_GAP, CONTAIN_TOL, Box, SlideLayout, analyse_slide, pin_geometry, shift_shapes
from sdgen.manifest import Binding, FieldSpec, Manifest

CONTINUATION_SUFFIX = " (cont.)"
SENTENCE_RE = re.compile(r"(?<=[.!?])\s+")
SPILL_RATIO = 1.25
SMALL_BOX_LINES = 4
PICTURE_OVERLAP = 0.3
SMALL_BOX_SHRINK_RATIO = 1.5
STALE_BUDGET_RATIO = 1.25
ROW_TOL = Inches(0.35)  # blocks whose tops differ by less sit on the same row of the slide
MissingMode = Literal["keep", "blank", "placeholder"]
PLACEHOLDER_ROW = "[To be completed]"


def placeholder_text(label: str) -> str:
    return f"[To be completed: {label}]"


class RenderIssue(BaseModel):
    level: Literal["info", "warning", "error"] = "warning"
    field: str | None = None
    shape: str = ""
    slide: int | None = None
    message: str

    def __str__(self) -> str:
        where = f" [slide {self.slide}]" if self.slide else ""
        who = f" {self.field}:" if self.field else ""
        return f"{self.level}{where}{who} {self.message}"


class ExtraSlide(BaseModel):
    """A slide that does not exist in the template: a clone of a prototype slide with one filled field."""

    key: str
    title: str
    spec: FieldSpec  # bound to the prototype slide and the shape to fill
    value: Any = None
    before: int = 0  # template slide number to insert in front of; 0 appends at the end
    after: int = 0  # template slide number to follow, past its copies and earlier detail slides; wins over `before`


class RenderResult(BaseModel):
    output: str
    slides: int
    issues: list[RenderIssue] = Field(default_factory=list)
    slide_map: list[int] = Field(default_factory=list)  # template slide number behind each output slide
    slide_keys: list[str] = Field(default_factory=list)  # extra-slide key per output slide, empty for template slides

    @property
    def errors(self) -> list[RenderIssue]:
        return [i for i in self.issues if i.level == "error"]


def render(
    template: str | Path,
    manifest: Manifest,
    content: Content,
    output: str | Path,
    missing: MissingMode = "placeholder",
    continue_on: set[int] | list[int] | None = None,
    field_modes: dict[str, MissingMode] | None = None,
    hidden: list[int] | set[int] | None = None,
    order: list[int] | None = None,
    titles: dict[int, str] | None = None,
    extras: list[ExtraSlide] | None = None,
    flows: dict[str, FlowSpec] | None = None,
    spill: bool = False,
    clear_shapes: list[tuple[int, int]] | None = None,
    subject_slides: set[int] | list[int] | None = None,
    details: dict[str, FieldSpec] | None = None,
    check: bool = True,
) -> RenderResult:
    prs = Presentation(str(template))
    slides = list(prs.slides)
    issues: list[RenderIssue] = []
    hidden_slides = {n for n in (hidden or []) if 1 <= n <= len(slides)}
    numbers = {slide.slide_id: n for n, slide in enumerate(slides, 1)}
    _reorder(prs, slides, order)
    prototypes = set(manifest.slides.prototypes.values())
    if continue_on is not None:
        prototypes |= set(continue_on)
    theme = theme_fonts(slides[0].part) if slides else ("Arial", "Arial")

    keep_subject = {1} if subject_slides is None else set(subject_slides)
    for spec in manifest.globals:
        value = content.globals.get(spec.key, "")
        stripped = _strip_subject_from_titles(slides, spec.replaces, keep_subject) if spec.key == "subject" else 0
        if not value:
            issues.append(RenderIssue(field=spec.key, message=f"no value; '{spec.replaces}' left in place"))
            continue
        hits = replace_literal_everywhere(prs, spec.replaces, value)
        if hits == 0 and not stripped:
            issues.append(RenderIssue(field=spec.key, message=f"'{spec.replaces}' not found in the template"))

    for number, title in (titles or {}).items():
        if not 1 <= number <= len(slides) or not title.strip():
            continue
        shape = slides[number - 1].shapes.title
        if shape is None:
            issues.append(RenderIssue(slide=number, message="no title placeholder to retitle"))
            continue
        set_rich_text(shape, _capped(title, issues, number))
        fit_text_shape(shape, theme=theme)

    # Cleared before anything is filled or copied, so continuation copies inherit the cleared state.
    for number, shape_id in clear_shapes or []:
        if not 1 <= number <= len(slides):
            continue
        target = find_shape(slides[number - 1], shape_id)
        if target is None or not getattr(target, "has_text_frame", False):
            issues.append(RenderIssue(slide=number, message=f"text to clear not found (shape {shape_id})"))
            continue
        set_rich_text(target, "")
        issues.append(RenderIssue(level="info", slide=number, message=f"template text cleared (shape {shape_id})"))

    skipped_slides = hidden_slides | {n for n in manifest.slides.exclude if 1 <= n <= len(slides)}
    drawn = _draw_flows(slides, manifest, content, flows or {}, skipped_slides, issues)
    kept = {k for k, m in (field_modes or {}).items() if m == "keep"} | {f.key for f in manifest.fields if f.static}
    layouts = _layouts(prs, slides, manifest, skipped_slides, issues)
    image_ids: dict[int, set[int]] = {}
    for spec in manifest.fields:
        if spec.kind == "image":
            for binding in spec.bindings:
                image_ids.setdefault(binding.slide, set()).add(binding.shape.id)
    pending: dict[int, dict] = {}
    wanted: dict[int, list[tuple]] = {}  # composite fields with more to say than their box holds
    for spec in manifest.fields:
        if spec.key in drawn:
            continue
        value = content.fields.get(spec.key)
        mode = (field_modes or {}).get(spec.key)
        if spec.static and mode != "blank":
            mode = "keep"
        if mode == "keep":
            issues.append(RenderIssue(level="info", field=spec.key, message="template content kept"))
            continue
        if mode == "blank":
            value = None
        written = value not in (None, "", []) and spec.kind in ("text", "bullets")
        mode = mode or missing
        if value is None or value == "" or value == []:
            if mode == "keep" or spec.kind == "image":
                level = "info" if spec.kind == "image" else "warning"
                issues.append(RenderIssue(level=level, field=spec.key, message="no value; template content left in place"))
                continue
            if mode == "blank":
                value = [] if spec.kind == "table" else ""
            else:
                value = [[PLACEHOLDER_ROW]] if spec.kind == "table" else placeholder_text(spec.label)
                issues.append(RenderIssue(level="info", field=spec.key, message="no value; placeholder shown"))
        for binding in spec.bindings:
            if binding.slide in skipped_slides:
                continue
            if not 1 <= binding.slide <= len(slides):
                issues.append(RenderIssue(level="error", field=spec.key, slide=binding.slide, message="slide does not exist"))
                continue
            slide = slides[binding.slide - 1]
            shape = _locate(slide, binding)
            if shape is None:
                issues.append(RenderIssue(level="error", field=spec.key, slide=binding.slide, message=f"shape {binding.shape.id} ({binding.shape.name}) not found"))
                continue
            try:
                detailed = spec.key in (details or {})
                allow = not detailed and (binding.slide in prototypes or (continue_on is None and shape.is_placeholder))
                room = _room(layouts.get(binding.slide), shape) if spec.kind == "table" else None
                remaining = _apply(prs, slide, shape, spec, binding, value, issues, allow, spill=spill or detailed, theme=theme, room=room)
                if written and binding.mode == "replace":
                    _remove_pictures_under(slide, shape, layouts.get(binding.slide), image_ids.get(binding.slide, set()), issues, binding.slide)
                if detailed:
                    if remaining and spec.kind != "table":
                        _trim_to_box(shape, spec, binding, remaining, theme, issues)
                    full = content.fields.get(detail_key(spec.key))
                    if remaining or full not in (None, "", []):
                        wanted.setdefault(binding.slide, []).append((spec, binding, details[spec.key], full if full not in (None, "", []) else value))
                    if remaining:
                        issues.append(RenderIssue(level="info", field=spec.key, slide=binding.slide, message="box shows the first part; the full content follows on a detail slide"))
                elif remaining:
                    pending.setdefault(binding.slide, {})[spec.key] = (spec, binding, remaining)
            except Exception as exc:  # keep rendering the rest of the document
                issues.append(RenderIssue(level="error", field=spec.key, slide=binding.slide, message=str(exc)))

    for number, fields_pending in pending.items():
        _continue_slide(prs, slides[number - 1], layouts.get(number), manifest, number, fields_pending, theme, issues, kept)

    extra_ids = _add_extras(prs, slides, _detail_slides(slides, layouts, wanted) + list(extras or []), issues, theme, numbers)

    for index in sorted(set(manifest.slides.exclude) | hidden_slides, reverse=True):
        if 1 <= index <= len(slides):
            remove_slide(prs, slides[index - 1])
        else:
            issues.append(RenderIssue(slide=index, message="excluded slide does not exist"))

    slide_map = _slide_map(prs, numbers)
    slide_keys = [extra_ids.get(slide.slide_id, "") for slide in prs.slides]
    if check:
        drawn_on: dict[int, FlowSpec] = {}
        for key in drawn:
            spec = manifest.field(key)
            template_number = spec.bindings[0].slide if spec is not None and spec.bindings else 0
            for position, (number, slide_key) in enumerate(zip(slide_map, slide_keys), 1):
                if number == template_number and not slide_key:
                    drawn_on[position] = (flows or {})[key]
        for finding in check_presentation(prs, flows=drawn_on):
            issues.append(RenderIssue(level=finding.level, slide=finding.slide, shape=finding.shape, message=f"check {finding.code}: {finding.message}"))
    prs.save(str(output))
    return RenderResult(output=str(output), slides=len(prs.slides), issues=issues, slide_map=slide_map, slide_keys=slide_keys)


def _strip_subject_from_titles(slides: list, replaces: str, keep: set[int]) -> int:
    """Titles outside the cover lose the subject marker, so the section title stands alone."""
    marker = " ".join(replaces.split())
    count = 0
    for number, slide in enumerate(slides, 1):
        title = slide.shapes.title
        if number in keep or title is None or not title.has_text_frame:
            continue
        text = title.text_frame.text
        if marker not in " ".join(text.split()):
            continue
        set_rich_text(title, cap_title(clean_title(text, marker)))
        count += 1
    return count


def _capped(title: str, issues: list[RenderIssue], slide: int | None = None) -> str:
    capped = cap_title(title)
    if capped != " ".join(title.split()):
        issues.append(RenderIssue(level="info", slide=slide, message=f"title shortened to '{capped}'"))
    return capped


def _set_header(shape, columns: list[str]) -> None:
    """Renames the header row of a cloned table so an extra slide shows its own columns."""
    cells = list(shape.table.rows[0].cells)
    for index, cell in enumerate(cells):
        set_rich_text(cell, columns[index] if index < len(columns) else "")


def _draw_flows(slides: list, manifest: Manifest, content: Content, flows: dict[str, FlowSpec], hidden: set[int], issues: list[RenderIssue]) -> set[str]:
    """Draws each flow into its image slot unless an uploaded image fills that slot."""
    drawn: set[str] = set()
    for key, flow in flows.items():
        spec = manifest.field(key)
        if spec is None or spec.kind != "image":
            issues.append(RenderIssue(field=key, message="flow refers to a field that is not an image slot"))
            continue
        if content.fields.get(key) not in (None, "", []) or not flow.nodes:
            continue
        for binding in spec.bindings:
            if binding.slide in hidden or not 1 <= binding.slide <= len(slides):
                continue
            slide = slides[binding.slide - 1]
            shape = _locate(slide, binding)
            if shape is None:
                issues.append(RenderIssue(field=key, slide=binding.slide, message="image slot not found for the flow"))
                continue
            draw_flow(slide, (shape.left, shape.top, shape.width, shape.height), flow, prefix=f"Flow {key}")
            remove_shapes(slide, [shape])
            issues.append(RenderIssue(level="info", field=key, slide=binding.slide, message=f"diagram drawn from the brief ({len(flow.nodes)} nodes)"))
            drawn.add(key)
    return drawn


def _trim_to_box(shape, spec: FieldSpec, binding: Binding, blocks: list, theme, issues: list[RenderIssue]) -> None:
    """The box keeps the first paragraphs that fit; the rest belongs to the detail slide."""
    chunks = _split_blocks(blocks, _capacity(shape, binding, blocks, theme), line_chars_of(shape, theme))
    set_rich_text(shape, chunks[0], keep_prefix=binding.keep_prefix)
    _shrink(shape, spec, binding, issues, theme)


def _detail_slides(slides: list, layouts: dict[int, SlideLayout], wanted: dict[int, list[tuple]]) -> list[ExtraSlide]:
    """One plain slide per composite field with more to say, in the reading order of the blocks on its slide."""
    result: list[ExtraSlide] = []
    for number in sorted(wanted):
        layout = layouts.get(number)
        placed = []
        for spec, binding, proto, full in wanted[number]:
            shape = find_shape(slides[number - 1], binding.shape.id)
            block = layout.block_of(binding.shape.id) if layout is not None else None
            if block is not None:
                top, left = block.box.top, block.box.left
            elif shape is not None and None not in (shape.top, shape.left):
                top, left = shape.top, shape.left
            else:
                top, left = 0, 0
            placed.append((top, left, spec, proto, full))
        rows: list[list] = []
        for item in sorted(placed, key=lambda p: (p[0], p[1])):
            if rows and abs(item[0] - rows[-1][0][0]) <= ROW_TOL:
                rows[-1].append(item)
            else:
                rows.append([item])
        for row in rows:
            for _, _, spec, proto, full in sorted(row, key=lambda p: p[1]):
                result.append(ExtraSlide(key=detail_key(spec.key), title=cap_title(spec.label), spec=proto, value=full, after=number))
    return result


def _add_extras(prs, slides: list, extras: list[ExtraSlide], issues: list[RenderIssue], theme: tuple[str, str] | None = None, numbers: dict[int, int] | None = None) -> dict[int, str]:
    ids: dict[int, str] = {}
    for extra in extras:
        binding = extra.spec.bindings[0] if extra.spec.bindings else None
        if binding is None or not 1 <= binding.slide <= len(slides):
            issues.append(RenderIssue(level="error", field=extra.key, message="extra slide has no prototype slide"))
            continue
        clone = clone_slide(prs, slides[binding.slide - 1])
        ids[clone.slide_id] = extra.key
        title_shape = clone.shapes.title
        if title_shape is not None:
            set_rich_text(title_shape, _capped(extra.title, issues))
            fit_text_shape(title_shape, theme=theme)
        shape = find_shape(clone, binding.shape.id)
        copies: list = []
        if shape is None:
            issues.append(RenderIssue(level="error", field=extra.key, message=f"shape {binding.shape.id} not found on the prototype slide"))
        else:
            value = extra.value
            if value in (None, "", []):
                value = [[PLACEHOLDER_ROW]] if extra.spec.kind == "table" else placeholder_text(extra.title)
                issues.append(RenderIssue(level="info", field=extra.key, message="no value; placeholder shown"))
            try:
                if extra.spec.kind == "table" and extra.spec.columns and getattr(shape, "has_table", False):
                    if len(extra.spec.columns) != len(shape.table.columns):
                        try:
                            resize_columns(shape, len(extra.spec.columns))
                        except ValueError as exc:
                            issues.append(RenderIssue(level="info", field=extra.key, message=f"table keeps the prototype's {len(shape.table.columns)} columns: {exc}"))
                    if len(extra.spec.columns) == len(shape.table.columns) and [h.strip().lower() for h in header_texts(shape)] != [c.strip().lower() for c in extra.spec.columns]:
                        fit_columns(shape, extra.spec.columns, _as_rows(value), theme)
                    _set_header(shape, extra.spec.columns)
                _drop_leftovers(clone, shape, extra.key, prs.slide_height)
                layout = analyse_slide(clone, {binding.shape.id: extra.key}, prs.slide_height)
                room = _room(layout, shape) if extra.spec.kind == "table" else None
                remaining = _apply(prs, clone, shape, extra.spec, binding, value, issues, False, spill=True, theme=theme, room=room)
                if extra.spec.kind != "table" and extra.value not in (None, "", []):
                    _remove_pictures_under(clone, shape, layout, set(), issues, binding.slide)
                if remaining:
                    copies = _continue_slide(prs, clone, layout, None, binding.slide, {extra.key: (extra.spec, binding, remaining)}, theme, issues, set())
            except Exception as exc:
                issues.append(RenderIssue(level="error", field=extra.key, message=str(exc)))
        chain = [clone] + copies
        for copy in copies:
            ids[copy.slide_id] = extra.key
        chain_ids = {s.slide_id for s in chain}
        others = [s.slide_id for s in prs.slides if s.slide_id not in chain_ids]
        if 1 <= extra.after <= len(slides):
            position = others.index(slides[extra.after - 1].slide_id) + 1
            while position < len(others) and others[position] not in (numbers or {}):
                position += 1  # past the anchor's own copies and the detail slides already placed after it
        else:
            target = slides[extra.before - 1] if 1 <= extra.before <= len(slides) else None
            position = others.index(target.slide_id) if target is not None else len(others)
        follower = others[position] if position < len(others) else None
        for member in chain:
            _move_before(prs, member, follower)
    return ids


def _move_before(prs, slide, follower: int | None) -> None:
    """Puts the slide right before the follower (last without one), wherever it sits now."""
    rest = [s.slide_id for s in prs.slides if s.slide_id != slide.slide_id]
    move_slide(prs, slide, rest.index(follower) if follower is not None else len(rest))


def _drop_leftovers(slide, shape, key: str, slide_height: int) -> None:
    """The prototype's other text and tables belonged to its own section; an extra slide shows only its title and its field."""
    layout = analyse_slide(slide, {shape.shape_id: key}, slide_height)
    bands = set(layout.top_band) | set(layout.bottom_band)
    doomed = []
    for other in slide.shapes:
        if other.shape_id == shape.shape_id or other.shape_id in bands or _placeholder_kind(other) in FOOTER_TYPES:
            continue
        has_text = getattr(other, "has_text_frame", False) and bool(other.text_frame.text.strip())
        if has_text or getattr(other, "has_table", False) or other.is_placeholder:
            doomed.append(other)
    if doomed:
        remove_shapes(slide, doomed)


def _placeholder_kind(shape):
    try:
        return shape.placeholder_format.type
    except (AttributeError, ValueError):
        return None


def _reorder(prs, slides: list, order: list[int] | None) -> None:
    wanted: list[int] = []
    for number in order or []:
        if 1 <= number <= len(slides) and number not in wanted:
            wanted.append(number)
    if not wanted:
        return
    sequence = wanted + [n for n in range(1, len(slides) + 1) if n not in wanted]
    for position, number in enumerate(sequence):
        move_slide(prs, slides[number - 1], position)


def _slide_map(prs, numbers: dict[int, int]) -> list[int]:
    # Clones are inserted right after their source, so they inherit the last seen origin.
    result: list[int] = []
    last = 0
    for slide in prs.slides:
        last = numbers.get(slide.slide_id, last)
        result.append(last)
    return result


def _locate(slide, binding: Binding):
    shape = find_shape(slide, binding.shape.id)
    if shape is not None and (not binding.shape.name or shape.name == binding.shape.name):
        return shape
    if binding.shape.name:
        by_name = [s for s in walk_shapes(slide.shapes) if s.name == binding.shape.name]
        if len(by_name) == 1:
            return by_name[0]
    return shape


def _apply(prs, slide, shape, spec: FieldSpec, binding: Binding, value: Any, issues: list[RenderIssue], prototype: bool, spill: bool = False, theme: tuple[str, str] | None = None, room: int | None = None) -> list:
    if spec.kind == "table":
        if not getattr(shape, "has_table", False):
            raise ValueError("bound shape is not a table")
        rows = _as_rows(value)
        if spec.static and not rows:
            clear_table_body(shape, binding.header_rows)
            return []
        fill_table(shape, rows, header_rows=binding.header_rows, keep_last_row_if=binding.keep_last_row_if, columns=spec.columns or None)
        if spill and room is not None and len(rows) > 1:
            fitting = _rows_that_fit(shape, binding, len(rows), room)
            if fitting < len(rows):
                fill_table(shape, rows[:fitting], header_rows=binding.header_rows, keep_last_row_if=binding.keep_last_row_if, columns=spec.columns or None)
                issues.append(RenderIssue(level="info", field=spec.key, slide=binding.slide, message=f"{len(rows) - fitting} row(s) continue on the next slide"))
                return list(rows[fitting:])
        return []
    if spec.kind == "image":
        images = images_of(value)
        missing = [i.path for i in images if not Path(i.path).is_file()]
        if missing:
            raise FileNotFoundError(f"image not found: {missing[0]}")
        if not images:
            return
        picture = replace_picture(slide, shape, images[0].path, binding.fit)
        current = slide
        for image in images[1:]:
            current = clone_slide(prs, current)
            _mark_continuation(current)
            target = next((s for s in walk_shapes(current.shapes) if s.name == picture.name), None)
            if target is None:
                issues.append(RenderIssue(field=spec.key, slide=binding.slide, message="could not place an extra image"))
                break
            picture = replace_picture(current, target, image.path, binding.fit)
        if len(images) > 1:
            issues.append(RenderIssue(level="info", field=spec.key, slide=binding.slide, message=f"{len(images) - 1} extra slide(s) for images"))
        return
    if not shape.has_text_frame:
        raise ValueError("bound shape has no text")
    text = _as_text(value)
    if binding.mode == "token":
        if not binding.token or replace_token(shape, binding.token, " ".join(text.split())) == 0:
            issues.append(RenderIssue(field=spec.key, slide=binding.slide, message=f"placeholder {binding.token} not found"))
        return

    if binding.keep_prefix:
        text = strip_leading_label(text, [binding.keep_prefix, spec.label])
    blocks = parse_blocks(text)
    set_rich_text(shape, blocks, keep_prefix=binding.keep_prefix)
    capacity = _budget(binding.max_chars, capacity_chars_of(shape, theme, len(binding.keep_prefix or "")))
    ratio = overflow_ratio(shape, theme)
    over = bool(capacity and len(text) > capacity * SPILL_RATIO) or ratio > SPILL_RATIO
    lines = capacity_lines_of(shape, theme)
    if over and lines is not None and lines < SMALL_BOX_LINES and ratio < SMALL_BOX_SHRINK_RATIO:
        over = False  # a small box that is only a little over shrinks instead of spawning a slide
    if over and prototype:
        chunks = _split_blocks(blocks, capacity or len(text), line_chars_of(shape, theme))
        issues.append(RenderIssue(level="info", field=spec.key, slide=binding.slide, message=f"continued on {len(chunks) - 1} extra slide(s)"))
        set_rich_text(shape, chunks[0], keep_prefix=binding.keep_prefix)
        _shrink(shape, spec, binding, issues, theme)
        current = slide
        for chunk in chunks[1:]:
            current = clone_slide(prs, current)
            _mark_continuation(current)
            target = find_shape(current, shape.shape_id)
            set_rich_text(target, chunk, keep_prefix=binding.keep_prefix)
            _shrink(target, spec, binding, issues, theme)
        return []
    if over and spill and (len(blocks) > 1 or len(SENTENCE_RE.split(text.strip())) > 1):
        return blocks
    if capacity and len(text) > capacity:
        issues.append(RenderIssue(level="info", field=spec.key, slide=binding.slide, message=f"text is {len(text)} characters, about {capacity} fit; shrunk to fit"))
    _shrink(shape, spec, binding, issues, theme)
    return []


def _layouts(prs, slides: list, manifest: Manifest, skipped: set[int], issues: list[RenderIssue]) -> dict[int, SlideLayout]:
    """Block layout per slide, taken before filling because tables grow while they are filled."""
    bound: dict[int, dict[int, str]] = {}
    for spec in manifest.fields:
        if spec.kind == "image":
            continue
        for binding in spec.bindings:
            if binding.mode == "replace" and binding.slide not in skipped and 1 <= binding.slide <= len(slides):
                bound.setdefault(binding.slide, {})[binding.shape.id] = spec.key
    layouts: dict[int, SlideLayout] = {}
    for number, shapes in bound.items():
        try:
            layouts[number] = analyse_slide(slides[number - 1], shapes, prs.slide_height)
        except Exception as exc:
            issues.append(RenderIssue(slide=number, message=f"layout not analysed: {exc}"))
    return layouts


def _continue_slide(prs, slide, layout: SlideLayout | None, manifest: Manifest | None, number: int, pending: dict, theme, issues: list[RenderIssue], kept: set[str]) -> list:
    """Continues the boxes that overflowed: grow and push when the top block overflowed, else cleaned copies."""
    try:
        blocks = {key: layout.block_of(binding.shape.id) for key, (_, binding, _) in pending.items()} if layout is not None else {}
        if layout is None or layout.content is None or any(block is None for block in blocks.values()):
            if manifest is None:
                raise ValueError("no block layout for the slide")
            _continue_composite(prs, slide, manifest, number, _chunk_pending(slide, pending, theme, {}, False, issues), issues, kept)
            return []
        top = layout.top_block()
        if top is not None and layout.is_full_width(top) and any(block is top for block in blocks.values()):
            top_pending = {k: v for k, v in pending.items() if blocks[k] is top}
            rest_pending = {k: v for k, v in pending.items() if blocks[k] is not top}
            copies = _expand_and_push(prs, slide, layout, top, top_pending, number, theme, issues)
            if rest_pending and copies:
                # The other overflowing boxes now sit on the pushed slide and continue from there.
                pushed = copies[-1]
                pushed_layout = analyse_slide(pushed, {i: k for i, k in layout.bound.items() if i not in top.members}, prs.slide_height)
                copies += _continue_slide(prs, pushed, pushed_layout, manifest, number, rest_pending, theme, issues, kept)
            return copies
        return _continue_cleaned(prs, slide, layout, pending, number, theme, issues)
    except Exception as exc:
        issues.append(RenderIssue(level="error", slide=number, message=f"could not continue the slide, text shrunk instead: {exc}"))
        for spec, binding, blocks in pending.values():
            shape = find_shape(slide, binding.shape.id)
            if shape is None:
                continue
            if spec.kind == "table":
                append_rows(shape, blocks, header_rows=binding.header_rows, keep_last_row_if=binding.keep_last_row_if, columns=spec.columns or None)
            else:
                set_rich_text(shape, blocks, keep_prefix=binding.keep_prefix)
                _shrink(shape, spec, binding, issues, theme)
        return []


def _expand_and_push(prs, slide, layout: SlideLayout, top, pending: dict, number: int, theme, issues: list[RenderIssue]) -> list:
    """Grows the top block to the bottom of the slide and moves every other block to a copy that follows."""
    others = [b for b in layout.blocks if b is not top]
    other_ids = [i for b in others for i in b.members]
    pushed = clone_slide(prs, slide) if others else None
    remove_shapes(slide, _shapes(slide, other_ids))
    delta = max(0, layout.floor - top.box.bottom)
    scales: dict[int, float] = {}
    for shape_id in _growing_members(slide, top, top.box.bottom, layout.bound):
        shape = find_shape(slide, shape_id)
        scales[shape_id] = (shape.height + delta) / shape.height
        _grow(shape, delta)
    if pushed is not None:
        remove_shapes(pushed, _shapes(pushed, top.members))
        shift_shapes(pushed, other_ids, top.box.top - min(b.box.top for b in others))
        _mark_continuation(pushed)
    chunked = _chunk_pending(slide, pending, theme, {key: scales.get(binding.shape.id, 1.0) for key, (_, binding, _) in pending.items()}, True, issues, _table_rooms(slide, pending, layout.floor, 0, True))
    copies = _spread(prs, slide, chunked, theme, issues)
    note = f"top block grown to the slide bottom, {len(others)} block(s) moved to the next slide"
    issues.append(RenderIssue(level="info", slide=number, message=note + (f", continued on {len(copies)} extra slide(s)" if copies else "")))
    return copies + ([pushed] if pushed is not None else [])


def _continue_cleaned(prs, slide, layout: SlideLayout, pending: dict, number: int, theme, issues: list[RenderIssue]) -> list:
    """Copies keep only the continued blocks, moved to the top of the content area and grown to its bottom."""
    keep: list = []
    for _, binding, _ in pending.values():
        block = layout.block_of(binding.shape.id)
        if not any(block is b for b in keep):
            keep.append(block)
    keep_ids = [i for b in keep for i in b.members]
    drop_ids = [i for b in layout.blocks if not any(b is k for k in keep) for i in b.members]
    shift = min(b.box.top for b in keep) - layout.content.top

    def below_of(block: Block) -> list:
        return [o for o in keep if o is not block and o.box.top >= block.box.bottom - CONTAIN_TOL and o.box.overlap_x(block.box) > 0]

    # A table whose remaining rows need more height pushes the kept blocks under it down, when the floor allows.
    extra: dict[int, int] = {}
    for key, (spec, binding, rows) in pending.items():
        shape = find_shape(slide, binding.shape.id)
        if spec.kind == "table" and shape is not None:
            block = layout.block_of(binding.shape.id)
            extra[id(block)] = max(0, _table_need(shape, binding, spec, rows, theme) + BLOCK_GAP - (block.box.bottom - shape.top))
    push: dict[int, int] = {}
    for b in sorted(keep, key=lambda b: b.box.top):
        for o in below_of(b):
            push[id(o)] = max(push.get(id(o), 0), push.get(id(b), 0) + extra.get(id(b), 0))
    if any(o.box.bottom - shift + push.get(id(o), 0) > layout.floor for o in keep):
        extra, push = {}, {}  # no room to push: the rows continue on further copies instead
    growth: dict[int, int] = {}
    growing: dict[int, list[int]] = {}
    limits: dict[int, int] = {}
    for b in keep:
        below = below_of(b)
        bottom = b.box.bottom - shift + push.get(id(b), 0)
        limits[id(b)] = min(o.box.top - shift + push.get(id(o), 0) for o in below) - BLOCK_GAP if below else layout.floor
        growth[id(b)] = extra.get(id(b), 0) if below else max(0, layout.floor - bottom, extra.get(id(b), 0))
        growing[id(b)] = _growing_members(slide, b, b.box.bottom, layout.bound) if growth[id(b)] > 0 else []
    scales: dict[str, float] = {}
    rooms: dict[str, tuple[int, int]] = {}
    for key, (spec, binding, _) in pending.items():
        block = layout.block_of(binding.shape.id)
        shape = find_shape(slide, binding.shape.id)
        if shape is None:
            continue
        if spec.kind == "table":
            rooms[key] = (0, limits[id(block)] - (shape.top - shift + push.get(id(block), 0)))
            continue
        grows = binding.shape.id in growing[id(block)]
        scales[key] = (shape.height + growth[id(block)]) / shape.height if grows else 1.0

    def prepare(copy) -> None:
        remove_shapes(copy, _shapes(copy, drop_ids))
        shift_shapes(copy, keep_ids, -shift)
        for b in keep:
            if push.get(id(b)):
                shift_shapes(copy, b.members, push[id(b)])
            for shape_id in growing[id(b)]:
                shape = find_shape(copy, shape_id)
                if shape is not None:
                    _grow(shape, growth[id(b)])

    chunked = _chunk_pending(slide, pending, theme, scales, False, issues, rooms)
    copies = _spread(prs, slide, chunked, theme, issues, prepare)
    if copies:
        issues.append(RenderIssue(level="info", slide=number, message=f"continued on {len(copies)} extra slide(s) that show only the continued block(s)"))
    return copies


def _chunk_pending(slide, pending: dict, theme, scales: dict[str, float], grow_first: bool, issues: list[RenderIssue], rooms: dict[str, int] | None = None) -> dict:
    """Splits each overflowing field into chunks and writes the first chunk into its box."""
    chunked: dict = {}
    for key, (spec, binding, blocks) in pending.items():
        shape = find_shape(slide, binding.shape.id)
        if shape is None:
            continue
        if spec.kind == "table":
            # Rows that still fit under the table stay on the slide; the rest is split by the room a copy offers.
            first, rest = (rooms or {}).get(key, (0, 0))
            chunks = _split_rows(shape, binding, spec, blocks, first, rest, theme)
            if chunks[0]:
                append_rows(shape, chunks[0], header_rows=binding.header_rows, keep_last_row_if=binding.keep_last_row_if, columns=spec.columns or None)
            chunked[key] = (spec, binding, chunks)
            continue
        capacity = _capacity(shape, binding, blocks, theme)
        grown = max(1, int(capacity * scales.get(key, 1.0)))
        chunks = _split_blocks(blocks, grown if grow_first else capacity, line_chars_of(shape, theme), grown)
        set_rich_text(shape, chunks[0], keep_prefix=binding.keep_prefix)
        _shrink(shape, spec, binding, issues, theme)
        chunked[key] = (spec, binding, chunks)
    return chunked


def _spread(prs, slide, chunked: dict, theme, issues: list[RenderIssue], prepare=None) -> list:
    """One copy per extra chunk; the first copy is prepared once, later copies clone it."""
    count = max((len(chunks) for _, _, chunks in chunked.values()), default=1)
    copies: list = []
    source = slide
    for index in range(1, count):
        copy = clone_slide(prs, source)
        if index == 1 and prepare is not None:
            prepare(copy)
        _mark_continuation(copy)
        for spec, binding, chunks in chunked.values():
            target = find_shape(copy, binding.shape.id)
            if target is None:
                continue
            if spec.kind == "table":
                fill_table(target, chunks[index] if index < len(chunks) else [], header_rows=binding.header_rows, keep_last_row_if=binding.keep_last_row_if, columns=spec.columns or None)
                continue
            set_rich_text(target, chunks[index] if index < len(chunks) else "", keep_prefix=binding.keep_prefix)
            if index < len(chunks):
                _shrink(target, spec, binding, issues, theme)
        copies.append(copy)
        source = copy
    return copies


def _remove_pictures_under(slide, shape, layout: SlideLayout | None, keep_ids: set[int], issues: list[RenderIssue], number: int) -> None:
    """A template picture under a written box illustrated another project's text; it goes."""
    if None in (shape.left, shape.top, shape.width, shape.height):
        return
    box = Box(shape.left, shape.top, shape.left + shape.width, shape.top + shape.height)
    bands = set(layout.top_band) | set(layout.bottom_band) if layout is not None else set()
    doomed = []
    for other in slide.shapes:
        if other.shape_type != MSO_SHAPE_TYPE.PICTURE or other.shape_id in keep_ids or other.shape_id in bands or other.name.startswith("Flow "):
            continue
        if None in (other.left, other.top, other.width, other.height) or not other.width or not other.height:
            continue
        picture = Box(other.left, other.top, other.left + other.width, other.top + other.height)
        overlap = picture.overlap_x(box) * max(0, min(picture.bottom, box.bottom) - max(picture.top, box.top))
        if overlap >= PICTURE_OVERLAP * picture.width * picture.height:
            doomed.append(other)
    if doomed:
        issues.append(RenderIssue(level="info", slide=number, message="template picture removed under written text (shape " + ", ".join(str(s.shape_id) for s in doomed) + ")"))
        remove_shapes(slide, doomed)


def _room(layout: SlideLayout | None, shape) -> int | None:
    return None if layout is None else layout.limit_below(shape.shape_id) - shape.top


def _rows_that_fit(shape, binding: Binding, count: int, room: int) -> int:
    """Body rows of the filled table that fit the room, at least one."""
    heights = [tr.h for tr in shape.table._tbl.tr_lst]
    footer = 1 if has_footer(shape, binding.keep_last_row_if, binding.header_rows) else 0
    used = sum(heights[:binding.header_rows]) + (heights[-1] if footer else 0)
    fitting = 0
    for height in heights[binding.header_rows:len(heights) - footer]:
        if fitting and used + height > room:
            break
        used += height
        fitting += 1
    return max(1, min(fitting, count))


def _table_need(shape, binding: Binding, spec: FieldSpec, rows: list, theme) -> int:
    """Height a copy of the table needs to show all the rows: header, rows and footer."""
    heights = [tr.h for tr in shape.table._tbl.tr_lst]
    footer = heights[-1] if has_footer(shape, binding.keep_last_row_if, binding.header_rows) else 0
    return sum(heights[: binding.header_rows]) + footer + sum(row_heights(shape, theme, rows=rows, header_rows=binding.header_rows, columns=spec.columns or None))


def _split_rows(shape, binding: Binding, spec: FieldSpec, rows: list, first: int, rest: int, theme) -> list[list]:
    """Row chunks: what still fits under the table, then one chunk per copy; without a known room the rest goes on one copy."""
    heights = [tr.h for tr in shape.table._tbl.tr_lst]
    footer = heights[-1] if has_footer(shape, binding.keep_last_row_if, binding.header_rows) else 0
    budget = rest - sum(heights[:binding.header_rows]) - footer
    estimates = row_heights(shape, theme, rows=rows, header_rows=binding.header_rows, columns=spec.columns or None)
    chunks: list[list] = [[]]
    used, limit = 0, first
    for index, (row, height) in enumerate(zip(rows, estimates)):
        if used + height > limit and (chunks[-1] or len(chunks) == 1):
            if budget <= 0:
                chunks.append(list(rows[index:]))
                return chunks
            chunks.append([])
            used, limit = 0, budget
        chunks[-1].append(row)
        used += height
    return chunks


def _table_rooms(slide, pending: dict, floor: int, shift: int, grow_first: bool) -> dict[str, tuple[int, int]]:
    """Per table: room left under it on this slide (when the block grows) and the room a copy offers."""
    rooms: dict[str, tuple[int, int]] = {}
    for key, (spec, binding, _) in pending.items():
        shape = find_shape(slide, binding.shape.id)
        if spec.kind == "table" and shape is not None:
            rest = floor - (shape.top - shift) - BLOCK_GAP
            first = floor - (shape.top + shape.height) - BLOCK_GAP if grow_first else 0
            rooms[key] = (first, rest)
    return rooms


def _capacity(shape, binding: Binding, blocks: list[Block], theme) -> int:
    measured = capacity_chars_of(shape, theme, len(binding.keep_prefix or ""))
    return _budget(binding.max_chars, measured) or max(1, sum(len(b.text) + 1 for b in blocks))


def _budget(stated: int | None, measured: int | None) -> int | None:
    """The stated budget, unless it is far above what the box measures: then the analysis is stale."""
    if stated and measured:
        return min(stated, int(measured * STALE_BUDGET_RATIO))
    return stated or measured


def _growing_members(slide, block: Block, bottom: int, bound: dict[int, str]) -> list[int]:
    """Members that get taller with the block: those touching its bottom and every bound text box inside it."""
    ids = []
    for shape_id in block.members:
        shape = find_shape(slide, shape_id)
        if shape is None or not shape.height:
            continue
        if getattr(shape, "has_table", False):
            continue  # a table is as tall as its rows
        touches = abs(shape.top + shape.height - bottom) <= CONTAIN_TOL
        if touches or (shape_id in bound and getattr(shape, "has_text_frame", False)):
            ids.append(shape_id)
    return ids


def _shapes(slide, shape_ids) -> list:
    return [s for s in (find_shape(slide, i) for i in shape_ids) if s is not None]


def _grow(shape, delta: int) -> None:
    """A box that gets taller reads from the top, whatever the template anchored it to."""
    pin_geometry(shape)
    shape.height = shape.height + delta
    if getattr(shape, "has_text_frame", False):
        shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP


def _continue_composite(prs, slide, manifest: Manifest, number: int, pending: dict, issues: list[RenderIssue], kept: set[str] | None = None) -> None:
    """Copies a slide as often as its longest box needs; every copy continues each long box and blanks the rest."""
    count = max(len(chunks) for _, _, chunks in pending.values())
    current = slide
    for index in range(count):
        current = clone_slide(prs, current)
        _mark_continuation(current)
        for other in manifest.fields:
            if other.kind == "image" or other.key in (kept or set()):
                continue
            for binding in other.bindings:
                if binding.slide != number or binding.mode == "token":
                    continue
                target = find_shape(current, binding.shape.id)
                if target is None:
                    continue
                entry = pending.get(other.key)
                if entry is not None and other.kind == "table" and getattr(target, "has_table", False):
                    fill_table(target, entry[2][index] if index < len(entry[2]) else [], header_rows=binding.header_rows, keep_last_row_if=binding.keep_last_row_if, columns=other.columns or None)
                elif entry is not None and index < len(entry[2]):
                    set_rich_text(target, entry[2][index], keep_prefix=binding.keep_prefix)
                    _shrink(target, other, binding, issues)
                elif other.kind == "table" and getattr(target, "has_table", False):
                    fill_table(target, [], header_rows=binding.header_rows, keep_last_row_if=binding.keep_last_row_if, columns=other.columns or None)
                elif getattr(target, "has_text_frame", False):
                    set_rich_text(target, "", keep_prefix=binding.keep_prefix)


def _shrink(shape, spec: FieldSpec, binding: Binding, issues: list[RenderIssue], theme: tuple[str, str] | None = None) -> None:
    scale = fit_text_shape(shape, theme=theme)
    if scale < 1.0:
        issues.append(RenderIssue(level="info", field=spec.key, slide=binding.slide, message=f"text shrunk to {int(scale * 100)} percent to fit the box"))


def _split_blocks(blocks: list[Block], first: int, line_chars: int = 0, rest: int | None = None) -> list[list[Block]]:
    """Cuts paragraphs into chunks; every paragraph costs at least one line so short bullets count."""
    widest = max(first, rest or first)
    pieces: list[Block] = []
    for block in blocks:
        pieces.extend(_sentence_pieces(block, widest) if len(block.text) + 1 > widest else [block])
    chunks: list[list[Block]] = [[]]
    used = 0
    limit = first
    for block in pieces:
        size = max(len(block.text) + 1, line_chars)
        if chunks[-1] and used + size > limit:
            chunks.append([])
            used = 0
            limit = rest if rest is not None else first
        chunks[-1].append(block)
        used += size
    return chunks


def _sentence_pieces(block: Block, limit: int) -> list[Block]:
    """A paragraph longer than any box is cut at sentence ends; inline formatting is dropped in that case."""
    sentences = [s for s in SENTENCE_RE.split(block.text.strip()) if s]
    if len(sentences) < 2:
        return [block]
    groups: list[str] = []
    for sentence in sentences:
        if groups and len(groups[-1]) + 1 + len(sentence) <= limit:
            groups[-1] = groups[-1] + " " + sentence
        else:
            groups.append(sentence)
    return [Block(spans=[Span(text=group)], bullet=block.bullet, level=block.level) for group in groups]


def _mark_continuation(slide) -> None:
    title = slide.shapes.title
    if title is None or title.text.endswith(CONTINUATION_SUFFIX):
        return
    paragraph = title.text_frame.paragraphs[-1]
    if paragraph.runs:
        paragraph.runs[-1].text = paragraph.runs[-1].text.rstrip() + CONTINUATION_SUFFIX
    else:
        paragraph.text = paragraph.text.rstrip() + CONTINUATION_SUFFIX


def _as_text(value: Any) -> str:
    if isinstance(value, ImageValue):
        return value.path
    if isinstance(value, list):
        return "\n".join(" | ".join(str(v) for v in row.values()) for row in value)
    return str(value)


def _as_rows(value: Any) -> list:
    if isinstance(value, list):
        return value
    if isinstance(value, ImageValue):
        return []
    return parse_pipe_table(str(value))
