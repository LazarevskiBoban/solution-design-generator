from __future__ import annotations

from pathlib import Path
from typing import Iterable, Iterator, Literal

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.shapes.base import BaseShape
from pptx.shapes.connector import Connector
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture
from pptx.slide import Slide
from pydantic import BaseModel, Field

EMU_PER_INCH = 914400
PREVIEW_CHARS = 80
MAX_TABLE_ROWS_SHOWN = 20

ShapeKind = Literal["text", "picture", "table", "chart", "group", "connector", "other"]


class ParagraphInfo(BaseModel):
    text: str
    level: int = 0
    has_bullet: bool = False
    bold: bool | None = None
    font_size: float | None = None


class TableInfo(BaseModel):
    rows: int
    cols: int
    cells: list[list[str]]


class ShapeInfo(BaseModel):
    id: int
    name: str
    kind: ShapeKind
    placeholder_type: str | None = None
    placeholder_idx: int | None = None
    preset: str | None = None
    left: float | None = None
    top: float | None = None
    width: float | None = None
    height: float | None = None
    rotation: float = 0.0
    text: str = ""
    paragraphs: list[ParagraphInfo] = Field(default_factory=list)
    table: TableInfo | None = None
    children: list[ShapeInfo] = Field(default_factory=list)

    @property
    def preview(self) -> str:
        flat = " ".join(self.text.split())
        return flat if len(flat) <= PREVIEW_CHARS else flat[: PREVIEW_CHARS - 1] + "…"

    @property
    def has_geometry(self) -> bool:
        return None not in (self.left, self.top, self.width, self.height)

    def walk(self) -> Iterator[ShapeInfo]:
        yield self
        for child in self.children:
            yield from child.walk()


class SlideInfo(BaseModel):
    index: int
    layout: str
    title: str | None = None
    shapes: list[ShapeInfo] = Field(default_factory=list)
    notes: str = ""

    def walk(self) -> Iterator[ShapeInfo]:
        for shape in self.shapes:
            yield from shape.walk()

    @property
    def shape_count(self) -> int:
        return sum(1 for _ in self.walk())

    @property
    def connector_count(self) -> int:
        return sum(1 for s in self.walk() if s.kind == "connector")


class DeckInfo(BaseModel):
    path: str
    width: float
    height: float
    slides: list[SlideInfo] = Field(default_factory=list)


def inspect_deck(path: str | Path) -> DeckInfo:
    prs = Presentation(str(path))
    slides = [_slide_info(i, slide) for i, slide in enumerate(prs.slides, start=1)]
    return DeckInfo(
        path=str(path),
        width=prs.slide_width / EMU_PER_INCH,
        height=prs.slide_height / EMU_PER_INCH,
        slides=slides,
    )


def walk_shapes(shapes: Iterable[BaseShape]) -> Iterator[BaseShape]:
    for shape in shapes:
        yield shape
        if isinstance(shape, GroupShape):
            yield from walk_shapes(shape.shapes)


def find_shape(slide: Slide, shape_id: int) -> BaseShape | None:
    for shape in walk_shapes(slide.shapes):
        if shape.shape_id == shape_id:
            return shape
    return None


def format_inventory(deck: DeckInfo) -> str:
    lines = [f"Deck: {deck.path} ({deck.width:.2f} x {deck.height:.2f} in), {len(deck.slides)} slides"]
    for slide in deck.slides:
        lines.append("")
        lines.append(f"== Slide {slide.index}  layout={slide.layout!r}  title={slide.title!r}")
        for shape in slide.shapes:
            _format_shape(shape, 1, lines)
        if slide.notes:
            lines.append(f"   notes: {' '.join(slide.notes.split())[:200]}")
    return "\n".join(lines)


def _slide_info(index: int, slide: Slide) -> SlideInfo:
    title_shape = slide.shapes.title
    notes = ""
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
        notes = slide.notes_slide.notes_text_frame.text
    return SlideInfo(
        index=index,
        layout=slide.slide_layout.name,
        title=title_shape.text if title_shape is not None else None,
        shapes=[_shape_info(s) for s in slide.shapes],
        notes=notes,
    )


def _shape_info(shape: BaseShape) -> ShapeInfo:
    kind = _kind(shape)
    info = ShapeInfo(
        id=shape.shape_id,
        name=shape.name,
        kind=kind,
        preset=_preset(shape),
        left=_inches(shape.left),
        top=_inches(shape.top),
        width=_inches(shape.width),
        height=_inches(shape.height),
        rotation=_rotation(shape),
    )
    if shape.is_placeholder:
        fmt = shape.placeholder_format
        info.placeholder_type = fmt.type.name.lower() if fmt.type is not None else "object"
        info.placeholder_idx = fmt.idx
    if kind == "group":
        info.children = [_shape_info(child) for child in shape.shapes]
    elif kind == "table":
        info.table = _table_info(shape)
    elif shape.has_text_frame:
        info.paragraphs = [_paragraph_info(p) for p in shape.text_frame.paragraphs]
        info.text = "\n".join(p.text for p in info.paragraphs)
    return info


def _kind(shape: BaseShape) -> ShapeKind:
    if isinstance(shape, GroupShape):
        return "group"
    if isinstance(shape, Connector):
        return "connector"
    if isinstance(shape, Picture):
        return "picture"
    if isinstance(shape, GraphicFrame):
        if shape.has_table:
            return "table"
        if shape.has_chart:
            return "chart"
        return "other"
    if shape.is_placeholder and shape.placeholder_format.type is not None:
        if shape.placeholder_format.type.name == "PICTURE":
            return "picture"
    if shape.has_text_frame:
        return "text"
    return "other"


def _paragraph_info(paragraph) -> ParagraphInfo:
    ppr = paragraph._p.pPr
    has_bullet = ppr is not None and (
        ppr.find(qn("a:buChar")) is not None or ppr.find(qn("a:buAutoNum")) is not None
    )
    bold = None
    size = None
    for run in paragraph.runs:
        if run.text.strip():
            bold = run.font.bold
            size = run.font.size.pt if run.font.size is not None else None
            break
    return ParagraphInfo(
        text=paragraph.text, level=paragraph.level, has_bullet=has_bullet, bold=bold, font_size=size
    )


def _table_info(shape: GraphicFrame) -> TableInfo:
    table = shape.table
    cells = [[cell.text for cell in row.cells] for row in table.rows]
    return TableInfo(rows=len(table.rows), cols=len(table.columns), cells=cells)


def _preset(shape: BaseShape) -> str | None:
    sp_pr = shape._element.find(qn("p:spPr"))
    if sp_pr is None:
        return None
    geom = sp_pr.find(qn("a:prstGeom"))
    if geom is not None:
        return geom.get("prst")
    return "custom" if sp_pr.find(qn("a:custGeom")) is not None else None


def _inches(value) -> float | None:
    return None if value is None else round(value / EMU_PER_INCH, 3)


def _rotation(shape: BaseShape) -> float:
    try:
        return float(shape.rotation)
    except (AttributeError, NotImplementedError):
        return 0.0


def _format_shape(shape: ShapeInfo, depth: int, lines: list[str]) -> None:
    pad = "   " * depth
    head = f"{pad}[{shape.id}] {shape.kind} {shape.name!r}"
    if shape.table:
        head += f" {shape.table.rows}x{shape.table.cols}"
    if shape.placeholder_type:
        head += f" ph={shape.placeholder_type}/{shape.placeholder_idx}"
    if shape.has_geometry:
        head += f" @({shape.left:.2f},{shape.top:.2f}) {shape.width:.2f}x{shape.height:.2f}"
    if shape.text:
        head += f" :: {shape.preview}"
    lines.append(head)
    if shape.table:
        for row in shape.table.cells[:MAX_TABLE_ROWS_SHOWN]:
            lines.append(f"{pad}   | " + " | ".join(" ".join(c.split())[:40] for c in row))
    for child in shape.children:
        _format_shape(child, depth + 1, lines)
