"""The font and spacing a paragraph really gets: its runs, the shape's list style, the placeholders it inherits from, the master text styles and the presentation default."""

from __future__ import annotations

from functools import lru_cache

from lxml import etree
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from sdgen.textmetrics import FontSpec

A_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
BASE_TYPES = {
    PP_PLACEHOLDER.CENTER_TITLE: PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.VERTICAL_TITLE: PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.DATE: PP_PLACEHOLDER.DATE,
    PP_PLACEHOLDER.FOOTER: PP_PLACEHOLDER.FOOTER,
    PP_PLACEHOLDER.SLIDE_NUMBER: PP_PLACEHOLDER.SLIDE_NUMBER,
    PP_PLACEHOLDER.TITLE: PP_PLACEHOLDER.TITLE,
}
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
OTHER_TYPES = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}


@lru_cache(maxsize=128)
def theme_fonts(part) -> tuple[str, str]:
    """(major, minor) latin typefaces behind a slide part; Arial when the theme cannot be read."""
    try:
        theme = part.slide_layout.slide_master.part.part_related_by(RT.THEME)
        root = etree.fromstring(theme.blob)
    except Exception:
        return ("Arial", "Arial")
    major = root.find(".//a:majorFont/a:latin", A_NS)
    minor = root.find(".//a:minorFont/a:latin", A_NS)
    return ((major.get("typeface") if major is not None else None) or "Arial", (minor.get("typeface") if minor is not None else None) or "Arial")


def style_chain(shape, level: int = 1) -> list[etree._Element]:
    """The level paragraph properties that apply to the shape's text, nearest first."""
    tag = qn(f"a:lvl{level}pPr")
    found: list[etree._Element] = []

    def add(container) -> None:
        node = container.find(tag) if container is not None else None
        if node is not None:
            found.append(node)

    if getattr(shape, "has_text_frame", False):
        add(shape.text_frame._txBody.find(qn("a:lstStyle")))
    if getattr(shape, "is_placeholder", False):
        for base in _base_placeholders(shape):
            if getattr(base, "has_text_frame", False):
                add(base.text_frame._txBody.find(qn("a:lstStyle")))
        master = _master_of(shape.part)
        styles = master._element.find(qn("p:txStyles")) if master is not None else None
        if styles is not None:
            add(styles.find(qn(_style_tag(shape.placeholder_format.type))))
    try:
        add(shape.part.package.presentation_part.presentation._element.find(qn("p:defaultTextStyle")))
    except Exception:
        pass
    return found


def resolve_font(shape, p: etree._Element, default_pt: float, theme: tuple[str, str]) -> FontSpec:
    """Size, weight and typeface of a paragraph after inheritance; `default_pt` only when nothing defines a size."""
    ppr = p.find(qn("a:pPr"))
    candidates = [_run_rpr(p), ppr.find(qn("a:defRPr")) if ppr is not None else None]
    candidates += [node.find(qn("a:defRPr")) for node in style_chain(shape, _level(ppr))]
    present = [c for c in candidates if c is not None]
    size = next((int(c.get("sz")) / 100 for c in present if c.get("sz")), default_pt)
    bold = next((c.get("b") == "1" for c in present if c.get("b") is not None), False)
    latin = next((c.find(qn("a:latin")).get("typeface") for c in present if c.find(qn("a:latin")) is not None and c.find(qn("a:latin")).get("typeface")), "")
    if latin.startswith("+mj"):
        family = theme[0]
    elif not latin or latin.startswith("+mn"):
        family = theme[1]
    else:
        family = latin
    return FontSpec(family=family, size_pt=size, bold=bold)


def resolve_spacing(shape, p: etree._Element, size_pt: float) -> tuple[float, float, float]:
    """(line spacing in percent, space before in points, space after in points) after inheritance."""
    ppr = p.find(qn("a:pPr"))
    sources = ([ppr] if ppr is not None else []) + style_chain(shape, _level(ppr))
    pct = next((v for v in (_line_pct(s, size_pt) for s in sources) if v is not None), 100.0)
    before = next((v for v in (_space(s, "a:spcBef", size_pt) for s in sources) if v is not None), 0.0)
    after = next((v for v in (_space(s, "a:spcAft", size_pt) for s in sources) if v is not None), 0.0)
    return pct, before, after


def _level(ppr) -> int:
    return int(ppr.get("lvl", "0")) + 1 if ppr is not None else 1


def _run_rpr(p: etree._Element):
    for run in p.findall(qn("a:r")):
        text = "".join(t.text or "" for t in run.findall(qn("a:t")))
        if text.strip():
            return run.find(qn("a:rPr"))
    first = p.find(qn("a:r"))
    return first.find(qn("a:rPr")) if first is not None else p.find(qn("a:endParaRPr"))


def _line_pct(ppr, size_pt: float) -> float | None:
    node = ppr.find(qn("a:lnSpc"))
    if node is None:
        return None
    pct = node.find(qn("a:spcPct"))
    if pct is not None and pct.get("val"):
        return int(pct.get("val")) / 1000
    pts = node.find(qn("a:spcPts"))
    if pts is not None and pts.get("val") and size_pt:
        return int(pts.get("val")) / 100 / (size_pt * 1.2) * 100
    return None


def _space(ppr, tag: str, size_pt: float) -> float | None:
    node = ppr.find(qn(tag))
    if node is None:
        return None
    pts = node.find(qn("a:spcPts"))
    if pts is not None and pts.get("val"):
        return int(pts.get("val")) / 100
    pct = node.find(qn("a:spcPct"))
    if pct is not None and pct.get("val"):
        return int(pct.get("val")) / 100000 * size_pt
    return None


def _base_placeholders(shape) -> list:
    fmt = shape.placeholder_format
    part = shape.part
    bases: list = []
    layout = getattr(part, "slide_layout", None)
    if layout is not None:
        base = layout.placeholders.get(fmt.idx)
        if base is not None:
            bases.append(base)
        master = layout.slide_master
        ph_type = base.placeholder_format.type if base is not None else fmt.type
    else:
        master = getattr(part, "slide_master", None)
        ph_type = fmt.type
    if master is not None and ph_type is not None:
        base = master.placeholders.get(BASE_TYPES.get(ph_type, PP_PLACEHOLDER.BODY))
        if base is not None and base._element is not shape._element:
            bases.append(base)
    return bases


def _master_of(part):
    layout = getattr(part, "slide_layout", None)
    if layout is not None:
        return layout.slide_master
    return getattr(part, "slide_master", None)


def _style_tag(ph_type) -> str:
    if ph_type in TITLE_TYPES:
        return "p:titleStyle"
    if ph_type is None or ph_type in OTHER_TYPES:
        return "p:otherStyle"
    return "p:bodyStyle"
