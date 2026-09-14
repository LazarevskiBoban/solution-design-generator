"""A look at the produced deck without PowerPoint: what sits outside the slide, overlaps, stays empty or would not fit."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Literal

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Inches
from pydantic import BaseModel

from sdgen.blueprint import TITLE_MAX
from sdgen.fill.text import effective_overflow_ratio, theme_fonts
from sdgen.flow import FlowSpec, lane_mismatches
from sdgen.layout import Box

OUTSIDE_TOLERANCE = Inches(0.02)
OVERLAP_SHARE = 0.3  # of the smaller shape's area
OVERFLOW_TOLERANCE = 1.1
CONTINUATION = " (cont.)"
FOOTER_TYPES = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER}
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
VERSION_RE = re.compile(r"version|author|contributor", re.IGNORECASE)
FLOW_PREFIX = "Flow "


class Finding(BaseModel):
    level: Literal["info", "warning", "error"] = "warning"
    slide: int
    shape: str = ""
    code: str
    message: str


def check_deck(path: str | Path, flows: dict[int, FlowSpec] | None = None, title_max: int = TITLE_MAX) -> list[Finding]:
    return check_presentation(Presentation(str(path)), flows=flows, title_max=title_max)


Baseline = dict[int, set[tuple[str, str]]]


def check_presentation(prs, flows: dict[int, FlowSpec] | None = None, title_max: int = TITLE_MAX, baseline: Baseline | None = None, origins: list[int] | None = None) -> list[Finding]:
    """Findings per output slide; `flows` maps a slide position to the flow drawn on it.

    `baseline` holds what the template showed per template slide and `origins` the template slide each output
    slide came from: a finding the template already had is left out.
    """
    slides = list(prs.slides)
    theme = theme_fonts(slides[0].part) if slides else None
    findings: list[Finding] = []
    for number, slide in enumerate(slides, 1):
        origin = origins[number - 1] if origins and number <= len(origins) else number
        known = (baseline or {}).get(origin, set())
        findings.extend(f for f in _check_slide(slide, number, prs.slide_width, prs.slide_height, theme, title_max) if (f.code, f.shape) not in known)
        for problem in lane_mismatches((flows or {}).get(number)) if flows and number in flows else []:
            findings.append(Finding(slide=number, code="lane", message=problem))
    return findings


def baseline_findings(prs, title_max: int = TITLE_MAX) -> Baseline:
    """What each slide of the untouched template shows, keyed by slide number, code and shape name."""
    slides = list(prs.slides)
    theme = theme_fonts(slides[0].part) if slides else None
    return {number: {(f.code, f.shape) for f in _check_slide(slide, number, prs.slide_width, prs.slide_height, theme, title_max)} for number, slide in enumerate(slides, 1)}


def _check_slide(slide, number: int, width: int, height: int, theme, title_max: int) -> list[Finding]:
    findings: list[Finding] = []
    shapes = [s for s in slide.shapes if None not in (s.left, s.top, s.width, s.height) and s.width > 0 and s.height > 0]
    boxes = {s.shape_id: _bounds(s) for s in shapes}
    title = slide.shapes.title
    title_box = boxes.get(title.shape_id) if title is not None else None

    for shape in shapes:
        box = boxes[shape.shape_id]
        if box.left < -OUTSIDE_TOLERANCE or box.top < -OUTSIDE_TOLERANCE or box.right > width + OUTSIDE_TOLERANCE or box.bottom > height + OUTSIDE_TOLERANCE:
            findings.append(Finding(slide=number, shape=shape.name, code="outside", message=f"'{shape.name}' reaches outside the slide"))

    flow_shapes = [s for s in shapes if s.name.startswith(FLOW_PREFIX)]
    nodes = [s for s in flow_shapes if " node " in s.name]
    labels = [s for s in flow_shapes if s.name.endswith(" label")]
    for index, first in enumerate(nodes):
        for second in nodes[index + 1 :]:
            if _overlap_share(boxes[first.shape_id], boxes[second.shape_id]) > OVERLAP_SHARE:
                findings.append(Finding(slide=number, shape=first.name, code="overlap", message=f"'{first.name}' overlaps '{second.name}'"))
    for label in labels:
        for node in nodes:
            if _overlap_share(boxes[label.shape_id], boxes[node.shape_id]) > OVERLAP_SHARE:
                findings.append(Finding(slide=number, shape=label.name, code="overlap", message=f"label '{label.text_frame.text}' sits on '{node.name}'"))
    if title_box is not None:
        for canvas in (s for s in flow_shapes if s.name.endswith(" canvas")):
            if _overlap_share(boxes[canvas.shape_id], title_box) > OVERLAP_SHARE:
                findings.append(Finding(slide=number, shape=canvas.name, code="overlap", message="the drawing overlaps the title"))

    content = False
    for shape in shapes:
        kind = _placeholder_type(shape)
        if kind in FOOTER_TYPES:
            continue
        is_title = title is not None and shape.shape_id == title.shape_id or kind in TITLE_TYPES
        has_table = getattr(shape, "has_table", False)
        has_text = getattr(shape, "has_text_frame", False) and bool(shape.text_frame.text.strip())
        if has_table:
            rows = list(shape.table.rows)
            header = " ".join(c.text for c in rows[0].cells) if rows else ""
            data = [r for r in rows[1:] if any(c.text.strip() for c in r.cells)]
            content = content or bool(data)
            if len(data) <= 1 and not VERSION_RE.search(header):
                findings.append(Finding(level="info", slide=number, shape=shape.name, code="thin_table", message=f"table '{shape.name}' has {len(data)} data row(s)"))
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or _has_image(shape) or shape.name.startswith(FLOW_PREFIX):
            content = True
            continue
        if shape.is_placeholder and not is_title and not has_text and getattr(shape, "has_text_frame", False):
            findings.append(Finding(slide=number, shape=shape.name, code="empty_box", message=f"placeholder '{shape.name}' is empty"))
        if has_text and not is_title:
            content = True
        if has_text and not (shape.is_placeholder and kind in FOOTER_TYPES):
            ratio = effective_overflow_ratio(shape, theme)
            if ratio > OVERFLOW_TOLERANCE:
                findings.append(Finding(slide=number, shape=shape.name, code="overflow", message=f"'{shape.name}' needs about {int(ratio * 100)} percent of its box height"))
    if not content:
        findings.append(Finding(slide=number, code="empty_slide", message="nothing but the title on the slide"))

    if title is not None and getattr(title, "has_text_frame", False):
        text = title.text_frame.text.strip().removesuffix(CONTINUATION)
        if len(text) > title_max:
            findings.append(Finding(slide=number, shape=title.name, code="title_long", message=f"title has {len(text)} characters, more than {title_max}"))
    return findings


def _bounds(shape) -> Box:
    """The box a shape covers on the slide, rotated shapes swapped around their centre."""
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    if round(getattr(shape, "rotation", 0.0) or 0.0) % 180 == 90:
        cx, cy = left + width // 2, top + height // 2
        return Box(cx - height // 2, cy - width // 2, cx + height // 2, cy + width // 2)
    return Box(left, top, left + width, top + height)


def _overlap_share(a: Box, b: Box) -> float:
    overlap = a.overlap_x(b) * max(0, min(a.bottom, b.bottom) - max(a.top, b.top))
    smaller = min(a.width * a.height, b.width * b.height)
    return overlap / smaller if smaller else 0.0


def _placeholder_type(shape):
    try:
        return shape.placeholder_format.type if shape.is_placeholder else None
    except (AttributeError, ValueError):
        return None


def _has_image(shape) -> bool:
    try:
        return shape.image is not None
    except (AttributeError, ValueError):
        return False
