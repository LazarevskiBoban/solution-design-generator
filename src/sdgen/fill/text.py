from __future__ import annotations

import copy
import re

from lxml import etree
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pydantic import BaseModel, Field

from sdgen.inventory import walk_shapes
from sdgen.textmetrics import FontSpec, capacity_chars, line_height_pt, wrapped_lines

BULLET_RE = re.compile(r"^(\s*)[-*•]\s+(.*)$")
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)")
DEFAULT_INDENT_EMU = 342900
MAX_LEVEL = 8
RUN_TAGS = (qn("a:r"), qn("a:br"), qn("a:fld"))
BULLET_TAGS = (qn("a:buChar"), qn("a:buAutoNum"), qn("a:buNone"), qn("a:buBlip"))
PPR_TAIL_TAGS = (qn("a:tabLst"), qn("a:defRPr"), qn("a:extLst"))


class Span(BaseModel):
    text: str
    bold: bool = False
    italic: bool = False


class Block(BaseModel):
    spans: list[Span] = Field(default_factory=list)
    bullet: bool = False
    level: int = 0

    @property
    def text(self) -> str:
        return "".join(s.text for s in self.spans)


def parse_blocks(value: str) -> list[Block]:
    blocks: list[Block] = []
    for raw in value.splitlines():
        if not raw.strip():
            continue
        match = BULLET_RE.match(raw)
        if match:
            indent = match.group(1).replace("\t", "  ")
            level = min(len(indent) // 2, MAX_LEVEL)
            blocks.append(Block(spans=parse_spans(match.group(2).strip()), bullet=True, level=level))
        else:
            blocks.append(Block(spans=parse_spans(raw.strip())))
    return blocks


def parse_spans(text: str) -> list[Span]:
    spans: list[Span] = []
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            spans.append(Span(text=part[2:-2], bold=True))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            spans.append(Span(text=part[1:-1], italic=True))
        else:
            spans.append(Span(text=part))
    return spans


def set_rich_text(target, value: str | list[Block], keep_prefix: str | None = None) -> None:
    blocks = parse_blocks(value) if isinstance(value, str) else [b.model_copy(deep=True) for b in value]
    tx_body = target.text_frame._txBody
    paragraphs = tx_body.findall(qn("a:p"))
    inherits_bullets = bool(getattr(target, "is_placeholder", False))

    base = _first(paragraphs, _has_text)
    if base is None:
        base = paragraphs[0]
    plain = _first(paragraphs, lambda p: _has_text(p) and not _has_bullet(p))
    if plain is None:
        plain = base
    bullet_templates = _bullet_templates(paragraphs)

    prefix_p = None
    prefix_rpr = None
    if keep_prefix:
        prefix_p, prefix_rpr = _split_prefix(paragraphs[0], keep_prefix)
        if prefix_p is None and blocks:
            blocks[0].spans.insert(0, Span(text=keep_prefix))

    new_paragraphs: list[etree._Element] = []
    for index, block in enumerate(blocks):
        if index == 0 and prefix_p is not None:
            _append_spans(prefix_p, block.spans, prefix_rpr)
            new_paragraphs.append(prefix_p)
            continue
        template = _pick_template(block, plain, bullet_templates)
        paragraph = _clone_paragraph(template)
        _apply_bullet(paragraph, block, _has_bullet(template), inherits_bullets)
        rpr = _run_rpr(template)
        _append_spans(paragraph, block.spans, rpr)
        _keep_style_when_empty(paragraph, rpr)
        new_paragraphs.append(paragraph)

    if not new_paragraphs:
        if prefix_p is None:
            prefix_p = _clone_paragraph(base)
            _keep_style_when_empty(prefix_p, _run_rpr(base))
        new_paragraphs.append(prefix_p)

    for paragraph in paragraphs:
        tx_body.remove(paragraph)
    for paragraph in new_paragraphs:
        tx_body.append(paragraph)


def replace_token(target, token: str, value: str, loose_spaces: bool = False) -> int:
    pattern = _loose_pattern(token) if loose_spaces else None
    return sum(_replace_in_paragraph(p._p, token, value, pattern) for p in target.text_frame.paragraphs)


def replace_literal_everywhere(prs, old: str, new: str, loose_spaces: bool = True) -> int:
    count = 0
    for slide in prs.slides:
        for shape in walk_shapes(slide.shapes):
            if shape.has_text_frame:
                count += replace_token(shape, old, new, loose_spaces)
            elif getattr(shape, "has_table", False):
                for cell in shape.table.iter_cells():
                    count += replace_token(cell, old, new, loose_spaces)
    return count


def _loose_pattern(token: str) -> re.Pattern | None:
    words = token.split()
    return re.compile(r"\s+".join(re.escape(w) for w in words)) if words else None


def _replace_in_paragraph(p: etree._Element, token: str, value: str, pattern: re.Pattern | None = None) -> int:
    count = 0
    search_from = 0
    while True:
        runs = [r for r in p if r.tag == qn("a:r")]
        texts = [_text_of(r) for r in runs]
        full = "".join(texts)
        if pattern is not None:
            match = pattern.search(full, search_from)
            start, end = (match.start(), match.end()) if match else (-1, -1)
        else:
            start = full.find(token, search_from)
            end = start + len(token)
        if start < 0:
            return count
        first = last = None
        local_start = local_end = 0
        offset = 0
        for i, text in enumerate(texts):
            if first is None and start < offset + len(text):
                first, local_start = i, start - offset
            if first is not None and end <= offset + len(text):
                last, local_end = i, end - offset
                break
            offset += len(text)
        if first is None or last is None:
            return count
        head = texts[first][:local_start]
        tail = texts[last][local_end:]
        _t(runs[first]).text = head + value + tail
        for run in runs[first + 1 : last + 1]:
            p.remove(run)
        count += 1
        search_from = start + len(value)


def _split_prefix(p: etree._Element, prefix: str):
    runs = [r for r in p if r.tag == qn("a:r")]
    if not "".join(_text_of(r) for r in runs).startswith(prefix):
        return None, None
    clone = _clone_paragraph(p)
    position = 0
    content_rpr = None
    for i, run in enumerate(runs):
        text = _text_of(run)
        kept = copy.deepcopy(run)
        if position + len(text) <= len(prefix):
            position += len(text)
            _insert_run(clone, kept)
            if position == len(prefix):
                source = runs[i + 1] if i + 1 < len(runs) else run
                content_rpr = _rpr_of(source)
                break
        else:
            _t(kept).text = text[: len(prefix) - position]
            _insert_run(clone, kept)
            content_rpr = _rpr_of(run)
            break
    return clone, content_rpr


def _pick_template(block: Block, plain: etree._Element, bullets: dict[int, etree._Element]) -> etree._Element:
    if not block.bullet or not bullets:
        return plain
    if block.level in bullets:
        return bullets[block.level]
    lower = [lvl for lvl in bullets if lvl <= block.level]
    return bullets[max(lower)] if lower else bullets[min(bullets)]


def _clone_paragraph(p: etree._Element) -> etree._Element:
    clone = copy.deepcopy(p)
    for child in list(clone):
        if child.tag in RUN_TAGS:
            clone.remove(child)
    return clone


def _apply_bullet(p: etree._Element, block: Block, template_has_bullet: bool, inherits: bool) -> None:
    ppr = p.find(qn("a:pPr"))
    if block.bullet:
        if ppr is None:
            ppr = etree.Element(qn("a:pPr"))
            p.insert(0, ppr)
        if block.level:
            ppr.set("lvl", str(block.level))
        else:
            ppr.attrib.pop("lvl", None)
        if not template_has_bullet and not inherits:
            _remove_bullet_elements(ppr)
            ppr.set("marL", str(DEFAULT_INDENT_EMU * (block.level + 1)))
            ppr.set("indent", str(-DEFAULT_INDENT_EMU))
            bullet = etree.Element(qn("a:buChar"))
            bullet.set("char", "•")
            _insert_in_ppr(ppr, bullet)
    elif template_has_bullet and ppr is not None:
        _remove_bullet_elements(ppr)
        ppr.attrib.pop("lvl", None)
        ppr.set("marL", "0")
        ppr.set("indent", "0")
        _insert_in_ppr(ppr, etree.Element(qn("a:buNone")))


def _append_spans(p: etree._Element, spans: list[Span], rpr: etree._Element | None) -> None:
    for span in spans:
        run = etree.Element(qn("a:r"))
        run_rpr = copy.deepcopy(rpr) if rpr is not None else None
        if run_rpr is None and (span.bold or span.italic):
            run_rpr = etree.Element(qn("a:rPr"))
        if run_rpr is not None:
            if span.bold:
                run_rpr.set("b", "1")
            if span.italic:
                run_rpr.set("i", "1")
            run.append(run_rpr)
        text = etree.SubElement(run, qn("a:t"))
        text.text = span.text
        _insert_run(p, run)


def _keep_style_when_empty(p: etree._Element, rpr: etree._Element | None) -> None:
    if rpr is None or p.find(qn("a:r")) is not None:
        return
    end = copy.deepcopy(rpr)
    end.tag = qn("a:endParaRPr")
    existing = p.find(qn("a:endParaRPr"))
    if existing is not None:
        p.replace(existing, end)
    else:
        p.append(end)


def _insert_run(p: etree._Element, run: etree._Element) -> None:
    end = p.find(qn("a:endParaRPr"))
    if end is not None:
        end.addprevious(run)
    else:
        p.append(run)


def _insert_in_ppr(ppr: etree._Element, element: etree._Element) -> None:
    for child in ppr:
        if child.tag in PPR_TAIL_TAGS:
            child.addprevious(element)
            return
    ppr.append(element)


def _remove_bullet_elements(ppr: etree._Element) -> None:
    for child in list(ppr):
        if child.tag in BULLET_TAGS:
            ppr.remove(child)


def _bullet_templates(paragraphs: list[etree._Element]) -> dict[int, etree._Element]:
    templates: dict[int, etree._Element] = {}
    for p in paragraphs:
        if _has_bullet(p):
            level = int(p.find(qn("a:pPr")).get("lvl", "0"))
            templates.setdefault(level, p)
    return templates


def _has_bullet(p: etree._Element) -> bool:
    ppr = p.find(qn("a:pPr"))
    return ppr is not None and (
        ppr.find(qn("a:buChar")) is not None or ppr.find(qn("a:buAutoNum")) is not None
    )


def _has_text(p: etree._Element) -> bool:
    return any((t.text or "").strip() for t in p.iter(qn("a:t")))


def _run_rpr(p: etree._Element) -> etree._Element | None:
    for run in p.findall(qn("a:r")):
        if _text_of(run).strip():
            return _rpr_of(run)
    end = p.find(qn("a:endParaRPr"))
    if end is not None:
        rpr = copy.deepcopy(end)
        rpr.tag = qn("a:rPr")
        return rpr
    return None


def _rpr_of(run: etree._Element) -> etree._Element | None:
    rpr = run.find(qn("a:rPr"))
    return copy.deepcopy(rpr) if rpr is not None else None


def _text_of(run: etree._Element) -> str:
    return run.findtext(qn("a:t")) or ""


def _t(run: etree._Element) -> etree._Element:
    t = run.find(qn("a:t"))
    if t is None:
        t = etree.SubElement(run, qn("a:t"))
    return t


def _first(items, predicate):
    return next((item for item in items if predicate(item)), None)


EMU_PER_PT = 12700
MIN_FONT_SCALE = 0.8
DEFAULT_FONT_PT = 14.0
DEFAULT_INSETS_EMU = (91440, 45720, 91440, 45720)  # left, top, right, bottom
A_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


class Measurement(BaseModel):
    needed_pt: float
    usable_width_pt: float
    usable_height_pt: float
    lines: int

    @property
    def ratio(self) -> float:
        return self.needed_pt / self.usable_height_pt if self.usable_height_pt > 0 else 0.0


def theme_fonts(part) -> tuple[str, str]:
    """(major, minor) latin typefaces behind a slide part; Arial when the theme cannot be read."""
    try:
        theme = part.slide_layout.slide_master.part.part_related_by(RT.THEME)
        root = etree.fromstring(theme.blob)
    except Exception:
        return ("Arial", "Arial")
    major = root.find(".//a:majorFont/a:latin", A_NS)
    minor = root.find(".//a:minorFont/a:latin", A_NS)
    return ((major.get("typeface") if major is not None else None) or "Arial", (minor.get("typeface") if minor is not None else None) or "Arial")


def measure_shape(shape, default_pt: float = DEFAULT_FONT_PT, theme: tuple[str, str] | None = None) -> Measurement | None:
    """Height the text needs against the height the box offers, with the deck fonts and paragraph spacing."""
    if not getattr(shape, "has_text_frame", False) or shape.width is None or shape.height is None:
        return None
    frame = shape.text_frame
    usable_w, usable_h = _usable_pt(shape, frame)
    if usable_w <= 0 or usable_h <= 0:
        return None
    theme = theme or theme_fonts(shape.part)
    tx_body = frame._txBody
    needed, lines = 0.0, 0
    for paragraph in frame.paragraphs:
        p = paragraph._p
        ppr = p.find(qn("a:pPr"))
        spec = _font_spec(p, default_pt, theme, tx_body)
        if ppr is not None and ppr.get("marL"):
            indent = int(ppr.get("marL")) / EMU_PER_PT
        else:
            indent = DEFAULT_INDENT_EMU * (paragraph.level + 1) / EMU_PER_PT if _has_bullet(p) else 0.0
        text = " ".join(paragraph.text.split())
        count = wrapped_lines(text, usable_w - indent, spec) if text else 1
        lines += count
        needed += count * line_height_pt(spec, _spacing_pct(ppr)) + _space_pt(ppr, "a:spcBef", spec.size_pt) + _space_pt(ppr, "a:spcAft", spec.size_pt)
    return Measurement(needed_pt=needed, usable_width_pt=usable_w, usable_height_pt=usable_h, lines=lines)


def overflow_ratio(shape, theme: tuple[str, str] | None = None) -> float:
    measure = measure_shape(shape, theme=theme)
    return measure.ratio if measure is not None else 0.0


def capacity_chars_of(shape, theme: tuple[str, str] | None = None, prefix_len: int = 0, default_pt: float = DEFAULT_FONT_PT) -> int | None:
    """Characters that fit the box comfortably, judged by the font and spacing of its first paragraph."""
    if not getattr(shape, "has_text_frame", False) or shape.width is None or shape.height is None:
        return None
    frame = shape.text_frame
    usable_w, usable_h = _usable_pt(shape, frame)
    if usable_w <= 0 or usable_h <= 0:
        return 0
    theme = theme or theme_fonts(shape.part)
    p = frame.paragraphs[0]._p
    ppr = p.find(qn("a:pPr"))
    spec = _font_spec(p, default_pt, theme, frame._txBody)
    return capacity_chars(usable_w, usable_h, spec, _spacing_pct(ppr), _space_pt(ppr, "a:spcAft", spec.size_pt), prefix_len)


def fit_text_shape(shape, default_pt: float = DEFAULT_FONT_PT, theme: tuple[str, str] | None = None) -> float:
    """Measures whether the text overflows its box and stores a PowerPoint font scale so it shrinks to fit.

    PowerPoint only recomputes shrink-to-fit when a user edits the text, so the scale is written explicitly.
    Returns the scale applied (1.0 when the text fits).
    """
    measure = measure_shape(shape, default_pt, theme)
    if measure is None:
        return 1.0
    scale = 1.0 if measure.needed_pt <= measure.usable_height_pt else max(MIN_FONT_SCALE, measure.usable_height_pt / measure.needed_pt)
    body = shape.text_frame._txBody.bodyPr
    existing = body.find(qn("a:normAutofit"))
    if scale >= 1.0 and existing is None:
        return 1.0
    for tag in ("a:noAutofit", "a:spAutoFit", "a:normAutofit"):
        for element in body.findall(qn(tag)):
            body.remove(element)
    autofit = etree.SubElement(body, qn("a:normAutofit"))
    if scale < 1.0:
        autofit.set("fontScale", str(int(round(scale * 100000))))
        autofit.set("lnSpcReduction", "10000")
    _order_body_children(body, autofit)
    return scale


def _usable_pt(shape, frame) -> tuple[float, float]:
    left, top, right, bottom = (
        frame.margin_left if frame.margin_left is not None else DEFAULT_INSETS_EMU[0],
        frame.margin_top if frame.margin_top is not None else DEFAULT_INSETS_EMU[1],
        frame.margin_right if frame.margin_right is not None else DEFAULT_INSETS_EMU[2],
        frame.margin_bottom if frame.margin_bottom is not None else DEFAULT_INSETS_EMU[3],
    )
    return (shape.width - left - right) / EMU_PER_PT, (shape.height - top - bottom) / EMU_PER_PT


def _font_spec(p: etree._Element, default_pt: float, theme: tuple[str, str], tx_body=None) -> FontSpec:
    """Size, weight and typeface of a paragraph: its first run, else the paragraph or list defaults, else the theme."""
    ppr = p.find(qn("a:pPr"))
    candidates = [_run_rpr(p), ppr.find(qn("a:defRPr")) if ppr is not None else None]
    lst = tx_body.find(qn("a:lstStyle")) if tx_body is not None else None
    if lst is not None:
        level = int(ppr.get("lvl", "0")) + 1 if ppr is not None else 1
        lvl = lst.find(qn(f"a:lvl{level}pPr"))
        candidates.append(lvl.find(qn("a:defRPr")) if lvl is not None else None)
    present = [c for c in candidates if c is not None]
    size = next((int(c.get("sz")) / 100 for c in present if c.get("sz")), default_pt)
    bold = next((c.get("b") == "1" for c in present if c.get("b") is not None), False)
    latin = next((c.find(qn("a:latin")).get("typeface") for c in present if c.find(qn("a:latin")) is not None and c.find(qn("a:latin")).get("typeface")), "")
    if latin.startswith("+mj"):
        family = theme[0]
    elif not latin or latin.startswith("+mn"):
        family = theme[1]
    else:
        family = latin
    return FontSpec(family=family, size_pt=size, bold=bold)


def _spacing_pct(ppr) -> float:
    node = ppr.find(qn("a:lnSpc")) if ppr is not None else None
    pct = node.find(qn("a:spcPct")) if node is not None else None
    return int(pct.get("val")) / 1000 if pct is not None and pct.get("val") else 100.0


def _space_pt(ppr, tag: str, size_pt: float) -> float:
    node = ppr.find(qn(tag)) if ppr is not None else None
    if node is None:
        return 0.0
    pts = node.find(qn("a:spcPts"))
    if pts is not None and pts.get("val"):
        return int(pts.get("val")) / 100
    pct = node.find(qn("a:spcPct"))
    if pct is not None and pct.get("val"):
        return int(pct.get("val")) / 100000 * size_pt
    return 0.0


def _order_body_children(body, autofit) -> None:
    # The autofit choice must precede scene3d, sp3d, flatTx and extLst in the body properties.
    for tag in ("a:scene3d", "a:sp3d", "a:flatTx", "a:extLst"):
        follower = body.find(qn(tag))
        if follower is not None:
            follower.addprevious(autofit)
            return
