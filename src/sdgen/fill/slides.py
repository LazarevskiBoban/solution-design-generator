from __future__ import annotations

import copy

from pptx.oxml.ns import qn
from pptx.slide import Slide

RELATIONSHIP_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
NOTES_RELTYPE_SUFFIX = "/notesSlide"
SHAPE_TAGS = {
    qn("p:sp"),
    qn("p:grpSp"),
    qn("p:graphicFrame"),
    qn("p:cxnSp"),
    qn("p:pic"),
    qn("p:contentPart"),
}


def slide_index(prs, slide: Slide) -> int:
    return prs.slides.index(slide)


def remove_slide(prs, slide: Slide) -> None:
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        if prs.part.related_part(sld_id.rId) is slide.part:
            sld_id_lst.remove(sld_id)
            prs.part.drop_rel(sld_id.rId)
            return
    raise ValueError("slide not found in presentation")


def move_slide(prs, slide: Slide, new_index: int) -> None:
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        if prs.part.related_part(sld_id.rId) is slide.part:
            sld_id_lst.remove(sld_id)
            sld_id_lst.insert(new_index, sld_id)
            return
    raise ValueError("slide not found in presentation")


def clone_slide(prs, slide: Slide) -> Slide:
    new = prs.slides.add_slide(slide.slide_layout)
    sp_tree = new.shapes._spTree
    for shape in list(new.shapes):
        sp_tree.remove(shape._element)

    background = slide._element.cSld.find(qn("p:bg"))
    if background is not None:
        new._element.cSld.insert(0, copy.deepcopy(background))

    for element in slide.shapes._spTree:
        if element.tag in SHAPE_TAGS:
            sp_tree.append(copy.deepcopy(element))

    remap = _copy_relationships(slide, new)
    if remap:
        _rewrite_relationship_ids(sp_tree, remap)

    move_slide(prs, new, slide_index(prs, slide) + 1)
    return new


def _copy_relationships(source: Slide, target: Slide) -> dict[str, str]:
    remap: dict[str, str] = {}
    for rel in list(source.part.rels.values()):
        if rel.reltype.endswith(NOTES_RELTYPE_SUFFIX):
            continue
        if rel.is_external:
            new_rid = target.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            new_rid = target.part.rels.get_or_add(rel.reltype, rel.target_part)
        if new_rid != rel.rId:
            remap[rel.rId] = new_rid
    return remap


def _rewrite_relationship_ids(root, remap: dict[str, str]) -> None:
    for element in root.iter():
        for attr, value in element.attrib.items():
            if attr.startswith(RELATIONSHIP_NS) and value in remap:
                element.set(attr, remap[value])
