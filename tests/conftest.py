from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

PROVIDER_VARS = ("SDGEN_LLM", "OPENAI_API_KEY", "SDGEN_OPENAI_MODEL", "AZURE_OPENAI_API_KEY", "AZURE_OPENAI_ENDPOINT", "AZURE_OPENAI_DEPLOYMENT", "AZURE_OPENAI_API_VERSION")


@pytest.fixture(autouse=True)
def no_real_provider(monkeypatch):
    # Streamlit copies secrets.toml into os.environ, so a UI test could leak real credentials into later tests.
    for var in PROVIDER_VARS:
        monkeypatch.delenv(var, raising=False)


@pytest.fixture
def sample_deck(tmp_path: Path) -> Path:
    image = tmp_path / "pic.png"
    Image.new("RGB", (40, 30), "white").save(image)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Executive Overview: Demo Integration"
    body = slide.placeholders[1].text_frame
    body.text = "First point"
    sub = body.add_paragraph()
    sub.text = "Sub point"
    sub.level = 1

    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1.5))
    box.name = "Business Need Box"
    box.text_frame.text = "Business Need: Something long enough to be treated as a real content section."
    second = box.text_frame.add_paragraph()
    second.text = "Second paragraph."
    second.runs[0].font.bold = True
    second.runs[0].font.size = Pt(12)

    label = slide.shapes.add_textbox(Inches(6), Inches(1.6), Inches(2), Inches(0.3))
    label.name = "Scope Label"
    label.text_frame.text = "Scope"

    frame = slide.shapes.add_table(3, 2, Inches(6), Inches(2), Inches(4), Inches(1.5))
    frame.name = "Scope Table"
    table = frame.table
    table.cell(0, 0).text = "Function"
    table.cell(0, 1).text = "Countries"
    table.cell(1, 0).text = "Finance"
    table.cell(1, 1).text = "ZA"

    slide.shapes.add_picture(str(image), Inches(1), Inches(5), Inches(2))

    group = slide.shapes.add_group_shape()
    group.name = "Systems Group"
    left = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8), Inches(5), Inches(1), Inches(0.5))
    left.text_frame.text = "ISPIC"
    right = group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10), Inches(5), Inches(1), Inches(0.5))
    right.text_frame.text = "BTP"
    slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(9), Inches(5.25), Inches(10), Inches(5.25))

    slide.notes_slide.notes_text_frame.text = "guidance note"
    prs.slides.add_slide(prs.slide_layouts[6])

    path = tmp_path / "sample.pptx"
    prs.save(path)
    return path
