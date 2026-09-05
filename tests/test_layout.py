from pptx import Presentation
from pptx.util import Inches

from sdgen.analyze import analyze_deck
from sdgen.inventory import inspect_deck
from sdgen.layout import analyse_slide


def _box(slide, name, left, top, width, height, text=""):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.name = name
    shape.text_frame.text = text
    return shape


def overview_deck():
    """Title, a full-width bar and box, two lower blocks side by side and a footer, like an executive overview."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Executive Overview"
    slide.shapes.title.height = Inches(0.6)
    _box(slide, "Need Bar", 0.5, 1.2, 9, 0.3, "Business Need")
    need = _box(slide, "Need Box", 0.5, 1.55, 9, 2.0, "need text")
    _box(slide, "Left Bar", 0.5, 3.8, 4.3, 0.3, "Solution Overview")
    left = _box(slide, "Left Box", 0.5, 4.15, 4.3, 1.5, "left text")
    _box(slide, "Right Bar", 5.2, 3.8, 4.3, 0.3, "Delivery")
    _box(slide, "Right Container", 5.2, 4.15, 4.3, 1.5, "")
    inner = _box(slide, "Right Inner", 5.3, 4.5, 4.0, 1.0, "inner text")
    _box(slide, "Footer", 0, 7.0, 2, 0.4, "Sensitivity")
    return prs, slide, {"need": need, "left": left, "inner": inner}


def test_blocks_bands_and_headers():
    prs, slide, shapes = overview_deck()
    bound = {shapes["need"].shape_id: "need", shapes["left"].shape_id: "left", shapes["inner"].shape_id: "inner"}
    layout = analyse_slide(slide, bound, prs.slide_height)
    names = {s.shape_id: s.name for s in slide.shapes}
    assert names[layout.title].startswith("Title") and [names[i] for i in layout.bottom_band] == ["Footer"]
    assert len(layout.blocks) == 3
    top = layout.top_block()
    assert top.fields == {"need"} and names[top.header] == "Need Bar" and layout.is_full_width(top)
    left = layout.block_of(shapes["left"].shape_id)
    assert names[left.header] == "Left Bar" and not layout.is_full_width(left)
    right = layout.block_of(shapes["inner"].shape_id)
    assert sorted(names[i] for i in right.members) == ["Right Bar", "Right Container", "Right Inner"]
    assert layout.content.top == top.box.top and layout.content.bottom == Inches(5.65) and layout.floor == Inches(6.95)


def test_side_by_side_top_row_is_not_full_width():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    a = _box(slide, "A", 0.5, 1.2, 4.3, 2, "a")
    b = _box(slide, "B", 5.2, 1.2, 4.3, 2, "b")
    layout = analyse_slide(slide, {a.shape_id: "a", b.shape_id: "b"}, prs.slide_height)
    assert len(layout.blocks) == 2 and not layout.is_full_width(layout.top_block())


def test_sample_deck_label_is_the_table_header(sample_deck):
    manifest = analyze_deck(inspect_deck(sample_deck)).to_manifest("demo")
    prs = Presentation(str(sample_deck))
    slide = prs.slides[0]
    bound = {b.shape.id: f.key for f in manifest.fields for b in f.bindings if b.slide == 1 and b.mode == "replace" and f.kind != "image"}
    layout = analyse_slide(slide, bound, prs.slide_height)
    names = {s.shape_id: s.name for s in slide.shapes}
    table = next(s for s in slide.shapes if s.has_table)
    scope = layout.block_of(table.shape_id)
    assert names[scope.header] == "Scope Label" and scope.fields == {"scope"}
    assert not layout.is_full_width(layout.top_block())
