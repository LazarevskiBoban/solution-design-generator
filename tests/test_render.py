from pathlib import Path

import pytest
from click.testing import CliRunner
from PIL import Image
from pptx import Presentation

from sdgen.analyze import analyze_deck
from sdgen.cli import main
from sdgen.content import Content, ImageValue, load_markdown, skeleton_markdown
from sdgen.inventory import inspect_deck
from sdgen.manifest import Binding, FieldSpec, GlobalSpec, Manifest, ShapeRef, SlideRules
from sdgen.registry import Registry
from sdgen.render import render

CARRIER_DECK = Path("D:/NTT-architectures/NTT DATA Inc I IT I Solution Design I Carrier AP Invoice Integration.pptx")


def _fixture_manifest(sample_deck) -> Manifest:
    analysis = analyze_deck(inspect_deck(sample_deck))
    for cand in analysis.candidates:
        if cand.kind == "image":
            cand.include = True
    manifest = analysis.to_manifest("demo", source="sample.pptx")
    manifest.globals.append(GlobalSpec(key="subject", replaces="Demo Integration"))
    manifest.slides.exclude = [2]
    return manifest


def _shape_text(slide, name):
    return next(s for s in slide.shapes if s.name == name).text_frame.text


def test_render_fills_all_kinds_and_excludes_slides(sample_deck, tmp_path):
    manifest = _fixture_manifest(sample_deck)
    image = tmp_path / "new.png"
    Image.new("RGB", (60, 30), "green").save(image)
    image_key = next(f.key for f in manifest.fields if f.kind == "image")
    content = Content(
        globals={"subject": "Carrier Invoices"},
        fields={
            "business_need": "Fresh need text\n- with a bullet",
            "scope": [{"Function": "Finance", "Countries": "ZA"}, {"Function": "Ops", "Countries": "KE"}],
            "first_point": "- alpha\n  - beta",
            image_key: ImageValue(path=str(image)),
        },
    )
    out = tmp_path / "out.pptx"
    result = render(sample_deck, manifest, content, out)

    assert result.slides == 1 and not result.errors
    slide = Presentation(str(out)).slides[0]
    assert slide.shapes.title.text == "Executive Overview: Carrier Invoices"
    assert _shape_text(slide, "Business Need Box") == "Business Need: Fresh need text\nwith a bullet"
    table = next(s for s in slide.shapes if s.has_table).table
    assert [[c.text for c in r.cells] for r in table.rows] == [["Function", "Countries"], ["Finance", "ZA"], ["Ops", "KE"]]
    body = slide.placeholders[1].text_frame
    assert [(p.text, p.level) for p in body.paragraphs] == [("alpha", 0), ("beta", 1)]
    picture = next(s for s in slide.shapes if hasattr(s, "image"))
    assert picture.image.size == (60, 30)


def test_static_fields_are_kept_even_with_values(sample_deck, tmp_path):
    manifest = _fixture_manifest(sample_deck)
    manifest.field("scope").static = True
    content = Content(globals={"subject": "X"}, fields={"scope": [{"Function": "Ops", "Countries": "KE"}], "business_need": "Fresh"})
    result = render(sample_deck, manifest, content, tmp_path / "static.pptx")
    slide = Presentation(str(tmp_path / "static.pptx")).slides[0]
    table = next(s for s in slide.shapes if s.has_table).table
    assert [[c.text for c in r.cells] for r in table.rows] == [["Function", "Countries"], ["Finance", "ZA"], ["", ""]]
    assert any(i.field == "scope" and "kept" in i.message for i in result.issues)
    assert _shape_text(slide, "Business Need Box") == "Business Need: Fresh"

    render(sample_deck, manifest, content, tmp_path / "blank.pptx", field_modes={"scope": "blank"})
    table = next(s for s in Presentation(str(tmp_path / "blank.pptx")).slides[0].shapes if s.has_table).table
    assert [[c.text for c in r.cells] for r in table.rows] == [["Function", "Countries"], ["", ""], ["", ""]]


def test_excluded_slides_are_not_filled(sample_deck, tmp_path):
    manifest = _fixture_manifest(sample_deck)
    manifest.slides.exclude = [1]
    manifest.field("business_need").bindings[0].max_chars = 40
    content = Content(globals={"subject": "X"}, fields={"business_need": "word " * 60})
    result = render(sample_deck, manifest, content, tmp_path / "ex.pptx", spill=True)
    assert result.slides == 1 and result.slide_map == [2]
    assert not any("continued" in i.message for i in result.issues)


def test_missing_values_keep_blank_or_placeholder(sample_deck, tmp_path):
    manifest = _fixture_manifest(sample_deck)
    result = render(sample_deck, manifest, Content(), tmp_path / "kept.pptx", missing="keep")
    messages = [str(i) for i in result.issues]
    assert any("subject" in m and "left in place" in m for m in messages)
    assert any("business_need" in m for m in messages)
    kept = Presentation(str(tmp_path / "kept.pptx")).slides[0]
    assert _shape_text(kept, "Business Need Box").startswith("Business Need: Something long")

    render(sample_deck, manifest, Content(), tmp_path / "blank.pptx", missing="blank")
    blank = Presentation(str(tmp_path / "blank.pptx")).slides[0]
    assert _shape_text(blank, "Business Need Box") == "Business Need: "
    table = next(s for s in blank.shapes if s.has_table).table
    assert [[c.text for c in r.cells] for r in table.rows] == [["Function", "Countries"], ["", ""]]

    result = render(sample_deck, manifest, Content(), tmp_path / "placeholder.pptx")
    shown = Presentation(str(tmp_path / "placeholder.pptx")).slides[0]
    assert _shape_text(shown, "Business Need Box") == "Business Need: [To be completed: Business Need]"
    table = next(s for s in shown.shapes if s.has_table).table
    assert [[c.text for c in r.cells] for r in table.rows] == [["Function", "Countries"], ["[To be completed]", ""]]
    assert shown.placeholders[1].text_frame.text == "[To be completed: First point]"
    assert any("placeholder shown" in str(i) for i in result.issues)
    assert not result.errors


def test_overflow_continues_on_cloned_slides(sample_deck, tmp_path):
    manifest = _fixture_manifest(sample_deck)
    manifest.slides.exclude = []
    spec = manifest.field("first_point")
    spec.bindings[0].max_chars = 40
    lines = [f"- point number {i} with some words" for i in range(1, 7)]
    result = render(sample_deck, manifest, Content(fields={"first_point": "\n".join(lines)}), tmp_path / "long.pptx")

    prs = Presentation(str(tmp_path / "long.pptx"))
    titles = [s.shapes.title.text if s.shapes.title else None for s in prs.slides]
    extra = titles.count("Executive Overview: Demo Integration (cont.)")
    assert extra >= 2 and result.slides == 2 + extra
    assert titles[0] == "Executive Overview: Demo Integration"
    collected = [p.text for s in list(prs.slides)[: 1 + extra] for p in s.placeholders[1].text_frame.paragraphs]
    assert collected == [line[2:] for line in lines]
    assert any("continued on" in str(i) for i in result.issues)


def test_render_appends_check_findings(sample_deck, tmp_path):
    manifest = _fixture_manifest(sample_deck)
    manifest.field("first_point").bindings[0].max_chars = 40
    content = Content(fields={"first_point": "\n".join(f"- point number {i} with some words" for i in range(1, 30))})
    checked = render(sample_deck, manifest, content, tmp_path / "checked.pptx", continue_on=[])
    assert any(i.message.startswith("check overflow") and i.slide == 1 and i.level == "warning" for i in checked.issues)
    quiet = render(sample_deck, manifest, content, tmp_path / "quiet.pptx", continue_on=[], check=False)
    assert not any(i.message.startswith("check ") for i in quiet.issues)


def test_continuation_is_limited_to_listed_slides(sample_deck, tmp_path):
    manifest = _fixture_manifest(sample_deck)
    manifest.slides.exclude = []
    manifest.field("first_point").bindings[0].max_chars = 40
    content = Content(fields={"first_point": "\n".join(f"- point number {i} with some words" for i in range(1, 7))})

    limited = render(sample_deck, manifest, content, tmp_path / "limited.pptx", continue_on=[])
    assert limited.slides == 2
    assert any("about 40 fit" in str(i) for i in limited.issues)

    allowed = render(sample_deck, manifest, content, tmp_path / "allowed.pptx", continue_on=[1])
    assert allowed.slides > 2


def test_bad_bindings_produce_errors_not_crashes(sample_deck, tmp_path):
    manifest = Manifest(
        name="broken",
        fields=[
            FieldSpec(key="a", label="A", bindings=[Binding(slide=9, shape=ShapeRef(id=1))]),
            FieldSpec(key="b", label="B", bindings=[Binding(slide=1, shape=ShapeRef(id=999, name="Nope"))]),
            FieldSpec(key="c", label="C", kind="table", bindings=[Binding(slide=1, shape=ShapeRef(id=3, name="Business Need Box"))]),
        ],
        slides=SlideRules(exclude=[7]),
    )
    content = Content(fields={"a": "x", "b": "y", "c": [{"k": "v"}]})
    result = render(sample_deck, manifest, content, tmp_path / "broken.pptx")
    levels = {i.field: i.level for i in result.issues if i.field}
    assert levels == {"a": "error", "b": "error", "c": "error"}
    assert (tmp_path / "broken.pptx").exists()


def test_cli_skeleton_and_render_via_registry(sample_deck, tmp_path):
    registry_dir = tmp_path / "templates"
    Registry(registry_dir).add("demo", sample_deck, _fixture_manifest(sample_deck))
    runner = CliRunner()

    skeleton = runner.invoke(main, ["skeleton", "demo", "--templates", str(registry_dir), "-o", str(tmp_path / "c.md")])
    assert skeleton.exit_code == 0, skeleton.output
    text = (tmp_path / "c.md").read_text(encoding="utf-8")
    text = text.replace('subject: ""', "subject: Demo Renamed").replace("## business_need\n", "## business_need\nHello there\n")
    (tmp_path / "c.md").write_text(text, encoding="utf-8")

    out = tmp_path / "rendered.pptx"
    rendered = runner.invoke(main, ["render", "demo", str(tmp_path / "c.md"), "-o", str(out), "--templates", str(registry_dir)])
    assert rendered.exit_code == 0, rendered.output
    assert "Wrote" in rendered.output and "warning:" in rendered.output
    slide = Presentation(str(out)).slides[0]
    assert slide.shapes.title.text == "Executive Overview: Demo Renamed"
    assert _shape_text(slide, "Business Need Box") == "Business Need: Hello there"


@pytest.mark.skipif(not CARRIER_DECK.exists(), reason="example deck not available")
def test_carrier_deck_end_to_end(tmp_path):
    manifest = analyze_deck(inspect_deck(CARRIER_DECK)).to_manifest("ntt", source=CARRIER_DECK.name)
    markdown = skeleton_markdown(manifest)
    markdown = markdown.replace('subject: ""', "subject: Lockbox Integration")
    markdown = markdown.replace("## business_need\n", "## business_need\nBank statements arrive daily.\n- CAMT.053 files\n  - one per account\n")
    head, _, tail = markdown.partition("## scope\n")
    section, _, rest = tail.partition("\n## ")
    section = section.replace("|  |  |  |  |", "| Finance | AR | ZA | 3 banks |")
    markdown = head + "## scope\n" + section + "\n## " + rest
    content = load_markdown(markdown, manifest)
    out = tmp_path / "carrier-out.pptx"
    result = render(CARRIER_DECK, manifest, content, out)

    assert not result.errors, [str(i) for i in result.errors]
    deck = inspect_deck(out)
    assert len(deck.slides) == 30
    assert deck.slides[4].title == "Executive Overview"
    need = next(s for s in deck.slides[4].walk() if s.text.startswith("Business Need: "))
    assert need.text == "Business Need: Bank statements arrive daily.\nCAMT.053 files\none per account"
    scope = next(s for s in deck.slides[4].walk() if s.kind == "table" and s.table.cells[0][0] == "Function")
    assert scope.table.cells == [["Function", "BU / Practice", "Operating Countries", "Carriers"], ["Finance", "AR", "ZA", "3 banks"]]
    assert not any("Carrier AP Invoice Integration" in s.title for s in deck.slides if s.title)


def test_field_modes_keep_and_blank_override_content(sample_deck, tmp_path):
    manifest = _fixture_manifest(sample_deck)
    content = Content(fields={"business_need": "Fresh need text", "first_point": "- alpha"})
    out = tmp_path / "modes.pptx"
    result = render(sample_deck, manifest, content, out, field_modes={"business_need": "keep", "scope": "blank", "first_point": "blank"})
    assert not result.errors
    assert any("template content kept" in str(i) for i in result.issues)
    slide = Presentation(str(out)).slides[0]
    assert _shape_text(slide, "Business Need Box").startswith("Business Need: Something long")
    table = next(s for s in slide.shapes if s.has_table).table
    assert [[c.text for c in r.cells] for r in table.rows] == [["Function", "Countries"], ["", ""]]
    body = slide.placeholders[1].text_frame.text
    assert "alpha" not in body and "To be completed" not in body


def test_hidden_and_reordered_slides_with_slide_map(sample_deck, tmp_path):
    manifest = _fixture_manifest(sample_deck)
    manifest.slides.exclude = []
    content = Content(fields={"business_need": "Need"})
    result = render(sample_deck, manifest, content, tmp_path / "order.pptx", order=[2, 1])
    assert result.slides == 2 and result.slide_map == [2, 1]
    prs = Presentation(str(tmp_path / "order.pptx"))
    assert prs.slides[0].shapes.title is None and prs.slides[1].shapes.title.text.startswith("Executive Overview")

    hidden = render(sample_deck, manifest, content, tmp_path / "hidden.pptx", hidden=[2])
    assert hidden.slides == 1 and hidden.slide_map == [1]

    manifest.field("first_point").bindings[0].max_chars = 40
    long = Content(fields={"first_point": "\n".join(f"- point number {i} with some words" for i in range(1, 7))})
    spread = render(sample_deck, manifest, long, tmp_path / "spread.pptx", order=[2, 1])
    assert spread.slide_map[0] == 2 and set(spread.slide_map[1:]) == {1} and len(spread.slide_map) >= 3


def test_titles_replace_slide_titles_without_the_subject(sample_deck, tmp_path):
    manifest = _fixture_manifest(sample_deck)
    out = tmp_path / "titles.pptx"
    result = render(sample_deck, manifest, Content(globals={"subject": "Carrier Invoices"}), out, titles={1: "Overview", 9: "Nothing"})
    assert not result.errors
    assert Presentation(str(out)).slides[0].shapes.title.text == "Overview"


def test_subject_stays_only_on_the_cover(sample_deck, tmp_path):
    manifest = _fixture_manifest(sample_deck)
    content = Content(globals={"subject": "Carrier Invoices"})
    render(sample_deck, manifest, content, tmp_path / "plain.pptx", subject_slides=[])
    assert Presentation(str(tmp_path / "plain.pptx")).slides[0].shapes.title.text == "Executive Overview"
    render(sample_deck, manifest, content, tmp_path / "cover.pptx")
    assert Presentation(str(tmp_path / "cover.pptx")).slides[0].shapes.title.text == "Executive Overview: Carrier Invoices"


def test_long_titles_are_capped_at_sixty_characters(sample_deck, tmp_path):
    from sdgen.render import ExtraSlide

    manifest = _fixture_manifest(sample_deck)
    manifest.slides.exclude = []
    spec = manifest.field("scope").model_copy(update={"key": "extra_long", "label": "Long"})
    title = "Interface inventory for lockbox, statements, payments and payment status across every bank"
    extra = ExtraSlide(key="extra_long", title=title, spec=spec, value=[{"Function": "A", "Countries": "B"}], before=2)
    result = render(sample_deck, manifest, Content(), tmp_path / "long.pptx", extras=[extra], titles={1: title})
    prs = Presentation(str(tmp_path / "long.pptx"))
    for slide in list(prs.slides)[:2]:
        text = slide.shapes.title.text
        assert len(text) <= 60 and title.startswith(text) and not text.endswith(" ")
    assert sum("title shortened" in str(i) for i in result.issues) == 2


def test_extra_slides_are_cloned_from_a_prototype(sample_deck, tmp_path):
    from sdgen.render import ExtraSlide

    manifest = _fixture_manifest(sample_deck)
    manifest.slides.exclude = []
    spec = manifest.field("scope").model_copy(update={"key": "extra_acc", "label": "Acceptance Criteria"})
    extra = ExtraSlide(key="extra_acc", title="Acceptance Criteria", spec=spec, value=[{"Function": "Test", "Countries": "ZA"}], before=2)
    out = tmp_path / "extras.pptx"
    result = render(sample_deck, manifest, Content(globals={"subject": "Carrier Invoices"}, fields={"scope": [{"Function": "Finance", "Countries": "ZA"}]}), out, extras=[extra])
    assert not result.errors
    assert result.slides == 3 and result.slide_map == [1, 1, 2] and result.slide_keys == ["", "extra_acc", ""]
    prs = Presentation(str(out))
    added = prs.slides[1]
    assert added.shapes.title.text == "Acceptance Criteria"
    table = next(s for s in added.shapes if s.has_table).table
    assert [[c.text for c in r.cells] for r in table.rows] == [["Function", "Countries"], ["Test", "ZA"]]

    renamed = extra.model_copy(update={"spec": spec.model_copy(update={"columns": ["Ref", "Scenario"]}), "value": [{"Ref": "1", "Scenario": "Happy path"}]})
    render(sample_deck, manifest, Content(globals={"subject": "X"}), tmp_path / "renamed.pptx", extras=[renamed])
    header = next(s for s in Presentation(str(tmp_path / "renamed.pptx")).slides[1].shapes if s.has_table).table
    assert [[c.text for c in r.cells] for r in header.rows] == [["Ref", "Scenario"], ["1", "Happy path"]]
    original = next(s for s in prs.slides[0].shapes if s.has_table).table
    assert [c.text for c in original.rows[1].cells] == ["Finance", "ZA"]

    tail = render(sample_deck, manifest, Content(), tmp_path / "tail.pptx", extras=[extra.model_copy(update={"before": 0, "value": None})])
    assert tail.slide_keys == ["", "", "extra_acc"] and any("placeholder shown" in str(i) for i in tail.issues)


def test_extra_table_gets_the_columns_of_its_field(sample_deck, tmp_path):
    from sdgen.render import ExtraSlide

    manifest = _fixture_manifest(sample_deck)
    manifest.slides.exclude = []
    columns = ["#", "Party", "Flow", "Direction", "Source", "Target", "Encryption", "Cut-off"]
    spec = manifest.field("scope").model_copy(update={"key": "extra_inv", "label": "Interface inventory", "columns": columns})
    rows = [{"#": "1", "Party": "BoA", "Flow": "Lockbox", "Direction": "in", "Source": "a", "Target": "b", "Encryption": "PGP", "Cut-off": "09:00"}]
    result = render(sample_deck, manifest, Content(), tmp_path / "wide.pptx", extras=[ExtraSlide(key="extra_inv", title="Interface inventory", spec=spec, value=rows, before=2)])
    assert not result.errors
    table = next(s for s in Presentation(str(tmp_path / "wide.pptx")).slides[1].shapes if s.has_table).table
    assert [c.text for c in table.rows[0].cells] == columns and [c.text for c in table.rows[1].cells] == list(rows[0].values())
    original = next(s for s in Presentation(str(sample_deck)).slides[0].shapes if s.has_table)
    assert sum(c.width for c in table.columns) == sum(c.width for c in original.table.columns)


def test_flows_are_drawn_into_the_image_slot(sample_deck, tmp_path):
    from pptx.shapes.picture import Picture

    from sdgen.flow import FlowEdge, FlowNode, FlowSpec

    manifest = _fixture_manifest(sample_deck)
    image_key = next(f.key for f in manifest.fields if f.kind == "image")
    flow = FlowSpec(nodes=[FlowNode(id="a", label="Bank", lane="source"), FlowNode(id="b", label="S/4HANA", lane="target")], edges=[FlowEdge(source="a", target="b", label="file")])
    out = tmp_path / "flow.pptx"
    result = render(sample_deck, manifest, Content(), out, flows={image_key: flow})
    assert not result.errors and any("diagram drawn from the brief" in str(i) for i in result.issues)
    slide = Presentation(str(out)).slides[0]
    assert not any(isinstance(s, Picture) for s in slide.shapes)
    assert sum(s.name.startswith(f"Flow {image_key} node") for s in slide.shapes) == 2

    image = tmp_path / "up.png"
    Image.new("RGB", (40, 30), "blue").save(image)
    kept = render(sample_deck, manifest, Content(fields={image_key: ImageValue(path=str(image))}), tmp_path / "upload.pptx", flows={image_key: flow})
    upload_slide = Presentation(str(tmp_path / "upload.pptx")).slides[0]
    assert any(isinstance(s, Picture) for s in upload_slide.shapes) and not any(s.name.startswith("Flow") for s in upload_slide.shapes)
    assert not any("diagram drawn" in str(i) for i in kept.issues)


def test_composite_box_spills_onto_cleaned_copies(sample_deck, tmp_path):
    manifest = _fixture_manifest(sample_deck)
    manifest.slides.exclude = []
    manifest.field("business_need").bindings[0].max_chars = 60
    manifest.field("first_point").bindings[0].max_chars = 20
    need = chr(10).join(f"Paragraph number {i} with enough words to matter." for i in range(1, 7))
    points = "- alpha" + chr(10) + "- beta gamma delta" + chr(10) + "- epsilon zeta eta theta"
    content = Content(fields={"business_need": need, "scope": [{"Function": "Finance", "Countries": "ZA"}], "first_point": points})

    kept = render(sample_deck, manifest, content, tmp_path / "shrunk.pptx", continue_on=[])
    assert kept.slides == 2 and any("shrunk to fit" in str(i) for i in kept.issues)

    spilled = render(sample_deck, manifest, content, tmp_path / "spill.pptx", continue_on=[], spill=True)
    prs = Presentation(str(tmp_path / "spill.pptx"))
    copies = len(prs.slides) - 2
    assert copies >= 1 and spilled.slide_map == [1] * (copies + 1) + [2]
    assert any(f"continued on {copies} extra" in str(i) for i in spilled.issues)
    original = prs.slides[0]
    assert [c.text for c in next(s for s in original.shapes if s.has_table).table.rows[1].cells] == ["Finance", "ZA"]
    assert _shape_text(original, "Business Need Box") == "Business Need: Paragraph number 1 with enough words to matter."
    need_box = next(s for s in original.shapes if s.name == "Business Need Box")
    copy = prs.slides[1]
    assert copy.shapes.title.text.endswith("(cont.)")
    names = {s.name for s in copy.shapes}
    assert not any(s.has_table for s in copy.shapes) and not any(hasattr(s, "image") for s in copy.shapes) and "Scope Label" not in names
    box = next(s for s in copy.shapes if s.name == "Business Need Box")
    assert box.top < need_box.top and box.height > need_box.height
    assert _shape_text(copy, "Business Need Box").startswith("Business Need: Paragraph number 2")
    texts = [_shape_text(s, "Business Need Box").removeprefix("Business Need: ") for s in list(prs.slides)[: copies + 1]]
    assert chr(10).join(t for t in texts if t) == need


def test_clear_shapes_empties_template_text_before_copies(sample_deck, tmp_path):
    manifest = _fixture_manifest(sample_deck)
    label_id = next(s.shape_id for s in Presentation(str(sample_deck)).slides[0].shapes if s.name == "Scope Label")
    out = tmp_path / "clear.pptx"
    result = render(sample_deck, manifest, Content(fields={"business_need": "Need"}), out, clear_shapes=[(1, label_id), (1, 9999)])
    assert not result.errors and any("template text cleared" in str(i) for i in result.issues) and any("not found" in str(i) for i in result.issues)
    assert _shape_text(Presentation(str(out)).slides[0], "Scope Label") == ""


def _overview_manifest(shapes) -> Manifest:
    def spec(key, shape, max_chars=None):
        return FieldSpec(key=key, label=key, bindings=[Binding(slide=1, shape=ShapeRef(id=shape.shape_id), max_chars=max_chars)])

    return Manifest(name="overview", fields=[spec("need", shapes["need"], 60), spec("left", shapes["left"]), spec("inner", shapes["inner"])])


def test_top_block_expands_and_pushes_the_rest(tmp_path):
    from pptx.util import Inches
    from test_layout import overview_deck

    prs, _, shapes = overview_deck()
    deck = tmp_path / "overview.pptx"
    prs.save(deck)
    need = "Lockbox files arrive daily from three banks over SWIFT. They must be posted automatically without manual work."
    content = Content(fields={"need": need, "left": "Left text", "inner": "Inner text"})
    result = render(deck, _overview_manifest(shapes), content, tmp_path / "out.pptx", spill=True)
    assert result.slides == 2 and result.slide_map == [1, 1] and not result.errors
    assert any("top block grown" in i.message and "1 block(s) moved" not in i.message for i in result.issues)
    first, second = Presentation(str(tmp_path / "out.pptx")).slides
    on_first = {s.name: s for s in first.shapes}
    assert {"Need Bar", "Need Box", "Footer"} <= set(on_first) and not {"Left Bar", "Left Box", "Right Inner"} & set(on_first)
    assert on_first["Need Box"].top + on_first["Need Box"].height == Inches(6.95) and on_first["Need Box"].text_frame.text == need
    assert not first.shapes.title.text.endswith("(cont.)")
    on_second = {s.name: s for s in second.shapes}
    assert {"Left Bar", "Left Box", "Right Bar", "Right Container", "Right Inner", "Footer"} <= set(on_second) and "Need Box" not in on_second
    assert on_second["Left Bar"].top == Inches(1.2) and on_second["Left Box"].top == Inches(1.55) and on_second["Footer"].top == Inches(7.0)
    assert second.shapes.title.text.endswith("(cont.)") and on_second["Left Box"].text_frame.text == "Left text"


def test_expand_continues_on_copies_before_the_pushed_slide(tmp_path):
    from test_layout import overview_deck

    prs, _, shapes = overview_deck()
    deck = tmp_path / "overview.pptx"
    prs.save(deck)
    need = "\n".join(f"Paragraph {i}: lockbox files arrive daily and are posted automatically." for i in range(1, 9))
    content = Content(fields={"need": need, "left": "Left text", "inner": "Inner text"})
    result = render(deck, _overview_manifest(shapes), content, tmp_path / "out.pptx", spill=True)
    slides = list(Presentation(str(tmp_path / "out.pptx")).slides)
    assert len(slides) >= 3 and result.slide_map == [1] * len(slides)
    boxes = [{s.name: s for s in sl.shapes}.get("Need Box") for sl in slides]
    assert all(box is not None for box in boxes[:-1]) and boxes[-1] is None
    assert "\n".join(box.text_frame.text for box in boxes[:-1]) == need
    assert all(sl.shapes.title.text.endswith("(cont.)") for sl in slides[1:]) and "Left Box" in {s.name for s in slides[-1].shapes}


def test_other_overflowing_block_continues_after_the_pushed_slide(tmp_path):
    from test_layout import overview_deck

    prs, _, shapes = overview_deck()
    deck = tmp_path / "overview.pptx"
    prs.save(deck)
    manifest = _overview_manifest(shapes)
    manifest.field("left").bindings[0].max_chars = 20
    need = "Lockbox files arrive daily from three banks over SWIFT. They must be posted automatically without manual work."
    left = chr(10).join(f"Left paragraph {i} with a few words." for i in range(1, 9))
    content = Content(fields={"need": need, "left": left, "inner": "Inner text"})
    result = render(deck, manifest, content, tmp_path / "out.pptx", spill=True)
    slides = list(Presentation(str(tmp_path / "out.pptx")).slides)
    assert result.slide_map == [1] * len(slides) and len(slides) >= 3 and not result.errors
    first, pushed, tail = slides[0], slides[1], slides[2:]
    assert "Need Box" in {s.name for s in first.shapes} and "Left Box" not in {s.name for s in first.shapes}
    pushed_names = {s.name for s in pushed.shapes}
    assert {"Left Bar", "Left Box", "Right Inner"} <= pushed_names and "Need Box" not in pushed_names
    for copy in tail:
        names = {s.name for s in copy.shapes}
        assert "Left Box" in names and "Right Inner" not in names and "Need Box" not in names
    texts = [{s.name: s for s in sl.shapes}["Left Box"].text_frame.text for sl in slides[1:]]
    assert chr(10).join(t for t in texts if t) == left


def test_long_paragraph_is_cut_at_sentence_ends():
    from sdgen.fill.text import parse_blocks
    from sdgen.render import _split_blocks

    text = "First sentence of the story. Second sentence follows it. Third one is here. Fourth closes it."
    chunks = _split_blocks(parse_blocks(text), 40)
    assert [" ".join(b.text for b in chunk) for chunk in chunks] == ["First sentence of the story.", "Second sentence follows it.", "Third one is here. Fourth closes it."]
    assert _split_blocks(parse_blocks("One sentence only that is long enough to exceed the limit by itself"), 40) == [parse_blocks("One sentence only that is long enough to exceed the limit by itself")]



def test_inner_text_box_grows_with_its_block_on_the_copy(tmp_path):
    from pptx.util import Inches
    from test_layout import overview_deck

    prs, _, shapes = overview_deck()
    deck = tmp_path / "overview.pptx"
    prs.save(deck)
    manifest = _overview_manifest(shapes)
    manifest.field("inner").bindings[0].max_chars = 45
    inner = chr(10).join(f"- Point {i} is short." for i in range(1, 7))
    content = Content(fields={"need": "Need text.", "left": "Left text", "inner": inner})
    result = render(deck, manifest, content, tmp_path / "out.pptx", spill=True)
    slides = list(Presentation(str(tmp_path / "out.pptx")).slides)
    assert not result.errors and len(slides) == 2
    original = {s.name: s for s in slides[0].shapes}["Right Inner"]
    copy = {s.name: s for s in slides[1].shapes}
    assert "Left Box" not in copy and copy["Right Inner"].height > original.height + Inches(3)
    assert copy["Right Container"].top + copy["Right Container"].height == copy["Right Inner"].top + copy["Right Inner"].height + Inches(0.15)
    texts = [{s.name: s for s in sl.shapes}["Right Inner"].text_frame.text for sl in slides]
    assert chr(10).join(texts) == chr(10).join(f"Point {i} is short." for i in range(1, 7))


def _table_deck(tmp_path):
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Scope"
    slide.shapes.title.height = Inches(0.6)
    frame = slide.shapes.add_table(2, 2, Inches(1), Inches(1.2), Inches(8), Inches(0.8))
    frame.name = "Scope Table"
    frame.table.cell(0, 0).text, frame.table.cell(0, 1).text = "Function", "Bank"
    frame.table.cell(1, 0).text = "x"
    below = slide.shapes.add_textbox(Inches(1), Inches(2.3), Inches(8), Inches(1))
    below.name = "Below Box"
    below.text_frame.text = "Success"
    side = slide.shapes.add_textbox(Inches(9.5), Inches(1.2), Inches(3), Inches(0.8))
    side.name = "Side Box"
    side.text_frame.text = "Legend"
    deck = tmp_path / "table.pptx"
    prs.save(deck)
    spec = FieldSpec(key="scope", label="Scope", kind="table", columns=["Function", "Bank"], bindings=[Binding(slide=1, shape=ShapeRef(id=frame.shape_id))])
    return deck, spec


def test_table_rows_continue_on_a_cleaned_copy(tmp_path):
    deck, spec = _table_deck(tmp_path)
    rows = [{"Function": f"Inbound {i}", "Bank": f"Bank {i}"} for i in range(1, 6)]
    result = render(deck, Manifest(name="t", fields=[spec]), Content(fields={"scope": rows}), tmp_path / "out.pptx", spill=True)
    slides = list(Presentation(str(tmp_path / "out.pptx")).slides)
    assert not result.errors and len(slides) == 2 and any("4 row(s) continue" in i.message for i in result.issues)
    first = next(s for s in slides[0].shapes if s.has_table)
    second = next(s for s in slides[1].shapes if s.has_table)
    assert len(first.table.rows) == 2 and [c.text for c in first.table.rows[1].cells] == ["Inbound 1", "Bank 1"]
    assert len(second.table.rows) == 5 and [c.text for c in second.table.rows[0].cells] == ["Function", "Bank"]
    assert [c.text for c in list(second.table.rows)[-1].cells] == ["Inbound 5", "Bank 5"]
    assert not {"Below Box", "Side Box"} & {s.name for s in slides[1].shapes} and slides[1].shapes.title.text.endswith("(cont.)")


def test_extra_table_rows_continue_too(tmp_path):
    from sdgen.render import ExtraSlide

    deck, spec = _table_deck(tmp_path)
    extra_spec = spec.model_copy(update={"key": "extra", "label": "Acceptance"})
    rows = [{"Function": f"Case {i}", "Bank": "ok"} for i in range(1, 25)]
    extra = ExtraSlide(key="extra", title="Acceptance", spec=extra_spec, value=rows, before=0)
    result = render(deck, Manifest(name="t", fields=[spec]), Content(fields={"scope": [{"Function": "One", "Bank": "B"}]}), tmp_path / "out.pptx", extras=[extra], spill=True)
    slides = list(Presentation(str(tmp_path / "out.pptx")).slides)
    assert not result.errors and len(slides) >= 3 and result.slide_keys == [""] + ["extra"] * (len(slides) - 1)
    assert slides[1].shapes.title.text == "Acceptance" and all(s.shapes.title.text == "Acceptance (cont.)" for s in slides[2:])
    tables = [next(s for s in sl.shapes if s.has_table) for sl in slides[1:]]
    assert sum(len(t.table.rows) - 1 for t in tables) == 24 and len(tables[0].table.rows) > 2
    # The prototype's note and legend belong to its own section: no extra slide shows them and none is spent on them alone.
    assert not any({"Below Box", "Side Box"} & {s.name for s in sl.shapes} for sl in slides[1:])


def _chain_deck(tmp_path):
    """A table prototype, a middle slide and a tail slide: an extra placed before the tail must travel past the middle."""
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    proto = prs.slides.add_slide(prs.slide_layouts[5])
    proto.shapes.title.text = "Scope"
    frame = proto.shapes.add_table(2, 2, Inches(1), Inches(1.2), Inches(8), Inches(0.8))
    frame.name = "Scope Table"
    frame.table.cell(0, 0).text, frame.table.cell(0, 1).text = "Function", "Bank"
    for title in ("Middle", "Tail"):
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = title
    deck = tmp_path / "chain.pptx"
    prs.save(deck)
    spec = FieldSpec(key="scope", label="Scope", kind="table", columns=["Function", "Bank"], bindings=[Binding(slide=1, shape=ShapeRef(id=frame.shape_id))])
    return deck, spec


def test_extra_chains_stay_together_and_in_order(tmp_path):
    from sdgen.render import ExtraSlide

    deck, spec = _chain_deck(tmp_path)
    extras = []
    for key in ("first", "second"):
        rows = [{"Function": f"{key} {i}", "Bank": "ok"} for i in range(1, 25)]
        extras.append(ExtraSlide(key=key, title=key.title(), spec=spec.model_copy(update={"key": key, "label": key.title()}), value=rows, before=3))
    result = render(deck, Manifest(name="t", fields=[spec]), Content(fields={"scope": [{"Function": "One", "Bank": "B"}]}), tmp_path / "out.pptx", extras=extras, spill=True)
    assert not result.errors
    keys = result.slide_keys
    firsts, seconds = keys.count("first"), keys.count("second")
    assert firsts >= 2 and seconds >= 2 and keys == ["", ""] + ["first"] * firsts + ["second"] * seconds + [""]
    slides = list(Presentation(str(tmp_path / "out.pptx")).slides)
    assert [s.shapes.title.text for s in slides][1::len(slides) - 2] == ["Middle", "Tail"]
    tables = [next(s for s in sl.shapes if s.has_table) for sl in slides[2:-1]]
    written = [r.cells[0].text for t in tables for r in list(t.table.rows)[1:]]
    assert written == [f"first {i}" for i in range(1, 25)] + [f"second {i}" for i in range(1, 25)]


def test_extra_table_columns_follow_their_text(tmp_path):
    from sdgen.render import ExtraSlide

    deck, spec = _table_deck(tmp_path)
    steps = spec.model_copy(update={"key": "steps", "label": "Cutover", "columns": ["#", "Step"]})
    rows = [{"#": "1", "Step": "Exchange SSH keys and pin the host keys on both doors"}]
    same = spec.model_copy(update={"key": "same", "label": "Same"})
    extras = [ExtraSlide(key="steps", title="Cutover", spec=steps, value=rows, before=0), ExtraSlide(key="same", title="Same", spec=same, value=[{"Function": "Lockbox", "Bank": "BoA"}], before=0)]
    result = render(deck, Manifest(name="t", fields=[spec]), Content(), tmp_path / "out.pptx", extras=extras)
    assert not result.errors and result.slide_keys == ["", "steps", "same"]
    slides = list(Presentation(str(tmp_path / "out.pptx")).slides)
    original = next(s for s in Presentation(str(deck)).slides[0].shapes if s.has_table).table
    weighted = next(s for s in slides[1].shapes if s.has_table).table
    kept = next(s for s in slides[2].shapes if s.has_table).table
    assert weighted.columns[0].width < weighted.columns[1].width / 4 and sum(c.width for c in weighted.columns) == sum(c.width for c in original.columns)
    assert [c.width for c in kept.columns] == [c.width for c in original.columns]


def test_subject_stripped_from_titles_is_not_reported_missing(sample_deck, tmp_path):
    manifest = _fixture_manifest(sample_deck)
    result = render(sample_deck, manifest, Content(globals={"subject": "Carrier Invoices"}), tmp_path / "plain.pptx", subject_slides=[])
    assert not any("not found in the template" in i.message for i in result.issues)


def test_template_picture_under_written_text_is_removed(tmp_path):
    from PIL import Image as PILImage
    from pptx.util import Inches
    from test_layout import overview_deck

    prs, slide, shapes = overview_deck()
    image = tmp_path / "shot.png"
    PILImage.new("RGB", (80, 40), "grey").save(image)
    under = slide.shapes.add_picture(str(image), Inches(1), Inches(2), Inches(4), Inches(1.2))
    under.name = "Screenshot"
    aside = slide.shapes.add_picture(str(image), Inches(5.5), Inches(6.2), Inches(1.5), Inches(0.5))
    aside.name = "Logo"
    deck = tmp_path / "pictures.pptx"
    prs.save(deck)
    manifest = _overview_manifest(shapes)
    result = render(deck, manifest, Content(fields={"need": "Written need.", "left": "Left", "inner": "Inner"}), tmp_path / "out.pptx")
    names = {s.name for s in Presentation(str(tmp_path / "out.pptx")).slides[0].shapes}
    assert "Screenshot" not in names and "Logo" in names
    assert any("template picture removed" in i.message and str(under.shape_id) in i.message for i in result.issues)

    render(deck, manifest, Content(fields={"left": "Left", "inner": "Inner"}), tmp_path / "kept.pptx")
    assert "Screenshot" in {s.name for s in Presentation(str(tmp_path / "kept.pptx")).slides[0].shapes}


def _detail_deck(tmp_path):
    """An executive overview with a small table, a plain text slide and a plain table slide to clone for the details."""
    from pptx.util import Inches
    from test_layout import overview_deck

    prs, slide, shapes = overview_deck()
    grid = slide.shapes.add_table(2, 2, Inches(5.2), Inches(5.85), Inches(4.3), Inches(0.6))
    grid.name = "Scope Table"
    grid.table.cell(0, 0).text, grid.table.cell(0, 1).text, grid.table.cell(1, 0).text = "Function", "Bank", "x"
    notes = prs.slides.add_slide(prs.slide_layouts[5])
    notes.shapes.title.text = "Notes"
    body = notes.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(9), Inches(5.5))
    body.name, body.text_frame.text = "Body", "body text"
    tables = prs.slides.add_slide(prs.slide_layouts[5])
    tables.shapes.title.text = "Table"
    frame = tables.shapes.add_table(2, 3, Inches(0.5), Inches(1.2), Inches(9), Inches(0.8))
    frame.name = "Big Table"
    for column, name in enumerate(["A", "B", "C"]):
        frame.table.cell(0, column).text = name
    deck = tmp_path / "detail.pptx"
    prs.save(deck)

    def field(key, shape, slide_number=1, max_chars=None, kind="text", columns=None):
        return FieldSpec(key=key, label=key.title(), kind=kind, columns=columns or [], bindings=[Binding(slide=slide_number, shape=ShapeRef(id=shape.shape_id), max_chars=max_chars)])

    manifest = Manifest(
        name="detail",
        fields=[
            field("need", shapes["need"], max_chars=60),
            field("left", shapes["left"]),
            field("inner", shapes["inner"], max_chars=45),
            field("scope", grid, kind="table", columns=["Function", "Bank"]),
            field("body", body, 2),
            field("big", frame, 3, kind="table", columns=["A", "B", "C"]),
        ],
    )
    text_proto = Binding(slide=2, shape=ShapeRef(id=body.shape_id))
    table_proto = Binding(slide=3, shape=ShapeRef(id=frame.shape_id))
    details = {
        "need": FieldSpec(key="need_details", label="Need", kind="bullets", bindings=[text_proto]),
        "left": FieldSpec(key="left_details", label="Left", kind="bullets", bindings=[text_proto.model_copy()]),
        "inner": FieldSpec(key="inner_details", label="Inner", kind="bullets", bindings=[text_proto.model_copy()]),
        "scope": FieldSpec(key="scope_details", label="Scope", kind="table", columns=["Function", "Bank"], bindings=[table_proto]),
    }
    return deck, manifest, details


def _plain_fields():
    return {"body": "Body text", "big": [{"A": "1", "B": "2", "C": "3"}]}


def test_composite_fields_get_detail_slides_in_reading_order(tmp_path):
    deck, manifest, details = _detail_deck(tmp_path)
    need = "\n".join(f"Paragraph {i}: lockbox files arrive daily and are posted automatically." for i in range(1, 9))
    inner = "\n".join(f"- Point {i} is short." for i in range(1, 7))
    content = Content(fields={"need": need, "left": "Left text", "left_details": "- more on the left\n- and more", "inner": inner, "scope": [{"Function": "One", "Bank": "B"}], **_plain_fields()})
    result = render(deck, manifest, content, tmp_path / "out.pptx", details=details, spill=True)
    slides = list(Presentation(str(tmp_path / "out.pptx")).slides)
    assert not result.errors
    assert result.slide_keys == ["", "need_details", "left_details", "inner_details", "", ""] and result.slide_map == [1, 1, 1, 1, 2, 3]
    first = {s.name: s for s in slides[0].shapes}
    assert {"Need Box", "Left Box", "Right Inner", "Scope Table"} <= set(first) and not slides[0].shapes.title.text.endswith("(cont.)")
    assert first["Need Box"].text_frame.text.startswith("Paragraph 1") and "Paragraph 8" not in first["Need Box"].text_frame.text
    assert [s.shapes.title.text for s in slides[1:4]] == ["Need", "Left", "Inner"]
    bodies = [{s.name: s for s in sl.shapes}["Body"].text_frame.text for sl in slides[1:4]]
    assert bodies == [need, "more on the left\nand more", "\n".join(f"Point {i} is short." for i in range(1, 7))]
    assert any("full content follows on a detail slide" in str(i) for i in result.issues)


def test_composite_table_is_trimmed_at_a_row_boundary_and_detailed(tmp_path):
    from pptx.util import Inches

    deck, manifest, details = _detail_deck(tmp_path)
    rows = [{"Function": f"Flow {i}", "Bank": f"Bank {i}"} for i in range(1, 7)]
    content = Content(fields={"need": "Need.", "left": "L", "inner": "I", "scope": rows, **_plain_fields()})
    result = render(deck, manifest, content, tmp_path / "out.pptx", details=details, spill=True)
    slides = list(Presentation(str(tmp_path / "out.pptx")).slides)
    assert not result.errors and result.slide_keys == ["", "scope_details", "", ""]
    first = next(s for s in slides[0].shapes if s.has_table).table
    assert 1 < len(first.rows) < 7 and [c.text for c in first.rows[1].cells] == ["Flow 1", "Bank 1"]
    detail = next(s for s in slides[1].shapes if s.has_table).table
    assert len(detail.columns) == 2 and sum(c.width for c in detail.columns) == Inches(9)
    assert [c.text for c in detail.rows[0].cells] == ["Function", "Bank"] and len(detail.rows) == 7
    assert [c.text for c in list(detail.rows)[-1].cells] == ["Flow 6", "Bank 6"] and slides[1].shapes.title.text == "Scope"


def test_detail_slide_still_continues_on_a_copy(tmp_path):
    deck, manifest, details = _detail_deck(tmp_path)
    details["need"].bindings[0].max_chars = 120
    need = "\n".join(f"Paragraph {i}: lockbox files arrive daily and are posted automatically." for i in range(1, 9))
    content = Content(fields={"need": need, "left": "L", "inner": "I", "scope": [{"Function": "One", "Bank": "B"}], **_plain_fields()})
    result = render(deck, manifest, content, tmp_path / "out.pptx", details=details, spill=True)
    slides = list(Presentation(str(tmp_path / "out.pptx")).slides)
    assert not result.errors and result.slide_keys[:2] == ["", "need_details"] and result.slide_keys.count("need_details") >= 2
    assert result.slide_keys[-2:] == ["", ""] and result.slide_map[-2:] == [2, 3]
    assert slides[1].shapes.title.text == "Need" and slides[2].shapes.title.text == "Need (cont.)"
    texts = [{s.name: s for s in sl.shapes}["Body"].text_frame.text for sl in slides[1 : result.slide_keys.count("need_details") + 1]]
    assert "\n".join(t for t in texts if t) == need


def test_planned_extras_follow_the_detail_slides(tmp_path):
    from sdgen.render import ExtraSlide

    deck, manifest, details = _detail_deck(tmp_path)
    spec = details["scope"].model_copy(update={"key": "extra_acc", "label": "Acceptance"})
    extra = ExtraSlide(key="extra_acc", title="Acceptance", spec=spec, value=[{"Function": "Case", "Bank": "ok"}], before=2)
    content = Content(fields={"need": "Need.", "left": "L", "left_details": "- more", "inner": "I", "scope": [{"Function": "One", "Bank": "B"}], **_plain_fields()})
    result = render(deck, manifest, content, tmp_path / "out.pptx", details=details, extras=[extra], spill=True)
    assert not result.errors
    assert result.slide_keys == ["", "left_details", "extra_acc", "", ""] and result.slide_map == [1, 1, 1, 2, 3]


def test_table_pushes_the_kept_block_below_on_the_copy(tmp_path):
    from pptx.util import Inches

    deck, spec = _table_deck(tmp_path)
    prs = Presentation(str(deck))
    below = next(s for s in prs.slides[0].shapes if s.name == "Below Box")
    note = FieldSpec(key="note", label="Note", kind="text", bindings=[Binding(slide=1, shape=ShapeRef(id=below.shape_id), max_chars=200)])
    rows = [{"Function": f"Inbound {i}", "Bank": f"Bank {i}"} for i in range(1, 6)]
    text = chr(10).join(f"Note paragraph {i} with a few words." for i in range(1, 7))
    result = render(deck, Manifest(name="t", fields=[spec, note]), Content(fields={"scope": rows, "note": text}), tmp_path / "out.pptx", spill=True)
    slides = list(Presentation(str(tmp_path / "out.pptx")).slides)
    assert not result.errors and len(slides) == 2
    copy = {s.name: s for s in slides[1].shapes}
    table = next(s for s in slides[1].shapes if s.has_table)
    assert len(table.table.rows) == 5 and "Below Box" in copy
    assert copy["Below Box"].top >= table.top + table.height and copy["Below Box"].top + copy["Below Box"].height <= Inches(7.15) + 1
