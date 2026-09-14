import importlib.util
import re
from pathlib import Path

import pytest
from pptx import Presentation
from streamlit.testing.v1 import AppTest

from sdgen.analyze import analyze_deck
from sdgen.blueprint import derive_blueprint
from sdgen.design import DesignStore
from sdgen.icons import catalogue
from sdgen.inventory import inspect_deck
from sdgen.manifest import GlobalSpec
from sdgen.registry import Registry

APP = Path(__file__).resolve().parent.parent / "ui" / "app.py"


@pytest.fixture
def registry_with_demo(sample_deck, tmp_path, monkeypatch):
    deck = inspect_deck(sample_deck)
    analysis = analyze_deck(deck)
    manifest = analysis.to_manifest("demo")
    manifest.globals.append(GlobalSpec(key="subject", label="Subject", replaces="Demo Integration"))
    blueprint = derive_blueprint(deck, analysis, manifest, "demo")
    root = tmp_path / "templates"
    Registry(root).add("demo", sample_deck, manifest, blueprint=blueprint)
    monkeypatch.setenv("SDGEN_TEMPLATES", str(root))
    monkeypatch.setenv("SDGEN_DESIGNS", str(tmp_path / "designs"))
    monkeypatch.setenv("SDGEN_LLM", "mock")
    return root


def test_templates_page_without_upload_shows_guidance(tmp_path, monkeypatch):
    monkeypatch.setenv("SDGEN_TEMPLATES", str(tmp_path / "empty"))
    app = AppTest.from_file(str(APP), default_timeout=60).run()
    assert not app.exception
    assert app.sidebar.radio[0].value == "Templates"
    assert any("Upload a deck" in i.value for i in app.info)


def test_brief_lint_offers_to_split_a_pasted_line(registry_with_demo, tmp_path):
    app = AppTest.from_file(str(APP), default_timeout=60).run()
    app.text_input(key="design_name:demo").input("Lint").run()
    state_key = "design:demo:lint"
    app.text_area(key=f"{state_key}:0:b:open_questions").input("Which bank? | Treasury Is PGP needed? | Security").run()
    assert any("one entry per line" in w.value for w in app.warning)
    app.button(key=f"{state_key}:split:open_questions").click().run()
    assert not app.exception
    assert app.session_state[state_key].brief.open_questions.count("\n") == 1
    assert app.text_area(key=f"{state_key}:1:b:open_questions").value.count("\n") == 1
    assert not any("one entry per line" in w.value for w in app.warning)
    assert any("split into 2 lines" in i.value for i in app.info)


def test_design_page_drafts_and_generates(registry_with_demo, tmp_path):
    app = AppTest.from_file(str(APP), default_timeout=60).run()
    assert not app.exception
    assert app.sidebar.radio[0].value == "New design"
    assert app.selectbox(key="llm_provider").value == "mock"
    assert app.selectbox(key="design_template").value == "demo"

    app.text_input(key="design_name:demo").input("CAMT 053").run()
    state_key = "design:demo:camt-053"
    assert state_key in app.session_state
    steps = _steps(app)
    assert steps[:3] == ["1. Brief", "2. Section plan", "3. Write the slides"] and steps[-2:] == ["5. Review sections", "6. Generate"]

    app.text_input(key=f"{state_key}:0:b:subject").input("Lockbox Integration")
    app.text_area(key=f"{state_key}:0:b:about").input("Bank statements arrive daily. They must be posted automatically.")
    app.text_input(key=f"{state_key}:0:fact:countries").input("ZA, KE")
    app.run()
    assert app.session_state[state_key].brief.facts == {"countries": "ZA, KE"}
    assert (tmp_path / "designs" / "camt-053" / "brief.md").read_text(encoding="utf-8").rstrip().endswith("## fact:countries\nZA, KE")
    app.button(key=f"{state_key}:complete:brief").click().run()
    assert app.session_state[f"{state_key}:step"] == "plan" and app.session_state[state_key].completed == ["brief"]
    assert next(e.icon for e in app.status if e.label.startswith("1.")) == chr(0x2705)
    assert (tmp_path / "designs" / "camt-053" / "design.yaml").read_text(encoding="utf-8").find("- brief") > 0
    app.button(key=f"{state_key}:draft").click().run()
    assert not app.exception
    design = app.session_state[state_key]
    assert "[Draft] Bank statements" in design.content_markdown
    assert any(s.value.startswith("Drafted 3 of 3 fields with mock") for s in app.success)
    assert (tmp_path / "designs" / "camt-053" / "content.md").is_file()
    labels = _labels(app)
    assert "3. Write the slides: written with mock" in labels
    assert app.session_state[f"{state_key}:step"] == "plan"
    assert any(e.label.startswith("3.") for e in app.expander)
    app.button(key=f"{state_key}:reopen:brief").click().run()
    assert app.session_state[f"{state_key}:step"] == "brief" and app.session_state[state_key].completed == []
    assert any(re.fullmatch(r"5\. Review sections: (\d+) of \1 written", label) for label in labels)

    app.button(key=f"{state_key}:generate").click().run()
    assert not app.exception
    assert any("Generated" in s.value for s in app.success)
    file_name, data, issues, slides = app.session_state[f"{state_key}:output"]
    assert file_name == "lockbox_integration.pptx" and slides == 2

    saved = tmp_path / "generated.pptx"
    saved.write_bytes(data)
    slide = Presentation(str(saved)).slides[0]
    assert slide.shapes.title.text == "Executive Overview"
    need = next(s for s in slide.shapes if s.name == "Business Need Box").text_frame.text
    assert need.startswith("Business Need: [Draft] Bank statements arrive daily.")

    app.text_area(key=f"{state_key}:1:b:open_questions").input("Which bank sends BAI2? | Treasury").run()
    app.button(key=f"{state_key}:generate").click().run()
    assert not app.exception
    _, data, _, slides = app.session_state[f"{state_key}:output"]
    assert slides == 3
    saved.write_bytes(data)
    titles = [s.shapes.title.text for s in Presentation(str(saved)).slides if s.shapes.title is not None]
    assert any(t.startswith("Open Questions") for t in titles)
    assert "## open_questions\nWhich bank sends BAI2? | Treasury" in (tmp_path / "designs" / "camt-053" / "brief.md").read_text(encoding="utf-8")

    section_key = Registry(registry_with_demo).load("demo").blueprint.sections[0].key
    app.radio(key=f"{state_key}:1:mode:{section_key}").set_value("keep").run()
    app.button(key=f"{state_key}:generate").click().run()
    assert not app.exception
    assert app.session_state[state_key].modes == {section_key: "keep"}
    saved.write_bytes(app.session_state[f"{state_key}:output"][1])
    kept = Presentation(str(saved)).slides[0]
    need = next(s for s in kept.shapes if s.name == "Business Need Box").text_frame.text
    assert need.startswith("Business Need: Something long")

    app.button(key=f"{state_key}:plan").click().run()
    assert not app.exception and app.session_state[f"{state_key}:proposed_plan"] is not None
    app.button(key=f"{state_key}:plan_confirm").click().run()
    assert not app.exception
    assert app.session_state[state_key].plan is not None and (tmp_path / "designs" / "camt-053" / "plan.yaml").is_file()
    assert f"{state_key}:proposed_plan" not in app.session_state
    app.button(key=f"{state_key}:complete:plan").click().run()
    assert app.session_state[f"{state_key}:step"] == "write" and "plan" in app.session_state[state_key].completed

    # The confirmation opens a dialog, which the test harness cannot drive; the trigger and the action are checked apart.
    app.selectbox(key="design_choice:demo").select("camt-053").run()
    app.button(key="design:demo:camt-053:delete").click().run()
    assert not app.exception
    app.sidebar.radio[0].set_value("Templates").run()
    app.button(key="template_remove").click().run()
    assert not app.exception
    assert "Remove a design" in [e.label for e in app.expander]
    assert app.selectbox(key="design_remove_choice").options == ["camt-053 (demo)"]
    app.button(key="design_remove").click().run()
    assert not app.exception

    ui = _app_module()
    blueprint = Registry(registry_with_demo).load("demo").blueprint
    keys = [s.key for s in blueprint.sections]
    board_design = ui.Design(name="x", template="demo")
    assert ui._section_order(board_design, blueprint) == keys
    ui._move_section(board_design, blueprint, keys[-1], -1)
    assert board_design.order == [keys[-1]] + keys[:-1]
    ui._move_section(board_design, blueprint, keys[-1], -1)
    assert board_design.order[0] == keys[-1]

    merged = ui._merge_draft({"a": "edited", "b": "old", "c": ""}, {"a": "old", "b": "old"}, {"a": "new", "b": "new", "c": "new", "d": "new"})
    assert merged == {"a": "edited", "b": "new", "c": "new", "d": "new"}

    fresh = ui.Design(name="y", template="demo")
    entries = [{"section": keys[0], "title": "A", "png": None, "text": ""}]
    assert [e["section"] for e in ui._visible_entries(entries, fresh, blueprint)] == keys
    fresh.hidden = [keys[1]]
    assert [e["section"] for e in ui._visible_entries(entries, fresh, blueprint)] == [keys[0]]
    twice = entries + [dict(entries[0])]
    assert [e["title"] for e in ui._visible_entries(twice, fresh, blueprint)] == ["A", "A (cont.)"]
    with_detail = [entries[0], {**entries[0], "title": "Need", "detail": "need_details"}, dict(entries[0])]
    assert [e["title"] for e in ui._visible_entries(with_detail, fresh, blueprint)] == ["A", "Need", "A (cont.)"]
    from types import SimpleNamespace

    from sdgen.render import RenderIssue

    response = SimpleNamespace(slide_map=[1, 2], slide_keys=["", ""], issues=[RenderIssue(level="warning", slide=2, message="check overflow: 'Box' needs about 130 percent of its box height"), RenderIssue(level="info", slide=1, message="check thin_table: x")])
    board = ui._board_entries(Registry(registry_with_demo).load("demo"), fresh, response, {}, {}, {})
    assert board[0]["checks"] == [] and board[1]["checks"] == ["overflow: 'Box' needs about 130 percent of its box height"]
    assert ui._with_overflow_note(board, ["Slide pictures are not available: x."])[-1].startswith("Check pass: 1 finding(s) on 1 slide(s)")
    assert ui.SHORTCUTS == {"previous": "Left", "next": "Right", "up": "Up", "down": "Down", "hide": "Delete"}
    assert "max-height" in ui.VIEWER_CSS and "100vw" in ui.FULL_VIEW_CSS and "viewer_details" in ui.FULL_VIEW_CSS and ui._write_title(fresh) == "3. Write the slides"
    assert f"{state_key}:draw_after_write" not in app.session_state
    store = DesignStore(tmp_path / "designs")
    loaded = Registry(registry_with_demo).load("demo")
    slot = next((s for s in blueprint.sections if any(loaded.manifest.field(k) and loaded.manifest.field(k).kind == "image" for k in s.fields)), None)
    if slot is not None:
        planned = ui.Design(name="p", template="demo", plan=ui.SectionPlan(flows=[ui.FlowRequest(section=slot.key, title="Flow", purpose="show it")]))
        assert [s.key for s in ui._pending_flow_sections(planned, loaded, store)] == [slot.key]
        planned.hidden = [slot.key]
        assert ui._pending_flow_sections(planned, loaded, store) == []
    assert ui._pending_flow_sections(fresh, loaded, store) == []
    assert ui.WALKTHROUGH_SUFFIX == "_walkthrough" and callable(ui.walkthrough_text)
    sizes = ui.slot_sizes(loaded)
    assert set(sizes) == {f.key for f in loaded.manifest.fields if f.kind == "image"}
    if slot is not None:
        sized = ui._sized_requests(planned, [slot], [ui.FlowRequest(section=slot.key, title="Flow", purpose="")])
        assert sized[0].width_in > 0 and sized[0].height_in > 0
    assert list(ui.DIAGRAM_FORMATS) == ["shapes", "drawio", "mermaid"]
    assert ui._flow_icons(ui.Design(name="s", template="demo", brief=ui.Brief(subject="x", about="SAP S/4HANA lockbox"))) == ui.icon_keys()
    assert ui._flow_icons(fresh) is None

    entry = Registry(registry_with_demo).load("demo")
    entry.manifest.field("scope").static = True
    kept_design = ui.Design(name="z", template="demo", modes={keys[0]: "keep"})
    values, modes = ui._render_fields(kept_design, entry)
    assert modes == {"scope": "keep"} and "scope" not in values and values["business_need"].startswith("Something long")
    values, modes = ui._render_fields(ui.Design(name="z", template="demo", modes={keys[0]: "blank"}), entry)
    assert modes["scope"] == "blank" and modes["business_need"] == "blank" and "scope" not in values
    entry.manifest.field("business_need").bindings[0].mode = "token"
    written_md = ui.dump_markdown(ui.Content(fields={"business_need": "Written insight"}), entry.manifest)
    values, _ = ui._render_fields(ui.Design(name="w", template="demo", modes={keys[0]: "keep"}, content_markdown=written_md), entry)
    assert values["business_need"] == "Written insight"
    assert ui._skipped_sections(ui.Design(name="s", template="demo", hidden=["a"], modes={"b": "blank", "c": "keep"})) == {"a", "b"}
    assert ui._kept_sections(ui.Design(name="s", template="demo", modes={"b": "blank", "c": "keep"})) == {"c"}

    store.save(ui.Design(name="second", template="demo"))
    ui._delete_design(store, "demo", "camt-053")
    assert not (tmp_path / "designs" / "camt-053").exists() and store.names("demo") == ["second"]
    ui._remove_template(Registry(registry_with_demo), store, "demo")
    assert not (registry_with_demo / "demo").exists() and store.names() == [] and not (tmp_path / "designs" / "second").exists()


def test_reference_material_can_be_added_and_removed(registry_with_demo, tmp_path, monkeypatch):
    import io

    from PIL import Image

    import sdgen.material as material

    monkeypatch.setattr(material, "fetch_link", lambda url, **_: ("ISO 20022", "Page text 12345", "fetched 2026-09-10"))
    app = AppTest.from_file(str(APP), default_timeout=60).run()
    app.text_input(key="design_name:demo").input("Lockbox").run()
    state_key = "design:demo:lockbox"
    assert any("Reference material" in m.value for m in app.markdown)
    app.text_area(key=f"{state_key}:0:mat:paste").input("reception, emission, tmp")
    app.text_input(key=f"{state_key}:0:mat:title").input("Folder layout")
    app.run()
    app.button(key=f"{state_key}:mat:add").click().run()
    assert not app.exception
    items = app.session_state[state_key].brief.material
    assert [(m.kind, m.title, m.text, m.status) for m in items] == [("text", "Folder layout", "reception, emission, tmp", "pasted")]
    app.text_input(key=f"{state_key}:1:mat:link").input("https://example.org/iso").run()
    app.button(key=f"{state_key}:mat:add").click().run()
    assert not app.exception
    items = app.session_state[state_key].brief.material
    assert (items[1].kind, items[1].title, items[1].text, items[1].status, items[1].url) == ("link", "ISO 20022", "Page text 12345", "fetched 2026-09-10", "https://example.org/iso")
    buffer = io.BytesIO()
    Image.new("RGB", (4, 4), "blue").save(buffer, format="PNG")
    app.file_uploader(key=f"{state_key}:2:mat:files").set_value([("sketch.png", buffer.getvalue(), "image/png")]).run()
    app.button(key=f"{state_key}:mat:add").click().run()
    assert not app.exception
    items = app.session_state[state_key].brief.material
    assert items[2].kind == "image" and items[2].file.endswith("-sketch.png") and items[2].status.startswith("not transcribed (mock provider)")
    folder = tmp_path / "designs" / "lockbox"
    material_yaml = (folder / "material.yaml").read_text(encoding="utf-8")
    assert "Folder layout" in material_yaml and "ISO 20022" in material_yaml and (folder / "material" / items[2].file).is_file()
    app.text_area(key=f"{state_key}:3:mat:text:{items[2].id}").input("Bank -> SFTP -> S/4").run()
    assert app.session_state[state_key].brief.material[2].text == "Bank -> SFTP -> S/4"
    app.button(key=f"{state_key}:mat:rm:{items[2].id}").click().run()
    assert not app.exception and not (folder / "material" / items[2].file).exists() and len(app.session_state[state_key].brief.material) == 2
    assert "sketch" not in (folder / "material.yaml").read_text(encoding="utf-8")


def test_design_page_picks_the_reasoning_deployment(registry_with_demo, monkeypatch):
    monkeypatch.setenv("SDGEN_LLM", "azure")
    monkeypatch.setenv("AZURE_OPENAI_DEPLOYMENT", "gpt-4.1,gpt-5,gpt-4o-mini")
    app = AppTest.from_file(str(APP), default_timeout=60).run()
    assert not app.exception and app.selectbox(key="llm_provider").value == "azure"
    assert app.session_state["llm"][1]["model"] == "gpt-5"
    app.text_input(key="design_name:demo").input("Lockbox").run()
    assert not app.exception
    picker = app.selectbox(key="design:demo:lockbox:deployment")
    assert picker.value == "gpt-5" and picker.options == ["gpt-4.1", "gpt-5", "gpt-4o-mini"]
    picker.select("gpt-4.1").run()
    assert not app.exception and app.selectbox(key="design:demo:lockbox:deployment").value == "gpt-4.1"


def test_icon_note_tells_what_the_drawings_show(tmp_path, monkeypatch):
    ui = _app_module()
    plain = ui.Design(name="n", template="demo")
    assert ui._icon_note(plain).startswith("Nodes draw as plain shapes")
    sap = ui.Design(name="s", template="demo", brief=ui.Brief(subject="x", about="SAP S/4HANA lockbox"))
    monkeypatch.setenv("SDGEN_ICONS", str(tmp_path))
    assert ui._icon_note(sap).startswith("No SAP icon")
    svg = tmp_path / catalogue()["cloud_integration"].file
    svg.write_text("<svg/>", encoding="utf-8")
    (tmp_path / "png").mkdir()
    (tmp_path / "png" / (svg.stem + ".png")).write_bytes(b"png")
    assert ui._icon_note(sap).startswith("Icons installed for 1 of")
    assert list(ui.DIAGRAM_FORMATS) == ["shapes", "drawio", "mermaid"]


def _labels(app):
    # The test tree lists an expander with an icon as a status box.
    return [e.label for e in list(app.expander) + list(app.status)]


def _steps(app):
    return sorted(label.split(":")[0] for label in _labels(app) if label[:1].isdigit())


def _app_module():
    spec = importlib.util.spec_from_file_location("sdgen_ui_app", APP)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_mappings_page_shows_grid_and_writes_workbook(registry_with_demo, tmp_path):
    from sdgen.design import Design, DesignStore
    from sdgen.mapping.model import FieldInfo, MappingEntry, MappingSet, SourceSpec, TargetSpec

    store = DesignStore(tmp_path / "designs")
    mapping = MappingSet(
        name="camt",
        target=TargetSpec(name="API", kind="edmx", fields=[FieldInfo(path="Stmt/Id", required=True), FieldInfo(path="Stmt/Amount")]),
        sources=[SourceSpec(name="Bank A", file="a.xml", kind="xml", fields=[FieldInfo(path="Doc/Id"), FieldInfo(path="Doc/Amt")])],
        entries=[MappingEntry(target_path="Stmt/Id", source="Bank A", source_path="Doc/Id")],
    )
    store.save(Design(name="camt", template="demo", mapping=mapping))

    app = AppTest.from_file(str(APP), default_timeout=60).run()
    app.sidebar.radio[0].set_value("Mappings").run()
    assert not app.exception
    assert app.selectbox(key="mapping_design:demo").value == "camt"
    assert any("Source Bank A: 1 of 2 target fields mapped" in t.value for t in app.text)

    app.button(key="design:demo:camt:summary_to_brief").click().run()
    assert not app.exception
    design = app.session_state["design:demo:camt"]
    assert "Detailed field mapping: camt-mapping.xlsx" in design.brief.mapping_summary
    assert (tmp_path / "designs" / "camt" / "mapping" / "camt-mapping.xlsx").is_file()
    assert store.load("camt").brief.mapping_summary == design.brief.mapping_summary
