from pathlib import Path

from click.testing import CliRunner
from openpyxl import load_workbook

from sdgen.cli import main
from sdgen.mapping.extract import detect_kind, extract_fields
from sdgen.mapping.model import FieldInfo, MappingEntry, MappingSet, SourceSpec, TargetSpec
from sdgen.mapping.workbook import write_workbook

CAMT = """<?xml version="1.0" encoding="UTF-8"?>
<Document xmlns="urn:iso:std:iso:20022:tech:xsd:camt.053.001.02">
  <BkToCstmrStmt>
    <GrpHdr><MsgId>STMT-2026-09-01</MsgId><CreDtTm>2026-09-01T06:00:00</CreDtTm></GrpHdr>
    <Stmt>
      <Id>0001</Id>
      <Acct><Id><IBAN>ZA1234567890</IBAN></Id><Ccy>ZAR</Ccy></Acct>
      <Ntry>
        <Amt Ccy="ZAR">1500.00</Amt><CdtDbtInd>CRDT</CdtDbtInd>
        <NtryDtls><TxDtls><Refs><EndToEndId>INV-1</EndToEndId></Refs></TxDtls></NtryDtls>
      </Ntry>
      <Ntry>
        <Amt Ccy="ZAR">250.50</Amt><CdtDbtInd>DBIT</CdtDbtInd>
        <NtryDtls><TxDtls><Refs><EndToEndId>INV-2</EndToEndId></Refs></TxDtls></NtryDtls>
      </Ntry>
    </Stmt>
  </BkToCstmrStmt>
</Document>
"""

XSD = """<?xml version="1.0"?>
<xs:schema xmlns:xs="http://www.w3.org/2001/XMLSchema">
  <xs:complexType name="ItemType">
    <xs:sequence>
      <xs:element name="Amount" type="xs:decimal"/>
      <xs:element name="Reference" type="xs:string" minOccurs="0"/>
    </xs:sequence>
    <xs:attribute name="Currency" type="xs:string" use="required"/>
  </xs:complexType>
  <xs:element name="BankStatement">
    <xs:complexType>
      <xs:sequence>
        <xs:element name="StatementId" type="xs:string"/>
        <xs:element name="Item" type="ItemType" maxOccurs="unbounded"/>
        <xs:element name="Remark" minOccurs="0">
          <xs:simpleType><xs:restriction base="xs:string"/></xs:simpleType>
        </xs:element>
      </xs:sequence>
    </xs:complexType>
  </xs:element>
</xs:schema>
"""

EDMX = """<?xml version="1.0" encoding="utf-8"?>
<edmx:Edmx Version="4.0" xmlns:edmx="http://docs.oasis-open.org/odata/ns/edmx">
  <edmx:DataServices>
    <Schema Namespace="API_BANKSTATEMENT" xmlns="http://docs.oasis-open.org/odata/ns/edm" xmlns:sap="http://www.sap.com/Protocols/SAPData">
      <EntityType Name="BankStatement">
        <Key><PropertyRef Name="StatementId"/></Key>
        <Property Name="StatementId" Type="Edm.String" Nullable="false" sap:label="Statement ID"/>
        <Property Name="HouseBank" Type="Edm.String" Nullable="false"/>
        <Property Name="StatementDate" Type="Edm.Date"/>
        <NavigationProperty Name="to_Item" Type="Collection(API_BANKSTATEMENT.Item)"/>
      </EntityType>
      <EntityType Name="Item">
        <Property Name="AmountInTransactionCurrency" Type="Edm.Decimal" Nullable="false"/>
        <Property Name="PaymentReference" Type="Edm.String"/>
      </EntityType>
    </Schema>
  </edmx:DataServices>
</edmx:Edmx>
"""


def _write(tmp_path, name, text):
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


def test_xml_sample_fields(tmp_path):
    kind, fields = extract_fields(_write(tmp_path, "camt.xml", CAMT))
    assert kind == "xml"
    by_path = {f.path: f for f in fields}
    amount = by_path["Document/BkToCstmrStmt/Stmt/Ntry/Amt"]
    assert amount.example == "1500.00" and amount.type == "decimal" and amount.occurs == 2 and amount.repeating
    assert by_path["Document/BkToCstmrStmt/Stmt/Ntry/Amt/@Ccy"].example == "ZAR"
    assert by_path["Document/BkToCstmrStmt/GrpHdr/CreDtTm"].type == "date"
    assert list(by_path)[0] == "Document/BkToCstmrStmt/GrpHdr/MsgId"


def test_json_and_csv_samples(tmp_path):
    kind, fields = extract_fields(_write(tmp_path, "s.json", '{"statement": {"id": "1", "items": [{"amount": 1.5, "ref": "A"}, {"amount": 2, "ref": null}]}}'))
    assert kind == "json"
    by_path = {f.path: f for f in fields}
    assert by_path["statement/items[]/amount"].occurs == 2 and by_path["statement/items[]/amount"].repeating
    assert by_path["statement/items[]/amount"].type == "decimal"
    assert by_path["statement/items[]/ref"].example == "A"

    kind, fields = extract_fields(_write(tmp_path, "s.csv", "Account,Amount,Date\nZA1,10.5,2026-09-01\nZA2,,2026-09-02\n"))
    assert kind == "csv"
    assert [(f.path, f.type, f.example) for f in fields] == [("Account", "string", "ZA1"), ("Amount", "decimal", "10.5"), ("Date", "date", "2026-09-01")]


def test_xsd_schema_fields(tmp_path):
    path = _write(tmp_path, "schema.xml", XSD)
    assert detect_kind(path) == "xsd"
    kind, fields = extract_fields(path)
    by_path = {f.path: f for f in fields}
    assert by_path["BankStatement/StatementId"].required is True and by_path["BankStatement/StatementId"].type == "string"
    assert "BankStatement/Item" not in by_path  # container elements are walked, not listed
    assert by_path["BankStatement/Item/Amount"].type == "decimal" and by_path["BankStatement/Item/Amount"].required
    assert by_path["BankStatement/Item/Reference"].required is False
    assert by_path["BankStatement/Item/@Currency"].required is True
    assert by_path["BankStatement/Remark"].type == "string" and by_path["BankStatement/Remark"].required is False


def test_edmx_fields(tmp_path):
    path = _write(tmp_path, "metadata.xml", EDMX)
    assert detect_kind(path) == "edmx"
    kind, fields = extract_fields(path)
    by_path = {f.path: f for f in fields}
    assert by_path["BankStatement/StatementId"].required and by_path["BankStatement/StatementId"].description == "Statement ID"
    assert by_path["BankStatement/StatementDate"].required is False and by_path["BankStatement/StatementDate"].type == "Edm.Date"
    assert by_path["BankStatement/to_Item"].type == "navigation"
    assert by_path["Item/AmountInTransactionCurrency"].required


def _mapping_set() -> MappingSet:
    target = TargetSpec(
        name="BankStatement API",
        file="metadata.xml",
        fields=[
            FieldInfo(path="BankStatement/StatementId", type="Edm.String", required=True),
            FieldInfo(path="BankStatement/HouseBank", type="Edm.String", required=True),
            FieldInfo(path="Item/AmountInTransactionCurrency", type="Edm.Decimal", required=True),
            FieldInfo(path="Item/PaymentReference", type="Edm.String", required=False),
        ],
    )
    sources = [
        SourceSpec(name="Bank A", file="a.xml", fields=[FieldInfo(path="Document/Stmt/Id"), FieldInfo(path="Document/Stmt/Ntry/Amt")]),
        SourceSpec(name="Bank B", file="b.xml", fields=[FieldInfo(path="Statement/Id")]),
    ]
    mapping = MappingSet(name="camt053", target=target, sources=sources)
    mapping.set_entry(MappingEntry(target_path="BankStatement/StatementId", source="Bank A", source_path="Document/Stmt/Id", example="0001"))
    mapping.set_entry(MappingEntry(target_path="Item/AmountInTransactionCurrency", source="Bank A", source_path="Document/Stmt/Ntry/Amt", rule="decimal, 2 places"))
    mapping.set_entry(MappingEntry(target_path="BankStatement/HouseBank", source="Bank A", rule="constant 'HB01'"))
    mapping.set_entry(MappingEntry(target_path="BankStatement/StatementId", source="Bank B", source_path="Statement/Id"))
    return mapping


def test_mapping_set_summary_and_roundtrip(tmp_path):
    mapping = _mapping_set()
    summary = {s.source: s for s in mapping.summary()}
    assert (summary["Bank A"].mapped, summary["Bank A"].total, summary["Bank A"].unmapped_required) == (3, 4, [])
    assert summary["Bank B"].mapped == 1 and summary["Bank B"].unmapped_required == ["BankStatement/HouseBank", "Item/AmountInTransactionCurrency"]
    assert "Source Bank B: 1 of 4 target fields mapped; required fields still unmapped: BankStatement/HouseBank" in mapping.summary_text()

    mapping.set_entry(MappingEntry(target_path="BankStatement/StatementId", source="Bank B", source_path=""))
    assert "BankStatement/StatementId" not in mapping.entries_for("Bank B")

    mapping.save(tmp_path / "m.yaml")
    assert MappingSet.load(tmp_path / "m.yaml") == mapping


def test_workbook_has_summary_and_one_sheet_per_source(tmp_path):
    mapping = _mapping_set()
    path = write_workbook(mapping, tmp_path / "mapping.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Summary", "Bank A", "Bank B"]
    summary = workbook["Summary"]
    assert summary["A1"].value == "Mapping" and summary["B1"].value == "camt053"
    rows = [[c.value for c in row] for row in summary.iter_rows(min_row=6, max_row=7)]
    assert rows[0][:4] == ["Bank A", "a.xml", 3, 4] and rows[1][4].startswith("BankStatement/HouseBank")

    sheet = workbook["Bank A"]
    assert [c.value for c in sheet[1]] == ["Target field", "Type", "Required", "Source field", "Transformation rule", "Example", "Notes"]
    data = {row[0].value: [c.value for c in row] for row in sheet.iter_rows(min_row=2)}
    assert data["BankStatement/StatementId"][2:6] == ["yes", "Document/Stmt/Id", None, "0001"]
    assert data["BankStatement/HouseBank"][4] == "constant 'HB01'"
    assert data["Item/PaymentReference"][2] == "no" and data["Item/PaymentReference"][3] is None
    assert sheet.freeze_panes == "A2"


def test_cli_mapping_commands(tmp_path):
    target = _write(tmp_path, "metadata.xml", EDMX)
    bank_a = _write(tmp_path, "bank-a.xml", CAMT)
    runner = CliRunner()
    listed = runner.invoke(main, ["mapping", "extract", str(bank_a)])
    assert listed.exit_code == 0 and "Document/BkToCstmrStmt/Stmt/Ntry/Amt  [decimal x2 repeating]" in listed.output

    created = runner.invoke(main, ["mapping", "new", "camt053", "--target", str(target), "--source", f"Bank A={bank_a}", "-o", str(tmp_path / "m.yaml")])
    assert created.exit_code == 0, created.output
    mapping = MappingSet.load(tmp_path / "m.yaml")
    assert mapping.target.kind == "edmx" and mapping.sources[0].name == "Bank A" and mapping.sources[0].kind == "xml"

    built = runner.invoke(main, ["mapping", "workbook", str(tmp_path / "m.yaml"), "-o", str(tmp_path / "m.xlsx")])
    assert built.exit_code == 0, built.output
    assert load_workbook(tmp_path / "m.xlsx").sheetnames == ["Summary", "Bank A"]
