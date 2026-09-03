import copy

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from sdgen.fill.text import (
    Block,
    Span,
    parse_blocks,
    replace_literal_everywhere,
    replace_token,
    set_rich_text,
)


def _slide(prs=None):
    prs = prs or Presentation()
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def _textbox(slide, text="template"):
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
    box.text_frame.text = text
    return box


def _paragraph_xml(paragraph):
    return paragraph._p


def _bullet_char(paragraph):
    ppr = _paragraph_xml(paragraph).find(qn("a:pPr"))
    bu = ppr.find(qn("a:buChar")) if ppr is not None else None
    return bu.get("char") if bu is not None else None


def _has_bu_none(paragraph):
    ppr = _paragraph_xml(paragraph).find(qn("a:pPr"))
    return ppr is not None and ppr.find(qn("a:buNone")) is not None


def test_parse_blocks_bullets_levels_and_inline_marks():
    blocks = parse_blocks("Intro **bold** and *it*\n\n- one\n  - two\n\t- tab\n1. numbered stays\nZZ1_C_FINAC_AIF")
    assert [(b.bullet, b.level, b.text) for b in blocks] == [
        (False, 0, "Intro bold and it"),
        (True, 0, "one"),
        (True, 1, "two"),
        (True, 1, "tab"),
        (False, 0, "1. numbered stays"),
        (False, 0, "ZZ1_C_FINAC_AIF"),
    ]
    assert [s.bold for s in blocks[0].spans] == [False, True, False, False]
    assert blocks[0].spans[3].italic is True


def test_paragraphs_keep_template_run_formatting():
    _, slide = _slide()
    box = _textbox(slide)
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x11, 0x22, 0x33)

    set_rich_text(box, "First line\nSecond with **bold** word\nThird")
    paragraphs = box.text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["First line", "Second with bold word", "Third"]
    assert all(r.font.size == Pt(11) for p in paragraphs for r in p.runs)
    assert all(str(r.font.color.rgb) == "112233" for p in paragraphs for r in p.runs)
    bold_flags = [r.font.bold for r in paragraphs[1].runs]
    assert bold_flags == [None, True, None]


def test_keep_prefix_preserves_label_run_and_uses_content_run_style():
    _, slide = _slide()
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
    paragraph = box.text_frame.paragraphs[0]
    label = paragraph.add_run()
    label.text = "Business Need: "
    label.font.bold = True
    content = paragraph.add_run()
    content.text = "old content"
    content.font.italic = True
    box.text_frame.add_paragraph().text = "old second"

    set_rich_text(box, "new content\nnew second", keep_prefix="Business Need: ")
    paragraphs = box.text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["Business Need: new content", "new second"]
    runs = paragraphs[0].runs
    assert (runs[0].text, runs[0].font.bold) == ("Business Need: ", True)
    assert (runs[1].text, runs[1].font.bold, runs[1].font.italic) == ("new content", None, True)
    assert paragraphs[1].runs[0].font.italic is None


def test_keep_prefix_inside_a_single_run_is_split():
    _, slide = _slide()
    box = _textbox(slide, "Scope: everything old")
    set_rich_text(box, "only new", keep_prefix="Scope: ")
    runs = box.text_frame.paragraphs[0].runs
    assert [r.text for r in runs] == ["Scope: ", "only new"]


def test_missing_prefix_is_prepended():
    _, slide = _slide()
    box = _textbox(slide, "no label here")
    set_rich_text(box, "value", keep_prefix="Label: ")
    assert box.text_frame.text == "Label: value"


def test_bullets_in_plain_text_box_get_explicit_bullets():
    _, slide = _slide()
    box = _textbox(slide)
    set_rich_text(box, "Heading\n- first\n  - nested\nTrailer")
    paragraphs = box.text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["Heading", "first", "nested", "Trailer"]
    assert _bullet_char(paragraphs[0]) is None
    assert _bullet_char(paragraphs[1]) == "•" and paragraphs[1].level == 0
    assert _bullet_char(paragraphs[2]) == "•" and paragraphs[2].level == 1
    assert _bullet_char(paragraphs[3]) is None and not _has_bu_none(paragraphs[3])


def test_bullets_reuse_template_bullet_paragraph():
    _, slide = _slide()
    box = _textbox(slide, "Heading")
    bullet_p = box.text_frame.add_paragraph()
    bullet_p.text = "template bullet"
    ppr = bullet_p._p.get_or_add_pPr()
    ppr.set("marL", "457200")
    ppr.set("indent", "-457200")
    bu = etree.SubElement(ppr, qn("a:buChar"))
    bu.set("char", "-")

    set_rich_text(box, "Title\n- a\n  - b\nPlain again")
    paragraphs = box.text_frame.paragraphs
    assert _bullet_char(paragraphs[0]) is None
    assert _bullet_char(paragraphs[1]) == "-" and paragraphs[1]._p.pPr.get("marL") == "457200"
    assert _bullet_char(paragraphs[2]) == "-" and paragraphs[2].level == 1
    assert _bullet_char(paragraphs[3]) is None


def test_body_placeholder_only_sets_levels():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    body = slide.placeholders[1]
    set_rich_text(body, "- top\n  - nested\nplain")
    paragraphs = body.text_frame.paragraphs
    assert [p.level for p in paragraphs] == [0, 1, 0]
    assert all(_bullet_char(p) is None for p in paragraphs)
    assert not _has_bu_none(paragraphs[2])


def test_empty_value_leaves_one_empty_paragraph():
    _, slide = _slide()
    box = _textbox(slide, "old")
    set_rich_text(box, "")
    assert len(box.text_frame.paragraphs) == 1 and box.text_frame.text == ""


def test_clearing_then_refilling_keeps_run_formatting():
    _, slide = _slide()
    frame = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    cell = frame.table.cell(1, 1)
    cell.text = "styled"
    cell.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    cell.text_frame.paragraphs[0].runs[0].font.bold = True

    set_rich_text(cell, "")
    assert cell.text_frame.text == "" and not cell.text_frame.paragraphs[0].runs
    set_rich_text(cell, "refilled")
    run = cell.text_frame.paragraphs[0].runs[0]
    assert (run.text, run.font.size, run.font.bold) == ("refilled", Pt(9), True)


def test_replace_token_within_and_across_runs():
    _, slide = _slide()
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    paragraph = box.text_frame.paragraphs[0]
    first = paragraph.add_run()
    first.text = "Initiative: <Internal"
    first.font.bold = True
    second = paragraph.add_run()
    second.text = " effort> and <external effort> twice <external effort>"

    assert replace_token(box, "<Internal effort>", "40h") == 1
    assert replace_token(box, "<external effort>", "10h") == 2
    assert paragraph.text == "Initiative: 40h and 10h twice 10h"
    assert paragraph.runs[0].font.bold is True
    assert replace_token(box, "<missing>", "x") == 0


def test_replace_literal_everywhere_covers_titles_tables_and_groups(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Overview: Old Name"
    _textbox(slide, "About Old Name and Old Name")
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(4), Inches(3), Inches(1)).table
    table.cell(0, 0).text = "Old Name row"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(5), Inches(5), Inches(2), Inches(1)).text_frame.text = "Old Name nested"

    assert replace_literal_everywhere(prs, "Old Name", "New Name") == 5
    assert slide.shapes.title.text == "Overview: New Name"
    assert table.cell(0, 0).text == "New Name row"
    texts = [s.text_frame.text for s in group.shapes]
    assert texts == ["New Name nested"]

    out = tmp_path / "out.pptx"
    prs.save(out)
    assert Presentation(str(out)).slides[0].shapes.title.text == "Overview: New Name"


def test_block_objects_are_accepted():
    _, slide = _slide()
    box = _textbox(slide)
    blocks = [Block(spans=[Span(text="Hello "), Span(text="there", bold=True)])]
    original = copy.deepcopy(blocks)
    set_rich_text(box, blocks)
    assert box.text_frame.text == "Hello there"
    assert blocks == original
