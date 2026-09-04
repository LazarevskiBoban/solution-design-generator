from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches

from sdgen.analyze import analyze_deck
from sdgen.blueprint import derive_blueprint
from sdgen.content import Content, ImageValue, load_markdown
from sdgen.diagrams import SLOT_PREFIX, add_image_slots
from sdgen.inventory import inspect_deck
from sdgen.registry import Registry
from sdgen.render import render

CARRIER_DECK = Path("D:/NTT-architectures/NTT DATA Inc I IT I Solution Design I Carrier AP Invoice Integration.pptx")


@pytest.fixture
def diagram_deck(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Integration Architecture: Demo Flow"
    for i in range(6):
        top = Inches(2) if i % 2 == 0 else Inches(4)
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1 + i * 1.5), top, Inches(1.2), Inches(0.8))
        box.text_frame.text = f"System {i}"
    for i in range(5):
        slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(2.2 + i * 1.5), Inches(2.4), Inches(2.5 + i * 1.5), Inches(4.4))
    frame = slide.shapes.add_table(3, 2, Inches(9), Inches(4.5), Inches(3.5), Inches(1.2))
    frame.name = "Steps Table"
    frame.table.cell(0, 0).text = "Step"
    frame.table.cell(0, 1).text = "Activity"
    frame.table.cell(1, 0).text = "1"
    frame.table.cell(1, 1).text = "Send file to middleware for validation and posting"
    logo = tmp_path / "logo.png"
    Image.new("RGB", (30, 30), "black").save(logo)
    slide.shapes.add_picture(str(logo), Inches(12), Inches(0.2), Inches(1))
    path = tmp_path / "diagram.pptx"
    prs.save(path)
    return path


def _prepare(path):
    deck = inspect_deck(path)
    analysis = analyze_deck(deck)
    manifest = analysis.to_manifest("demo")
    blueprint = derive_blueprint(deck, analysis, manifest, "demo")
    return deck, analysis, manifest, blueprint


def test_slot_replaces_drawing_but_keeps_title_and_bound_table(diagram_deck, tmp_path):
    deck, analysis, manifest, blueprint = _prepare(diagram_deck)
    assert blueprint.sections[0].kind == "diagram"
    prs = Presentation(str(diagram_deck))
    manifest, blueprint = add_image_slots(prs, manifest, blueprint)

    slide = prs.slides[0]
    names = [s.name for s in slide.shapes]
    assert names[0].startswith("Title") and "Steps Table" in names
    slots = [s for s in slide.shapes if s.name.startswith(SLOT_PREFIX)]
    assert len(slots) == 1 and len(slide.shapes) == 3
    slot = slots[0]
    assert slot.left == Inches(1) and slot.top == Inches(2)
    assert slot.width == pytest.approx(Inches(8.7), rel=0.01)
    assert slot.height == pytest.approx(Inches(2.8), rel=0.01)

    field = next(f for f in manifest.fields if f.kind == "image")
    assert field.bindings[0].shape.id == slot.shape_id and field.bindings[0].slide == 1
    assert field.key in blueprint.sections[0].fields
    assert len(slide.part.rels) == 2  # layout + slot image; the logo relationship is gone
    out = tmp_path / "slots.pptx"
    prs.save(out)
    Presentation(str(out))


def test_render_places_one_or_several_images(diagram_deck, tmp_path):
    deck, analysis, manifest, blueprint = _prepare(diagram_deck)
    registry = Registry(tmp_path / "templates")
    entry = registry.add("demo", diagram_deck, manifest, blueprint=blueprint)
    field = next(f for f in entry.manifest.fields if f.kind == "image")

    red, blue = tmp_path / "red.png", tmp_path / "blue.png"
    Image.new("RGB", (800, 200), "red").save(red)
    Image.new("RGB", (200, 800), "blue").save(blue)

    single = render(entry.template_path, entry.manifest, Content(fields={field.key: ImageValue(path=str(red))}), tmp_path / "one.pptx")
    assert not single.errors and single.slides == 1

    content = load_markdown(f"## {field.key}\n![]({red})\n![]({blue})\n", entry.manifest)
    result = render(entry.template_path, entry.manifest, content, tmp_path / "two.pptx")
    assert not result.errors and result.slides == 2
    prs = Presentation(str(tmp_path / "two.pptx"))
    sizes = [next(s for s in slide.shapes if hasattr(s, "image")).image.size for slide in prs.slides]
    assert sizes == [(800, 200), (200, 800)]
    assert prs.slides[1].shapes.title.text.endswith("(cont.)")


@pytest.mark.skipif(not CARRIER_DECK.exists(), reason="example deck not available")
def test_carrier_diagram_slides_get_slots(tmp_path):
    deck, analysis, manifest, blueprint = _prepare(CARRIER_DECK)
    prs = Presentation(str(CARRIER_DECK))
    manifest, blueprint = add_image_slots(prs, manifest, blueprint)
    out = tmp_path / "carrier-slots.pptx"
    prs.save(out)

    info = inspect_deck(out)
    diagram_slides = [s.slide for s in blueprint.sections if s.kind == "diagram"]
    assert 8 in diagram_slides and 16 in diagram_slides and 12 in diagram_slides
    for index in diagram_slides:
        slide = info.slides[index - 1]
        kinds = [s.kind for s in slide.shapes]
        assert kinds.count("picture") == 1 and slide.connector_count == 0
        assert any(s.name.startswith(SLOT_PREFIX) for s in slide.shapes)
    slide12 = info.slides[11]
    assert any(s.kind == "table" for s in slide12.shapes)
    assert sum(1 for f in manifest.fields if f.kind == "image") == len(diagram_slides)
