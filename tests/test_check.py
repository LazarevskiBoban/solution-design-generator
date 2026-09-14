from pptx import Presentation
from pptx.util import Inches, Pt

from sdgen.check import check_presentation
from sdgen.fill.text import effective_overflow_ratio, fit_text_shape
from sdgen.flow import FlowNode, FlowSpec, clean_flow, parse_systems


def _deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def _titled(prs, title: str):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    return slide


def _box(slide, name: str, left: float, top: float, width: float, height: float, text: str = "x"):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.name = name
    box.text_frame.text = text
    return box


def _codes(findings):
    return sorted(f.code for f in findings)


def test_flags_shapes_outside_the_slide_and_long_titles():
    prs = _deck()
    slide = _titled(prs, "A title that runs on far longer than sixty characters allow on one slide")
    _box(slide, "Left", -1.0, 2.0, 2.0, 1.0)
    _box(slide, "Right", 12.5, 2.0, 2.0, 1.0)
    _box(slide, "Inside", 1.0, 2.0, 2.0, 1.0)
    findings = check_presentation(prs)
    assert _codes(findings) == ["outside", "outside", "title_long"]
    assert {f.shape for f in findings if f.code == "outside"} == {"Left", "Right"} and findings[-1].slide == 1
    clean = _deck()
    fine = _titled(clean, "Executive Overview (cont.)")
    _box(fine, "Body", 1.0, 2.0, 2.0, 1.0)
    assert check_presentation(clean) == []


def test_flags_flow_shapes_over_the_title_and_overlapping_nodes():
    prs = _deck()
    slide = _titled(prs, "Flows")
    title = slide.shapes.title
    _box(slide, "Flow x canvas", title.left / 914400, title.top / 914400, 6.0, 5.0, "")
    _box(slide, "Flow x node a", 1.0, 3.0, 2.0, 0.7, "A")
    _box(slide, "Flow x node b", 2.0, 3.2, 2.0, 0.7, "B")
    _box(slide, "Flow x node c", 5.0, 3.0, 2.0, 0.7, "C")
    label = _box(slide, "Flow x edge 1 label", 5.5, 3.1, 0.6, 0.5, "sFTP")
    label.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    findings = check_presentation(prs)
    assert _codes(findings) == ["overlap", "overlap", "overlap"]
    messages = [f.message for f in findings]
    assert any("node a' overlaps 'Flow x node b'" in m for m in messages) and any("label 'sFTP' sits on" in m for m in messages) and "the drawing overlaps the title" in messages

    # A node whose label does not fit is measured like any other box.
    tight_deck = _deck()
    crammed = _titled(tight_deck, "Nodes")
    tight = _box(crammed, "Flow y node long", 1.0, 2.0, 0.8, 0.3, "Integration Support and Business")
    for paragraph in tight.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(14)
    assert [(f.code, f.shape) for f in check_presentation(tight_deck)] == [("overflow", "Flow y node long")]


def test_flags_empty_placeholders_and_empty_slides():
    prs = _deck()
    content = prs.slides.add_slide(prs.slide_layouts[1])
    content.shapes.title.text = "Empty body"
    only_title = _titled(prs, "Only a title")
    filled = prs.slides.add_slide(prs.slide_layouts[1])
    filled.shapes.title.text = "Filled"
    filled.placeholders[1].text_frame.text = "Some text"
    findings = check_presentation(prs)
    assert [(f.slide, f.code) for f in findings] == [(1, "empty_box"), (1, "empty_slide"), (2, "empty_slide")]
    assert only_title.shapes.title.text == "Only a title"


def test_flags_short_tables_but_not_version_control():
    prs = _deck()
    slide = _titled(prs, "Tables")
    thin = slide.shapes.add_table(2, 2, Inches(1), Inches(1.5), Inches(6), Inches(0.8))
    thin.name = "Thin"
    thin.table.cell(0, 0).text, thin.table.cell(1, 0).text = "Ref", "1"
    version = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(0.8))
    version.name = "Version"
    version.table.cell(0, 0).text, version.table.cell(1, 0).text = "Version", "1.0"
    full = slide.shapes.add_table(4, 2, Inches(1), Inches(4.5), Inches(6), Inches(1.6))
    full.name = "Full"
    for row in range(4):
        full.table.cell(row, 0).text = str(row)
    findings = check_presentation(prs)
    assert [(f.code, f.shape, f.level) for f in findings] == [("thin_table", "Thin", "info")]


def test_estimated_overflow_honours_the_stored_font_scale():
    prs = _deck()
    slide = _titled(prs, "Overflow")
    moderate = _box(slide, "Moderate", 1.0, 1.5, 6.0, 1.0, "\n".join(["one line of text"] * 6))
    huge = _box(slide, "Huge", 1.0, 3.0, 6.0, 1.0, "\n".join(["one line of text"] * 40))
    fits = _box(slide, "Fits", 1.0, 4.5, 6.0, 1.0, "one line")
    for box in (moderate, huge, fits):
        for paragraph in box.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(12)
    before = check_presentation(prs)
    assert {f.shape for f in before if f.code == "overflow"} == {"Moderate", "Huge"}
    fit_text_shape(moderate)
    fit_text_shape(huge)
    assert effective_overflow_ratio(moderate) < 1.1 < effective_overflow_ratio(huge)
    after = check_presentation(prs)
    assert {f.shape for f in after if f.code == "overflow"} == {"Huge"}


def test_lane_mismatches_become_findings():
    lanes = parse_systems("SAP BTP Integration Suite | middleware\nS/4HANA on RISE | target\nAzure VM | source")
    spec = clean_flow(FlowSpec(nodes=[FlowNode(id="a", label="S/4 to VM copy", lane="sap_btp_integration_suite")], systems=lanes), lanes)
    prs = _deck()
    slide = _titled(prs, "Flow")
    _box(slide, "Flow x node a", 1.0, 2.0, 2.0, 0.7, "S/4 to VM copy")
    findings = check_presentation(prs, flows={1: spec})
    assert [f.code for f in findings] == ["lane"] and "S/4 to VM copy" in findings[0].message
    assert check_presentation(prs, flows={2: spec}) == []


def test_baseline_hides_what_the_template_already_shows():
    from sdgen.check import baseline_findings

    prs = _deck()
    slide = _titled(prs, "Template")
    _box(slide, "Off", -1.0, 2.0, 2.0, 1.0)
    _box(slide, "Body", 1.0, 2.0, 2.0, 1.0)
    baseline = baseline_findings(prs)
    assert baseline == {1: {("outside", "Off")}}
    assert check_presentation(prs, baseline=baseline) == []
    _box(slide, "New", 13.0, 2.0, 2.0, 1.0)
    assert [f.shape for f in check_presentation(prs, baseline=baseline, origins=[1])] == ["New"]
    assert [f.shape for f in check_presentation(prs, baseline=baseline, origins=[2])] == ["Off", "New"]
