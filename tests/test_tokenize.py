from pptx import Presentation

from sdgen.analyze import analyze_deck
from sdgen.content import Content
from sdgen.inventory import inspect_deck
from sdgen.manifest import GlobalSpec
from sdgen.registry import Registry
from sdgen.render import render
from sdgen.tokenize import tokenize_deck


def _manifest(sample_deck):
    manifest = analyze_deck(inspect_deck(sample_deck)).to_manifest("demo")
    manifest.globals.append(GlobalSpec(key="subject", replaces="Demo Integration"))
    return manifest


def _shape_text(slide, name):
    return next(s for s in slide.shapes if s.name == name).text_frame.text


def test_tokenize_replaces_bound_content_with_markers(sample_deck):
    prs = Presentation(str(sample_deck))
    manifest = tokenize_deck(prs, _manifest(sample_deck))
    slide = prs.slides[0]

    assert slide.shapes.title.text == "Executive Overview: {{subject}}"
    assert manifest.globals[0].replaces == "{{subject}}"
    assert _shape_text(slide, "Business Need Box") == "Business Need: {{business_need}}"
    assert slide.placeholders[1].text_frame.text == "{{first_point}}"
    table = next(s for s in slide.shapes if s.has_table).table
    assert [[c.text for c in r.cells] for r in table.rows] == [["Function", "Countries"], ["{{scope}}", ""]]
    assert _shape_text(slide, "Scope Label") == "Scope"


def test_registry_stores_tokenized_template_that_still_renders(sample_deck, tmp_path):
    registry = Registry(tmp_path / "templates")
    entry = registry.add("demo", sample_deck, _manifest(sample_deck))
    stored = Presentation(str(entry.template_path)).slides[0]
    assert "{{business_need}}" in _shape_text(stored, "Business Need Box")
    assert registry.load("demo").manifest.globals[0].replaces == "{{subject}}"

    content = Content(
        globals={"subject": "Lockbox"},
        fields={"business_need": "Real text", "scope": [{"Function": "AR", "Countries": "KE"}], "first_point": "- one"},
    )
    result = render(entry.template_path, entry.manifest, content, tmp_path / "out.pptx")
    assert not result.errors
    out = Presentation(str(tmp_path / "out.pptx")).slides[0]
    assert out.shapes.title.text == "Executive Overview: Lockbox"
    assert _shape_text(out, "Business Need Box") == "Business Need: Real text"
    table = next(s for s in out.shapes if s.has_table).table
    assert [[c.text for c in r.cells] for r in table.rows] == [["Function", "Countries"], ["AR", "KE"]]
    assert "{{" not in "".join(s.text_frame.text for s in out.shapes if s.has_text_frame)


def test_registry_can_keep_the_original_deck(sample_deck, tmp_path):
    entry = Registry(tmp_path / "templates").add("raw", sample_deck, _manifest(sample_deck), tokenize=False)
    stored = Presentation(str(entry.template_path)).slides[0]
    assert "{{" not in _shape_text(stored, "Business Need Box")


def test_capture_content_reads_text_bullets_and_tables(sample_deck):
    from pptx import Presentation

    from sdgen.analyze import analyze_deck
    from sdgen.inventory import inspect_deck
    from sdgen.tokenize import capture_content

    manifest = analyze_deck(inspect_deck(sample_deck)).to_manifest("demo")
    original = capture_content(Presentation(str(sample_deck)), manifest)
    assert original.fields["business_need"] == "Something long enough to be treated as a real content section.\nSecond paragraph."
    assert original.fields["scope"] == [{"Function": "Finance", "Countries": "ZA"}]
    assert original.fields["first_point"] == "- First point\n  - Sub point"
