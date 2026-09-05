import os
from pathlib import Path

import pytest

from sdgen.preview import preview

POWERPOINT = Path("C:/Program Files/Microsoft Office/root/Office16/POWERPNT.EXE")


@pytest.mark.skipif(
    not POWERPOINT.exists() or os.environ.get("SDGEN_COM_TESTS") != "1",
    reason="set SDGEN_COM_TESTS=1 on a machine with PowerPoint",
)
def test_preview_opens_deck_and_exports_pdf(sample_deck, tmp_path):
    result = preview(sample_deck, tmp_path / "sample.pdf")
    assert result.opened, result.message
    assert result.slides == 2
    assert Path(result.pdf).is_file()


def test_preview_reports_unreadable_file(tmp_path):
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a deck")
    result = preview(bad)
    assert not result.opened
    assert result.message


@pytest.mark.skipif(
    not POWERPOINT.exists() or os.environ.get("SDGEN_COM_TESTS") != "1",
    reason="set SDGEN_COM_TESTS=1 on a machine with PowerPoint",
)
def test_export_slide_images_writes_one_png_per_slide(sample_deck, tmp_path):
    from sdgen.preview import export_slide_images

    files = export_slide_images(sample_deck, tmp_path / "png", width=640)
    assert [f.name for f in files] == ["slide-01.png", "slide-02.png"]
    assert all(f.stat().st_size > 0 for f in files)
    second = export_slide_images(sample_deck, tmp_path / "one", width=320, only=[2])
    assert [f.name for f in second] == ["slide-02.png"]


def test_export_slide_images_raises_on_unreadable_file(tmp_path):
    from sdgen.preview import export_slide_images

    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a deck")
    with pytest.raises(RuntimeError):
        export_slide_images(bad, tmp_path / "png", attempts=1)


def test_preview_rows_describe_each_field(sample_deck):
    from sdgen.analyze import analyze_deck
    from sdgen.blueprint import derive_blueprint
    from sdgen.content import Content
    from sdgen.inventory import inspect_deck
    from sdgen.preview import preview_rows

    deck = inspect_deck(sample_deck)
    analysis = analyze_deck(deck)
    manifest = analysis.to_manifest("demo")
    blueprint = derive_blueprint(deck, analysis, manifest, "demo")
    content = Content(fields={"business_need": "Fresh need text", "scope": [{"Function": "Finance", "Countries": "ZA"}]})
    rows = preview_rows(blueprint, manifest, content, {})
    by_field = {row["field"]: row for row in rows}
    assert by_field["Business Need"]["content"] == "Fresh need text" and by_field["Business Need"]["slide"] == 1
    assert by_field["Scope"]["content"] == "1 rows"
    assert by_field["First point"]["content"].startswith("(empty")
    blank = preview_rows(blueprint, manifest, content, {blueprint.sections[0].key: "blank"})
    assert all(row["content"] == "(blank)" for row in blank if row["slide"] == 1)


def test_overflow_report_raises_on_unreadable_file(tmp_path):
    from sdgen.preview import overflow_report

    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a deck")
    with pytest.raises(RuntimeError):
        overflow_report(bad, attempts=1)


@pytest.mark.skipif(
    not POWERPOINT.exists() or os.environ.get("SDGEN_COM_TESTS") != "1",
    reason="set SDGEN_COM_TESTS=1 on a machine with PowerPoint",
)
def test_export_slides_flags_text_that_overflows(sample_deck, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    from sdgen.preview import export_slides, overflow_report

    prs = Presentation(str(sample_deck))
    box = prs.slides[0].shapes.add_textbox(Inches(1), Inches(6.5), Inches(2), Inches(0.4))
    box.name = "Tiny Box"
    box.text_frame.word_wrap = True
    box.text_frame.text = "word " * 60
    deck = tmp_path / "overflow.pptx"
    prs.save(deck)
    exported = export_slides(deck, tmp_path / "png", width=320)
    assert [f.name for f in exported.files] == ["slide-01.png", "slide-02.png"]
    flagged = {(o.slide, o.shape) for o in exported.overflows}
    assert (1, "Tiny Box") in flagged
    assert {(o.slide, o.shape) for o in overflow_report(deck)} == flagged
    narrowed = {o.shape for o in export_slides(deck, tmp_path / "png2", width=320, only=[1], interest={1: {"Tiny Box"}}).overflows}
    assert "Tiny Box" in narrowed and "Scope Label" not in narrowed and "Scope Label" in {o.shape for o in exported.overflows}
