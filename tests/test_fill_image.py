import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from sdgen.fill.image import replace_picture


@pytest.fixture
def images(tmp_path):
    wide = tmp_path / "wide.png"
    tall = tmp_path / "tall.png"
    Image.new("RGB", (400, 100), "red").save(wide)
    Image.new("RGB", (100, 400), "blue").save(tall)
    return wide, tall


def _deck_with_picture(image, inside_group=False):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(1)).text_frame.text = "before"
    owner = slide.shapes.add_group_shape().shapes if inside_group else slide.shapes
    pic = owner.add_picture(str(image), Inches(2), Inches(2), Inches(4), Inches(3))
    pic.name = "Diagram Picture"
    slide.shapes.add_textbox(Inches(7), Inches(0), Inches(1), Inches(1)).text_frame.text = "after"
    return prs, slide, pic


def test_contain_fits_and_centres_inside_the_old_box(images):
    wide, tall = images
    prs, slide, old = _deck_with_picture(tall)
    new = replace_picture(slide, old, wide, fit="contain")
    assert new.width == Inches(4) and new.height == Inches(1)
    assert new.left == Inches(2) and new.top == Inches(3)
    assert new.name == "Diagram Picture"
    assert [s.name for s in slide.shapes][1] == "Diagram Picture"
    assert len([s for s in slide.shapes if s.shape_type == old.shape_type]) == 1
    assert len(slide.part.rels) == 2  # layout + one image; the old image relationship is gone


def test_cover_fills_the_box_and_crops(images):
    wide, tall = images
    prs, slide, old = _deck_with_picture(wide)
    new = replace_picture(slide, old, tall, fit="cover")
    assert (new.width, new.height, new.left, new.top) == (Inches(4), Inches(3), Inches(2), Inches(2))
    assert new.crop_top == pytest.approx(new.crop_bottom) and new.crop_top > 0.3
    assert new.crop_left == 0


def test_replacement_inside_a_group_stays_in_the_group(images, tmp_path):
    wide, tall = images
    prs, slide, old = _deck_with_picture(tall, inside_group=True)
    group = next(s for s in slide.shapes if s.shape_type is not None and hasattr(s, "shapes"))
    new = replace_picture(slide, old, wide)
    assert [s.name for s in group.shapes] == ["Diagram Picture"]
    assert new._element.getparent() is group._element
    out = tmp_path / "g.pptx"
    prs.save(out)
    Presentation(str(out))


def test_empty_picture_placeholder_is_filled(images):
    wide, _ = images
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[8])
    placeholder = slide.placeholders[1]
    new = replace_picture(slide, placeholder, wide)
    assert new.image.size == (400, 100)
