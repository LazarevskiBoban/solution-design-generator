from pathlib import Path

import pytest
from click.testing import CliRunner
from pptx import Presentation

from sdgen.analyze import analyze_deck, format_analysis, slugify
from sdgen.cli import main
from sdgen.inventory import inspect_deck
from sdgen.manifest import Manifest

CARRIER_DECK = Path("D:/NTT-architectures/NTT DATA Inc I IT I Solution Design I Carrier AP Invoice Integration.pptx")


def _titled_deck(tmp_path, titles):
    prs = Presentation()
    for title in titles:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
    path = tmp_path / "titled.pptx"
    prs.save(path)
    return path


def test_slugify():
    assert slugify("Business Need, Problem & Expected Outcomes") == "business_need_problem_expected_outcomes"
    assert slugify("2 Way Match") == "field_2_way_match"
    assert slugify("   ") == "field"


def test_fixture_candidates(sample_deck):
    analysis = analyze_deck(inspect_deck(sample_deck))
    by_key = {c.key: c for c in analysis.candidates}

    need = by_key["business_need"]
    assert need.include and need.kind == "text" and need.keep_prefix == "Business Need: "
    assert need.max_chars and need.max_chars >= 100

    scope = by_key["scope"]
    assert scope.kind == "table" and scope.columns == ["Function", "Countries"]
    assert "scope_label" not in by_key and not any(c.shape_name == "Scope Label" for c in analysis.candidates)

    body = by_key["first_point"]
    assert body.kind == "bullets" and body.include

    images = [c for c in analysis.candidates if c.kind == "image"]
    assert len(images) == 1 and not images[0].include
    assert not any(c.shape_name in ("ISPIC", "BTP") for c in analysis.candidates)

    assert analysis.globals == []
    assert {s.index: s.kind for s in analysis.slides} == {1: "content", 2: "static"}


def test_manifest_from_analysis(sample_deck, tmp_path):
    analysis = analyze_deck(inspect_deck(sample_deck))
    manifest = analysis.to_manifest("demo", source="sample.pptx")
    assert set(f.key for f in manifest.fields) == {"business_need", "scope", "first_point"}
    assert manifest.field("scope").columns == ["Function", "Countries"]
    assert "characters fit" in manifest.field("business_need").guidance
    manifest.save(tmp_path / "m.yaml")
    assert Manifest.load(tmp_path / "m.yaml") == manifest


def test_subject_detected_from_titles(tmp_path):
    titles = [
        "Overview: Demo Integration",
        "Design: Demo Integration (draft)",
        "Flows: Demo Integration | v1",
        "Contents",
    ]
    analysis = analyze_deck(inspect_deck(_titled_deck(tmp_path, titles)))
    assert len(analysis.globals) == 1
    assert analysis.globals[0].key == "subject"
    assert analysis.globals[0].replaces == "Demo Integration"


def test_cli_analyze_writes_manifest(sample_deck, tmp_path):
    out = tmp_path / "manifest.yaml"
    result = CliRunner().invoke(main, ["analyze", str(sample_deck), "-o", str(out), "--all"])
    assert result.exit_code == 0, result.output
    assert "business_need" in result.output
    assert "3 fields written" in result.output
    assert Manifest.load(out).name == "sample"
    assert format_analysis(analyze_deck(inspect_deck(sample_deck))).startswith("slides ")


@pytest.mark.skipif(not CARRIER_DECK.exists(), reason="example deck not available")
def test_carrier_deck_analysis():
    analysis = analyze_deck(inspect_deck(CARRIER_DECK))
    assert analysis.globals[0].replaces == "Carrier AP Invoice Integration"

    included = {(c.slide, c.key): c for c in analysis.candidates if c.include}
    assert included[(5, "business_need")].keep_prefix == "Business Need: "
    assert included[(5, "solution_overview")].shape_name == "TextBox 16"
    assert included[(5, "scope")].columns[0] == "Function"
    assert included[(5, "internal_effort")].mode == "token"
    assert included[(2, "document_version_control")].kind == "table"
    assert included[(1, "subtitle")].kind == "bullets"

    kinds = {s.index: s.kind for s in analysis.slides}
    assert kinds[8] == "diagram" and kinds[16] == "diagram"
    assert kinds[11] == "static" and kinds[13] == "content"
    assert {30, 31} <= set(analysis.exclude)

    manifest = analysis.to_manifest("ntt-solution-design")
    assert manifest.field("success_measurement_quantitative") is not None
    assert manifest.slides.exclude == [30, 31]


def test_short_text_under_a_label_is_a_field(tmp_path):
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Solution Design Areas: Demo"
    label = slide.shapes.add_textbox(Inches(1), Inches(1.6), Inches(3), Inches(0.35))
    label.text_frame.text = "DATA ARCHITECTURE"
    box = slide.shapes.add_textbox(Inches(1), Inches(2.0), Inches(3.5), Inches(1.9))
    box.text_frame.text = "None identified"
    lonely = slide.shapes.add_textbox(Inches(8), Inches(5.0), Inches(3), Inches(0.5))
    lonely.text_frame.text = "Just a remark"
    path = tmp_path / "areas.pptx"
    prs.save(path)

    analysis = analyze_deck(inspect_deck(path))
    by_key = {c.key: c for c in analysis.candidates}
    assert by_key["data_architecture"].include and by_key["data_architecture"].reason == "named after nearby label"
    assert not any(c.include for c in analysis.candidates if c.preview.startswith("Just a remark"))


def test_max_chars_uses_font_metrics():
    import pytest

    from sdgen.analyze import _max_chars
    from sdgen.inventory import ParagraphInfo, ShapeInfo
    from sdgen.textmetrics import load_font

    def shape(**kw):
        return ShapeInfo(id=1, name="box", kind="text", left=1, top=1, width=4, height=1.5, paragraphs=[ParagraphInfo(text="x", font_size=10, **kw)])

    arial = _max_chars(shape(font_name="Arial"))
    assert arial and arial % 10 == 0
    assert _max_chars(shape(font_name="Arial", line_spacing=1.5)) < arial
    assert _max_chars(shape(font_name="Arial"), prefix_len=20) == arial - 20
    if load_font("Arial") is not None:
        assert _max_chars(shape(font_name="Arial", bold=True)) < arial



def test_placeholder_budget_uses_the_inherited_size(tmp_path):
    from pptx import Presentation

    from sdgen.analyze import _max_chars
    from sdgen.textmetrics import FontSpec, capacity_chars

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.placeholders[1].text_frame.text = "Lockbox files arrive daily from three banks over SWIFT and are posted by the standard programs."
    path = tmp_path / "placeholder.pptx"
    prs.save(str(path))
    shape = next(s for s in inspect_deck(path).slides[0].shapes if s.placeholder_type in ("body", "object"))
    paragraph = shape.paragraphs[0]
    assert paragraph.font_size == 32.0 and paragraph.font_name == "Calibri" and paragraph.line_spacing == 1.0
    assert _max_chars(shape) < capacity_chars(shape.width * 72 - 14.4, shape.height * 72 - 7.2, FontSpec("Calibri", 14, False), 100, 0)


def test_line_budget_accompanies_the_character_budget(sample_deck):
    analysis = analyze_deck(inspect_deck(sample_deck))
    need = next(c for c in analysis.candidates if c.key == "business_need")
    assert need.max_lines and 1 < need.max_lines < 40
    binding = analysis.to_manifest("demo").field("business_need").bindings[0]
    assert binding.max_lines == need.max_lines and binding.max_chars == need.max_chars
