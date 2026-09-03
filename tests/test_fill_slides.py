from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from sdgen.fill.slides import clone_slide, move_slide, remove_slide, slide_index
from sdgen.inventory import inspect_deck

CARRIER_DECK = Path("D:/NTT-architectures/NTT DATA Inc I IT I Solution Design I Carrier AP Invoice Integration.pptx")


def _deck(tmp_path):
    red = tmp_path / "red.png"
    blue = tmp_path / "blue.png"
    Image.new("RGB", (20, 20), "red").save(red)
    Image.new("RGB", (30, 10), "blue").save(blue)

    prs = Presentation()
    for title in ("A", "B", "C"):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
    first = prs.slides[0]
    first.shapes.add_picture(str(red), Inches(1), Inches(2))
    first.shapes.add_picture(str(blue), Inches(3), Inches(2))
    first.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1)).text_frame.text = "body"
    first.notes_slide.notes_text_frame.text = "private notes"
    return prs


def _titles(prs):
    return [s.shapes.title.text for s in prs.slides]


def _pictures(slide):
    return [s for s in slide.shapes if hasattr(s, "image")]


def test_clone_inserts_after_source_with_working_images(tmp_path):
    prs = _deck(tmp_path)
    source = prs.slides[0]
    clone = clone_slide(prs, source)

    assert _titles(prs) == ["A", "A", "B", "C"]
    assert slide_index(prs, clone) == 1
    assert clone.slide_layout is source.slide_layout
    assert [s.name for s in clone.shapes] == [s.name for s in source.shapes]
    assert [p.image.blob for p in _pictures(clone)] == [p.image.blob for p in _pictures(source)]
    assert not clone.has_notes_slide

    out = tmp_path / "cloned.pptx"
    prs.save(out)
    reopened = Presentation(str(out))
    assert _titles(reopened) == ["A", "A", "B", "C"]
    assert [p.image.size for p in _pictures(reopened.slides[1])] == [(20, 20), (30, 10)]


def test_remove_drops_slide_and_its_part(tmp_path):
    prs = _deck(tmp_path)
    remove_slide(prs, prs.slides[1])
    assert _titles(prs) == ["A", "C"]

    out = tmp_path / "removed.pptx"
    prs.save(out)
    reopened = Presentation(str(out))
    assert _titles(reopened) == ["A", "C"]
    slide_parts = [p for p in reopened.part.package.iter_parts() if "slides/slide" in str(p.partname)]
    assert len(slide_parts) == 2


def test_move_slide_reorders(tmp_path):
    prs = _deck(tmp_path)
    move_slide(prs, prs.slides[2], 0)
    assert _titles(prs) == ["C", "A", "B"]
    move_slide(prs, prs.slides[0], 2)
    assert _titles(prs) == ["A", "B", "C"]


@pytest.mark.skipif(not CARRIER_DECK.exists(), reason="example deck not available")
def test_real_diagram_slide_clones_and_stale_slides_drop(tmp_path):
    prs = Presentation(str(CARRIER_DECK))
    slides = list(prs.slides)
    clone = clone_slide(prs, slides[15])
    remove_slide(prs, slides[30])
    remove_slide(prs, slides[29])
    out = tmp_path / "carrier-edited.pptx"
    prs.save(out)

    deck = inspect_deck(out)
    assert len(deck.slides) == 31
    original, copied = deck.slides[15], deck.slides[16]
    assert copied.title == original.title
    assert copied.shape_count == original.shape_count
    assert copied.connector_count == original.connector_count
    assert [s.kind for s in copied.walk()] == [s.kind for s in original.walk()]
    assert not any("Replace with slide" in s.text for slide in deck.slides for s in slide.walk())
    assert slide_index(prs, clone) == 16
