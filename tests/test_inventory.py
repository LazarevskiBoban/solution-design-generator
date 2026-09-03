import pytest
from click.testing import CliRunner
from pptx import Presentation

from sdgen.cli import main
from sdgen.inventory import DeckInfo, find_shape, format_inventory, inspect_deck


def _by_name(slide, name):
    return next(s for s in slide.walk() if s.name == name)


def test_deck_and_slide_basics(sample_deck):
    deck = inspect_deck(sample_deck)
    assert deck.width == pytest.approx(13.333, abs=0.01)
    assert deck.height == pytest.approx(7.5, abs=0.01)
    assert [s.index for s in deck.slides] == [1, 2]
    first = deck.slides[0]
    assert first.layout == "Title and Content"
    assert first.title == "Executive Overview: Demo Integration"
    assert first.notes == "guidance note"
    assert deck.slides[1].shapes == []


def test_placeholders_and_paragraph_levels(sample_deck):
    first = inspect_deck(sample_deck).slides[0]
    title = next(s for s in first.shapes if s.placeholder_type == "title")
    body = next(s for s in first.shapes if s.placeholder_idx == 1)
    assert title.kind == "text"
    assert body.placeholder_type == "object"
    assert [(p.text, p.level) for p in body.paragraphs] == [("First point", 0), ("Sub point", 1)]


def test_text_box_paragraph_formatting(sample_deck):
    first = inspect_deck(sample_deck).slides[0]
    box = _by_name(first, "Business Need Box")
    assert box.kind == "text"
    assert box.text.startswith("Business Need: ")
    assert box.paragraphs[1].bold is True
    assert box.paragraphs[1].font_size == 12
    assert box.has_geometry and box.width == pytest.approx(4.0, abs=0.01)
    assert box.preset == "rect"


def test_table_cells(sample_deck):
    first = inspect_deck(sample_deck).slides[0]
    table = _by_name(first, "Scope Table")
    assert table.kind == "table"
    assert (table.table.rows, table.table.cols) == (3, 2)
    assert table.table.cells[0] == ["Function", "Countries"]
    assert table.table.cells[1] == ["Finance", "ZA"]


def test_picture_group_and_connector(sample_deck):
    first = inspect_deck(sample_deck).slides[0]
    pictures = [s for s in first.shapes if s.kind == "picture"]
    assert len(pictures) == 1 and pictures[0].width == pytest.approx(2.0, abs=0.01)
    group = _by_name(first, "Systems Group")
    assert group.kind == "group"
    assert [c.text for c in group.children] == ["ISPIC", "BTP"]
    assert group.children[1].preset == "roundRect"
    assert first.connector_count == 1
    assert first.shape_count == len(first.shapes) + len(group.children)


def test_find_shape_reaches_nested_shapes(sample_deck):
    info = inspect_deck(sample_deck).slides[0]
    nested_id = _by_name(info, "Systems Group").children[0].id
    slide = Presentation(str(sample_deck)).slides[0]
    assert find_shape(slide, nested_id).text_frame.text == "ISPIC"
    assert find_shape(slide, 99999) is None


def test_json_roundtrip_and_text_format(sample_deck):
    deck = inspect_deck(sample_deck)
    assert DeckInfo.model_validate_json(deck.model_dump_json()) == deck
    text = format_inventory(deck)
    assert "== Slide 1" in text
    assert "'Scope Table' 3x2" in text
    assert "| Function | Countries" in text
    assert "notes: guidance note" in text


def test_cli_inspect_filters_slides(sample_deck):
    result = CliRunner().invoke(main, ["inspect", str(sample_deck), "--slide", "2"])
    assert result.exit_code == 0, result.output
    assert "== Slide 2" in result.output
    assert "== Slide 1" not in result.output
    as_json = CliRunner().invoke(main, ["inspect", str(sample_deck), "--json"])
    assert as_json.exit_code == 0
    assert '"slides"' in as_json.output
