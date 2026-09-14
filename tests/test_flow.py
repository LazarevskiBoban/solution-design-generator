from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from sdgen import palette
from sdgen.brief import Brief
from sdgen.flow import SYSTEM_PROMPT, FlowEdge, FlowNode, FlowSpec, clean_flow, draw_flow, plan_flows, to_mermaid
from sdgen.llm import MockLLM

SPEC = FlowSpec(
    title="Lockbox transport",
    nodes=[
        FlowNode(id="bank", label="Bank (SWIFT FileAct)", kind="external", lane="source"),
        FlowNode(id="autoclient", label="AutoClient VM", kind="system", lane="source"),
        FlowNode(id="sftp", label="Azure Blob SFTP", kind="store", lane="middleware"),
        FlowNode(id="btp", label="SAP BTP Cloud Integration", kind="system", lane="middleware"),
        FlowNode(id="s4", label="S/4HANA lockbox processing", kind="system", lane="target"),
    ],
    edges=[
        FlowEdge(source="bank", target="autoclient", label="lockbox file", kind="file"),
        FlowEdge(source="autoclient", target="sftp", label="SFTP put", kind="file"),
        FlowEdge(source="sftp", target="btp", label="poll", kind="async"),
        FlowEdge(source="btp", target="s4", label="file drop", kind="sync"),
    ],
)


def _blank_slide():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def test_draw_flow_in_bands_stacks_lanes_and_reads_left_to_right():
    prs, slide = _blank_slide()
    box = (Inches(0.5), Inches(1.0), Inches(12.0), Inches(5.5))
    created = draw_flow(slide, box, SPEC, prefix="Flow x")
    names = [s.name for s in created]
    assert names[0] == "Flow x canvas" and created[0].width == box[2]
    assert [n for n in names if n.startswith("Flow x lane") and not n.endswith(" mark")] == ["Flow x lane source", "Flow x lane middleware", "Flow x lane target"]
    by_name = {s.name: s for s in created}
    source, middleware, target = (by_name[f"Flow x lane {lane}"] for lane in ("source", "middleware", "target"))
    assert source.top < middleware.top < target.top and source.left == middleware.left == target.left and source.text_frame.text == "Source"
    bank, autoclient, sftp, btp, s4 = (by_name[f"Flow x node {n}"] for n in ("bank", "autoclient", "sftp", "btp", "s4"))
    assert bank.left < autoclient.left < sftp.left < btp.left < s4.left and bank.top == autoclient.top < sftp.top == btp.top < s4.top
    assert all(s.rotation == 0.0 for s in created) and bank.text_frame.text == "Bank (SWIFT FileAct)"
    connectors = [s for s in created if s.name.startswith("Flow x edge") and not s.name.endswith("label")]
    labels = [s for s in created if s.name.endswith("label")]
    assert len(connectors) == 4 and len(labels) == 4 and labels[0].text_frame.text == "lockbox file"
    glued = by_name["Flow x edge 1"]._element  # bank to autoclient: neighbours in one band
    assert glued.find(".//" + qn("a:stCxn")) is not None and glued.find(".//" + qn("a:tailEnd")).get("type") == "triangle"
    assert by_name["Flow x edge 2"]._element.find(".//" + qn("a:stCxn")) is None  # autoclient to sftp crosses the seam
    left, top, width, height = box
    for shape in created:
        assert shape.left >= left - 1 and shape.top >= top - 1 and shape.left + shape.width <= left + width + 1 and shape.top + shape.height <= top + height + 1
    assert _labels_clear_of_nodes(created)


def test_draw_flow_columns_option_keeps_the_side_by_side_look():
    prs, slide = _blank_slide()
    box = (Inches(0.5), Inches(1.0), Inches(12.0), Inches(5.5))
    created = draw_flow(slide, box, SPEC.model_copy(update={"layout": "columns"}), prefix="Flow x")
    names = [s.name for s in created]
    assert [n for n in names if n.startswith("Flow x lane") and not n.endswith(" mark")] == ["Flow x lane source", "Flow x lane middleware", "Flow x lane target"]
    connectors = [s for s in created if s.name.startswith("Flow x edge") and not s.name.endswith("label")]
    assert len(connectors) == 4 and all(c._element.find(".//" + qn("a:stCxn")) is not None and c._element.find(".//" + qn("a:tailEnd")).get("type") == "triangle" for c in connectors)
    by_name = {s.name: s for s in created}
    bank, autoclient, s4 = by_name["Flow x node bank"], by_name["Flow x node autoclient"], by_name["Flow x node s4"]
    assert bank.left == autoclient.left and bank.top < autoclient.top and bank.left < s4.left
    assert by_name["Flow x lane source"].left < by_name["Flow x lane target"].left and by_name["Flow x lane source"].top == by_name["Flow x lane target"].top
    left, top, width, height = box
    for shape in created:
        assert shape.left >= left - 1 and shape.top >= top - Inches(0.35)
        assert shape.left + shape.width <= left + width + Inches(0.75) and shape.top + shape.height <= top + height + 1


def test_narrow_box_keeps_the_bands_and_shrinks_the_nodes():
    prs, slide = _blank_slide()
    created = draw_flow(slide, (Inches(0.5), Inches(1.0), Inches(6.0), Inches(5.5)), SPEC)
    by_name = {s.name: s for s in created}
    bank, sftp, s4 = by_name["Flow node bank"], by_name["Flow node sftp"], by_name["Flow node s4"]
    assert bank.top < sftp.top < s4.top and bank.left < sftp.left < s4.left and bank.width < Inches(3.0)
    assert all(s.left + s.width <= Inches(6.5) + 1 for s in created)


def test_sequence_layout_draws_lifelines_and_numbered_arrows():
    from sdgen.flow import looks_like_sequence

    prs, slide = _blank_slide()
    edges = [FlowEdge(source="a" if i % 2 else "b", target="b" if i % 2 else "a", label=f"step {i}") for i in range(1, 6)]
    spec = FlowSpec(nodes=[FlowNode(id="a", label="BTP sFTP adapter", lane="middleware"), FlowNode(id="b", label="sFTP door", lane="source")], edges=edges, layout="sequence")
    assert looks_like_sequence(spec) and not looks_like_sequence(SPEC)
    created = draw_flow(slide, (Inches(0.5), Inches(1.0), Inches(12.0), Inches(5.5)), spec, prefix="Flow s")
    by_name = {s.name: s for s in created}
    assert "Flow s lifeline a" in by_name and "Flow s lifeline b" in by_name and not any(" lane " in n for n in by_name)
    a, b = by_name["Flow s node a"], by_name["Flow s node b"]
    assert a.top == b.top and a.left < b.left and by_name["Flow s lifeline a"].top >= a.top + a.height - 1
    arrows = [by_name[f"Flow s edge {i}"] for i in range(1, 6)]
    assert [s.top for s in arrows] == sorted(s.top for s in arrows) and all(s._element.find(".//" + qn("a:stCxn")) is None for s in arrows)
    labels = [by_name[f"Flow s edge {i} label"] for i in range(1, 6)]
    assert [label.text_frame.text for label in labels] == [f"{i}. step {i}" for i in range(1, 6)]
    assert _labels_clear_of_nodes(created)


def test_plan_flows_marks_a_handshake_as_a_sequence():
    from sdgen.plan import FlowRequest

    handshake = {"flows": [{"section": "ssh", "nodes": [{"id": "a", "label": "BTP adapter", "lane": "middleware"}, {"id": "b", "label": "sFTP server", "lane": "source"}], "edges": [{"source": "a", "target": "b", "label": "connect"}, {"source": "b", "target": "a", "label": "host key"}, {"source": "a", "target": "b", "label": "sign"}]}]}
    llm = _stub(handshake)
    flows = plan_flows(Brief(subject="s"), [FlowRequest(section="ssh", title="SSH handshake", purpose="", kind="sequence")], llm)
    assert flows["ssh"].layout == "sequence" and "a sequence: two to four participants" in llm.user
    detected = plan_flows(Brief(subject="s"), [FlowRequest(section="ssh", title="SSH", purpose="")], _stub(handshake))
    assert detected["ssh"].layout == "sequence"
    plain = _stub({"flows": [{"section": "x", "nodes": [{"id": "a", "label": "A", "lane": "source"}, {"id": "b", "label": "B", "lane": "target"}], "edges": [{"source": "a", "target": "b"}]}]})
    assert plan_flows(Brief(subject="s"), [FlowRequest(section="x", title="X", purpose="")], plain)["x"].layout == "bands"


def test_parallel_edges_between_two_nodes_do_not_stack_labels():
    prs, slide = _blank_slide()
    edges = [FlowEdge(source="a" if i % 2 else "b", target="b" if i % 2 else "a", label=f"step {i}") for i in range(1, 8)]
    spec = FlowSpec(nodes=[FlowNode(id="a", label="BTP adapter", lane="middleware"), FlowNode(id="b", label="sFTP door", lane="middleware")], edges=edges)
    created = draw_flow(slide, (Inches(0.5), Inches(1.0), Inches(12.0), Inches(5.5)), spec, prefix="Flow p")
    labels = [s for s in created if s.name.endswith(" label")]
    lines = [s for s in created if s.name.startswith("Flow p edge") and not s.name.endswith("label")]
    assert len(labels) == 7 and len({s.top for s in labels}) == 7 and len({s.top for s in lines}) == 7
    assert not any(_overlaps(_box(x), _box(y)) for i, x in enumerate(labels) for y in labels[i + 1 :])
    assert _labels_clear_of_nodes(created)


def test_mermaid_and_yaml_roundtrip(tmp_path):
    text = to_mermaid(SPEC)
    assert text.startswith("flowchart LR\n  subgraph source[Source]")
    assert '    bank{{"Bank (SWIFT FileAct)"}}' in text and '    sftp[("Azure Blob SFTP")]' in text
    assert "  bank ==>|lockbox file| autoclient" in text and "  sftp -.->|poll| btp" in text and "  btp -->|file drop| s4" in text
    SPEC.save(tmp_path / "flow.yaml")
    assert FlowSpec.load(tmp_path / "flow.yaml") == SPEC


def test_plan_flows_keeps_known_sections_and_cleans_edges():
    class JsonLLM:
        name = "j"

        def complete(self, system, user):
            return ""

        def complete_json(self, system, user, schema, name="result"):
            assert name == "flows" and "level_2 | Level 2 | end to end" in user
            return {
                "flows": [
                    {"section": "level_2", "title": "T", "nodes": [{"id": "a b", "label": "A", "lane": "source"}, {"id": "c", "label": "C", "lane": "target"}, {"id": "c", "label": "dup", "lane": "target"}], "edges": [{"source": "a b", "target": "c", "label": "x"}, {"source": "c", "target": "zzz"}]},
                    {"section": "unknown", "nodes": [{"id": "n", "label": "N", "lane": "source"}], "edges": []},
                ]
            }

    class Request:
        def __init__(self, section, title, purpose):
            self.section, self.title, self.purpose = section, title, purpose

    flows = plan_flows(Brief(subject="S", about="A"), [Request("level_2", "Level 2", "end to end")], JsonLLM())
    assert list(flows) == ["level_2"]
    spec = flows["level_2"]
    assert [n.id for n in spec.nodes] == ["a_b", "c"] and len(spec.edges) == 1 and spec.edges[0].source == "a_b"
    assert plan_flows(Brief(subject="S"), [Request("level_2", "", "")], MockLLM()) == {}
    assert plan_flows(Brief(subject="S"), [], MockLLM()) == {}
    assert clean_flow(FlowSpec(nodes=[FlowNode(id="x", label=" ", lane="source")])).nodes == []


def test_sap_detection_and_icon_keys(tmp_path):
    from sdgen.brief import Brief
    from sdgen.flow import clean_flow, plan_flows, uses_sap
    from sdgen.icons import icon_keys
    from sdgen.plan import FlowRequest

    assert uses_sap(Brief(subject="Lockbox", about="Files go from the bank to SAP S/4HANA."))
    assert not uses_sap(Brief(subject="Payroll", about="Files go from the bank to the payroll provider."))
    spec = FlowSpec(nodes=[FlowNode(id="s4", label="S/4HANA", icon="s4hana"), FlowNode(id="x", label="X", icon="nonsense")])
    cleaned = clean_flow(spec)
    assert [n.icon for n in cleaned.nodes] == ["s4hana", ""]
    cleaned.save(tmp_path / "f.yaml")
    assert FlowSpec.load(tmp_path / "f.yaml") == cleaned

    class Catcher:
        name = "fake"

        def complete(self, system, user):
            return ""

        def complete_json(self, system, user, schema, name="result"):
            self.user, self.schema = user, schema
            return {"flows": [{"section": "level_2", "nodes": [{"id": "a", "label": "A", "lane": "source", "icon": "bank"}, {"id": "b", "label": "B", "lane": "target", "icon": "s4hana"}], "edges": [{"source": "a", "target": "b"}]}]}

    request = FlowRequest(section="level_2", title="L2", purpose="")
    llm = Catcher()
    flows = plan_flows(Brief(subject="x"), [request], llm, icons=icon_keys())
    node_schema = llm.schema["properties"]["flows"]["items"]["properties"]["nodes"]["items"]["properties"]
    assert "Icon keys" in llm.user and node_schema["icon"]["enum"] == icon_keys()
    assert [n.icon for n in flows["level_2"].nodes] == ["bank", "s4hana"]
    plain = Catcher()
    plan_flows(Brief(subject="x"), [request], plain)
    assert "Icon keys" not in plain.user and "icon" not in plain.schema["properties"]["flows"]["items"]["properties"]["nodes"]["items"]["properties"]


def test_draw_flow_places_icon_pictures(tmp_path, monkeypatch):
    from PIL import Image
    from pptx.util import Inches

    import sdgen.flow as flow_module

    png = tmp_path / "icon.png"
    Image.new("RGB", (32, 32), "blue").save(png)
    monkeypatch.setattr(flow_module, "icon_png", lambda key: png if key == "s4hana" else None)
    spec = FlowSpec(nodes=[FlowNode(id="s4", label="S/4HANA", lane="target", icon="s4hana"), FlowNode(id="bank", label="Bank", lane="source")], edges=[FlowEdge(source="bank", target="s4")])
    created = draw_flow(_blank_slide()[1], (Inches(1), Inches(1), Inches(10), Inches(4)), spec, prefix="Flow x")
    names = [s.name for s in created]
    assert "Flow x icon s4" in names and "Flow x icon bank" not in names
    node = next(s for s in created if s.name == "Flow x node s4")
    assert node.text_frame.margin_left > Inches(0.3)
    # A node too narrow for icon plus label keeps the label width and gets a corner badge instead.
    narrow: list = []
    slide = _blank_slide()[1]
    small = flow_module._node(slide, FlowNode(id="s4", label="S/4HANA", icon="s4hana"), 0, 0, Inches(0.6), Inches(0.5), "Flow y", narrow)
    assert small.text_frame.margin_left == Inches(0.05) and narrow[0].width == flow_module.ICON_BADGE


def _box(shape):
    return (shape.left, shape.top, shape.left + shape.width, shape.top + shape.height)


def _overlaps(a, b) -> bool:
    return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]


def _labels_clear_of_nodes(created) -> bool:
    nodes = [s for s in created if " node " in s.name]
    labels = [s for s in created if s.name.endswith(" label")]
    return not any(_overlaps(_box(label), _box(node)) for label in labels for node in nodes)


def test_edges_that_skip_a_node_go_around_it():
    prs, slide = _blank_slide()
    spec = FlowSpec(
        nodes=[
            FlowNode(id="a", label="A", lane="middleware"),
            FlowNode(id="b", label="B", lane="middleware"),
            FlowNode(id="c", label="C", lane="middleware"),
            FlowNode(id="s", label="S", lane="source"),
            FlowNode(id="t", label="T", lane="target"),
        ],
        edges=[
            FlowEdge(source="a", target="b", label="one"),
            FlowEdge(source="b", target="c", label="two"),
            FlowEdge(source="a", target="c", label="skip"),
            FlowEdge(source="s", target="t", label="far"),
            FlowEdge(source="s", target="a", label="in"),
        ],
    )
    created = draw_flow(slide, (Inches(0.5), Inches(1.0), Inches(12.0), Inches(5.5)), spec, prefix="Flow x")
    by_name = {s.name: s for s in created}
    a, b, c = by_name["Flow x node a"], by_name["Flow x node b"], by_name["Flow x node c"]
    assert a.left < b.left < c.left and a.top == b.top == c.top
    skip = by_name["Flow x edge 3"]
    assert skip._element.find(".//" + qn("a:stCxn")) is None and skip._element.find(".//" + qn("a:tailEnd")).get("type") == "triangle"
    assert skip.top + skip.height > b.top + b.height and not _overlaps(_box(skip), _box(b))  # under the band, not through b
    far = by_name["Flow x edge 4"]
    rightmost = max(s.left + s.width for s in created if " node " in s.name)
    assert far.left + far.width > rightmost and far._element.find(".//" + qn("a:stCxn")) is None  # source to target through the right channel
    assert by_name["Flow x edge 1"]._element.find(".//" + qn("a:stCxn")) is not None
    assert _labels_clear_of_nodes(created)
    canvas = _box(by_name["Flow x canvas"])
    for shape in created:
        if " node " in shape.name or " edge " in shape.name and not shape.name.endswith("label"):
            box = _box(shape)
            assert box[0] >= canvas[0] - 1 and box[1] >= canvas[1] - 1 and box[2] <= canvas[2] + 1 and box[3] <= canvas[3] + 1


def test_wrapped_flow_draws_the_bands_twice_and_links_the_rows():
    prs, slide = _blank_slide()
    nodes = [FlowNode(id=f"n{i}", label=f"Step {i}", lane="middleware") for i in range(1, 10)]
    edges = [FlowEdge(source=f"n{i}", target=f"n{i + 1}", label=f"e{i}") for i in range(1, 9)]
    created = draw_flow(slide, (Inches(0.5), Inches(1.0), Inches(6.0), Inches(5.5)), FlowSpec(nodes=nodes, edges=edges), prefix="Flow w")
    by_name = {s.name: s for s in created}
    assert "Flow w lane middleware" in by_name and "Flow w lane middleware row 2" in by_name
    rows = [s for s in created if s.name.startswith("Flow w lane middleware")]
    assert rows[0].top < rows[1].top and rows[0].left == rows[1].left and rows[1].text_frame.text == "Middleware"
    first_row = [by_name[f"Flow w node n{i}"] for i in range(1, 10) if by_name[f"Flow w node n{i}"].top == by_name["Flow w node n1"].top]
    assert 2 <= len(first_row) < 9 and all(n.width >= Inches(1.1) - 1 for n in first_row)
    last, first_next = first_row[-1], by_name[f"Flow w node n{len(first_row) + 1}"]
    assert first_next.top > last.top and first_next.left < last.left
    link = by_name[f"Flow w edge {len(first_row)}"]
    assert link._element.find(".//" + qn("a:stCxn")) is None
    assert link.left + link.width > last.left + last.width and link.left < first_next.left
    assert link.top < last.top + last.height and link.top + link.height > first_next.top
    canvas = _box(by_name["Flow w canvas"])
    for shape in created:
        box = _box(shape)
        assert box[0] >= canvas[0] - 1 and box[1] >= canvas[1] - 1 and box[2] <= canvas[2] + 1 and box[3] <= canvas[3] + 1
    assert _labels_clear_of_nodes(created)


def test_nodes_and_lanes_take_the_sap_look():
    from pptx.enum.dml import MSO_LINE

    import sdgen.flow as flow_module

    prs, slide = _blank_slide()
    named = SPEC.model_copy(update={"lanes": {"target": "SAP Landscape (RISE)"}, "nodes": [SPEC.nodes[0].model_copy(update={"subtitle": "SWIFT FileAct"})] + SPEC.nodes[1:]})
    created = draw_flow(slide, (Inches(0.5), Inches(1.0), Inches(12.0), Inches(5.5)), named, prefix="Flow x")
    by_name = {s.name: s for s in created}
    source, target = by_name["Flow x lane source"], by_name["Flow x lane target"]
    assert source.fill.fore_color.rgb == palette.rgb(palette.GREY_FILL) and target.fill.fore_color.rgb == palette.rgb(palette.SAP_FILL)
    assert source.text_frame.text == "Source" and target.text_frame.text == "SAP Landscape (RISE)"
    assert "Flow x lane source mark" not in by_name and by_name["Flow x lane target mark"].text_frame.text == "SAP"
    bank, s4 = by_name["Flow x node bank"], by_name["Flow x node s4"]
    assert bank.line.dash_style == MSO_LINE.ROUND_DOT and bank.line.color.rgb == palette.rgb(palette.SLATE)
    assert s4.line.dash_style is None and s4.line.color.rgb == palette.rgb(palette.SAP_BLUE) and s4.fill.fore_color.rgb == palette.rgb(palette.WHITE)
    assert [p.text for p in bank.text_frame.paragraphs] == ["Bank (SWIFT FileAct)", "SWIFT FileAct"]
    lane_of = {n.id: n.lane for n in named.nodes}
    for shape in created:
        if " node " in shape.name:
            lane = _box(by_name[f"Flow x lane {lane_of[shape.name.split(' node ')[1]]}"])
            node = _box(shape)
            assert node[0] >= lane[0] and node[1] >= lane[1] and node[2] <= lane[2] and node[3] <= lane[3]
    short = flow_module._node(_blank_slide()[1], FlowNode(id="a", label="A", subtitle="tiny"), 0, 0, Inches(2), Inches(0.45), "Flow y")
    assert [p.text for p in short.text_frame.paragraphs] == ["A"]


def test_flat_box_falls_back_to_columns():
    prs, slide = _blank_slide()
    created = draw_flow(slide, (Inches(0.5), Inches(1.0), Inches(6.0), Inches(1.8)), SPEC, prefix="Flow f")
    bank = next(s for s in created if s.name == "Flow f node bank")
    s4 = next(s for s in created if s.name == "Flow f node s4")
    assert bank.left < s4.left and bank.top == s4.top and bank.height >= Inches(0.4)


def test_plan_flows_tells_the_model_the_drawing_area():
    from sdgen.flow import node_cap
    from sdgen.plan import FlowRequest

    class Catcher:
        name = "fake"

        def complete(self, system, user):
            return ""

        def complete_json(self, system, user, schema, name="result"):
            self.user = user
            return {"flows": []}

    llm = Catcher()
    requests = [
        FlowRequest(section="a", title="A", purpose="p", width_in=12.7, height_in=5.5),
        FlowRequest(section="b", title="B", purpose="q", width_in=6.3, height_in=2.6),
        FlowRequest(section="c", title="C", purpose=""),
    ]
    plan_flows(Brief(subject="x"), requests, llm)
    assert "drawing area 12.7 x 5.5 in, at most 9 nodes" in llm.user
    assert "drawing area 6.3 x 2.6 in, at most 4 nodes, two lanes at most" in llm.user
    assert "size unknown, at most 9 nodes" in llm.user
    assert node_cap(12.7, 5.5) == 9 and node_cap(6.3, 2.6) == 4 and node_cap(7, 4) == 8


def test_steps_are_kept_or_derived_from_the_edges():
    from sdgen.flow import flow_steps, walkthrough_text
    from sdgen.plan import FlowRequest

    assert flow_steps(SPEC)[:2] == ["Bank (SWIFT FileAct) to AutoClient VM: lockbox file.", "AutoClient VM to Azure Blob SFTP: SFTP put."]
    told = SPEC.model_copy(update={"steps": ["First.", " Second step. ", ""]})
    assert walkthrough_text(told) == "1. First.\n2. Second step."
    assert walkthrough_text(SPEC).startswith("1. Bank (SWIFT FileAct) to AutoClient VM: lockbox file.\n2. ")
    assert clean_flow(FlowSpec(nodes=[FlowNode(id="a", label="A")], steps=["  two   words ", " "])).steps == ["two words"]

    class Catcher:
        name = "fake"

        def complete(self, system, user):
            return ""

        def complete_json(self, system, user, schema, name="result"):
            assert "steps" in schema["properties"]["flows"]["items"]["properties"]
            return {"flows": [{"section": "x", "nodes": [{"id": "a", "label": "A", "lane": "source"}], "edges": [], "steps": ["A starts.", "A ends."]}]}

    flows = plan_flows(Brief(subject="s"), [FlowRequest(section="x", title="X", purpose="")], Catcher())
    assert flows["x"].steps == ["A starts.", "A ends."]


def test_sap_nodes_lanes_and_headings():
    from sdgen.flow import lane_is_sap, lane_title, node_is_sap

    assert [node_is_sap(n) for n in SPEC.nodes] == [False, False, False, True, True]
    assert [lane_is_sap(SPEC, lane) for lane in ("source", "middleware", "target")] == [False, True, True]
    assert not node_is_sap(FlowNode(id="x", label="SAP S/4HANA", icon="bank")) and node_is_sap(FlowNode(id="y", label="ERP", icon="s4hana"))
    assert node_is_sap(FlowNode(id="z", label="Core ERP", subtitle="SAP ECC 6.0"))
    named = SPEC.model_copy(update={"lanes": {"source": "Bank side", "target": "SAP Landscape (On-Premise)"}})
    assert lane_title(named, "source") == "Bank side" and lane_title(named, "middleware") == "Middleware"
    assert not lane_is_sap(named, "source") and lane_is_sap(named.model_copy(update={"lanes": {"source": "SAP side"}}), "source")
    text = to_mermaid(named)
    assert "  subgraph source[Bank side]" in text and '  subgraph target["SAP Landscape (On-Premise)"]' in text


def _stub(payload):
    class Stub:
        name = "fake"

        def complete(self, system, user):
            return ""

        def complete_json(self, system, user, schema, name="result"):
            self.schema, self.user = schema, user
            return payload

    return Stub()


def test_parse_systems_orders_by_role_and_slugs_ids():
    from sdgen.flow import OTHER_LANE, parse_systems

    text = "SAP BTP Integration Suite | middleware | new\nS/4HANA on RISE | target | change\nWindows VM on Azure | source | keep\nS/4HANA on RISE | receiver | keep\n\n- Banks | sender"
    lanes = parse_systems(text)
    assert [(lane.id, lane.role, lane.sap) for lane in lanes] == [
        ("windows_vm_on_azure", "source", False),
        ("banks", "source", False),
        ("sap_btp_integration_suite", "middleware", True),
        ("s_4hana_on_rise", "target", True),
        ("s_4hana_on_rise_2", "target", True),
    ]
    assert lanes[2].title == "SAP BTP Integration Suite" and lanes[2].change == "new" and OTHER_LANE.id == "other"
    assert parse_systems("") == []


def test_plan_flows_builds_the_lane_enum_from_the_systems_fact(tmp_path):
    from sdgen.flow import lane_is_sap, lane_order, lane_title
    from sdgen.plan import FlowRequest

    brief = Brief(subject="s", facts={"systems": "Windows VM on Azure | source | keep\nSAP BTP Integration Suite | middleware | new\nS/4HANA on RISE | target | change"})
    request = FlowRequest(section="x", title="X", purpose="")
    payload = {
        "flows": [
            {
                "section": "x",
                "lanes": {"s_4hana_on_rise": "RISE S/4HANA"},
                "nodes": [
                    {"id": "vm", "label": "VM sFTP door", "lane": "windows_vm_on_azure"},
                    {"id": "flow", "label": "Inbound iFlow", "kind": "step", "lane": "sap_btp_integration_suite"},
                    {"id": "s4out", "label": "S/4 OUT", "kind": "store", "lane": "windows_vm_on_azure"},
                    {"id": "ops", "label": "Operations team", "kind": "external", "lane": "other"},
                    {"id": "lost", "label": "Alert mail", "kind": "step", "lane": "nowhere"},
                ],
                "edges": [{"source": "vm", "target": "flow", "kind": "error", "label": "login failed"}],
            }
        ]
    }
    llm = _stub(payload)
    flows = plan_flows(brief, [request], llm)
    node_schema = llm.schema["properties"]["flows"]["items"]["properties"]["nodes"]["items"]["properties"]
    assert node_schema["lane"]["enum"] == ["windows_vm_on_azure", "sap_btp_integration_suite", "s_4hana_on_rise", "other"]
    assert "# Lanes" in llm.user and "- s_4hana_on_rise | S/4HANA on RISE | target | change" in llm.user
    spec = flows["x"]
    assert {n.id: n.lane for n in spec.nodes} == {"vm": "windows_vm_on_azure", "flow": "sap_btp_integration_suite", "s4out": "s_4hana_on_rise", "ops": "other", "lost": "sap_btp_integration_suite"}
    assert [lane.id for lane in spec.systems] == ["windows_vm_on_azure", "sap_btp_integration_suite", "s_4hana_on_rise", "other"]
    assert lane_title(spec, "s_4hana_on_rise") == "RISE S/4HANA" and lane_title(spec, "windows_vm_on_azure") == "Windows VM on Azure"
    assert lane_is_sap(spec, "s_4hana_on_rise") and lane_is_sap(spec, "sap_btp_integration_suite") and not lane_is_sap(spec, "windows_vm_on_azure")
    assert spec.edges[0].kind == "error" and lane_order(SPEC) == ["source", "middleware", "target"]
    spec.save(tmp_path / "f.yaml")
    assert FlowSpec.load(tmp_path / "f.yaml") == spec
    text = to_mermaid(spec)
    assert "subgraph windows_vm_on_azure[Windows VM on Azure]" in text and "vm -.->|login failed| flow" in text


def test_lane_mismatches_are_reported_when_ambiguous():
    from sdgen.flow import lane_mismatches, parse_systems

    lanes = parse_systems("SAP BTP Integration Suite | middleware\nS/4HANA on RISE | target\nAzure VM | source")
    spec = clean_flow(FlowSpec(nodes=[FlowNode(id="a", label="S/4 to VM copy", lane="sap_btp_integration_suite"), FlowNode(id="b", label="Timer", lane="sap_btp_integration_suite")], systems=lanes), lanes)
    assert [n.lane for n in spec.nodes] == ["sap_btp_integration_suite"] * 2
    assert lane_mismatches(spec) == ["'S/4 to VM copy' sits in SAP BTP Integration Suite but names Azure VM, S/4HANA on RISE"]
    assert lane_mismatches(SPEC) == []


def test_plan_flows_keeps_subtitles_and_lane_headings(tmp_path):
    from sdgen.plan import FlowRequest

    request = FlowRequest(section="x", title="X", purpose="")
    llm = _stub({"flows": [{"section": "x", "lanes": {"source": " Bank  side ", "middleware": "", "bogus": "No"}, "nodes": [{"id": "a", "label": "Bank", "subtitle": " SWIFT  FileAct ", "lane": "source"}, {"id": "b", "label": "S/4HANA", "lane": "target"}], "edges": [{"source": "a", "target": "b"}]}]})
    flows = plan_flows(Brief(subject="s"), [request], llm)
    flow_schema = llm.schema["properties"]["flows"]["items"]["properties"]
    assert "subtitle" in flow_schema["nodes"]["items"]["properties"] and list(flow_schema["lanes"]["properties"]) == ["source", "middleware", "target"]
    assert "subtitle" in SYSTEM_PROMPT and "Lanes: a heading" in SYSTEM_PROMPT
    spec = flows["x"]
    assert spec.lanes == {"source": "Bank side"} and [n.subtitle for n in spec.nodes] == ["SWIFT FileAct", ""]
    spec.save(tmp_path / "f.yaml")
    assert FlowSpec.load(tmp_path / "f.yaml") == spec
    odd = _stub({"flows": [{"section": "x", "lanes": "none", "nodes": [{"id": "a", "label": "A", "lane": "source"}], "edges": []}]})
    assert plan_flows(Brief(subject="s"), [request], odd)["x"].lanes == {}


def test_plan_flows_offers_the_matching_reference_pack(tmp_path, monkeypatch):
    from sdgen.plan import FlowRequest

    monkeypatch.setenv("SDGEN_REFS", str(tmp_path))
    request = FlowRequest(section="x", title="X", purpose="")
    sap_brief = Brief(subject="Lockbox", about="Files reach SAP S/4HANA through Cloud Integration as IDocs.")
    answer = {"flows": [{"section": "x", "reference": "sap:RA0021", "nodes": [{"id": "a", "label": "A", "lane": "source"}], "edges": []}]}
    llm = _stub(answer)
    flows = plan_flows(sap_brief, [request], llm)
    enum = llm.schema["properties"]["flows"]["items"]["properties"]["reference"]["enum"]
    assert "sap:RA0021" in enum and "" in enum and len(enum) == 34
    assert "# SAP Architecture Center reference architectures" in llm.user and "sap:RA0021 | Application to Application Integration" in llm.user
    assert "How SAP Architecture Center names" not in llm.user and "Reference:" in SYSTEM_PROMPT
    assert flows["x"].reference == "sap:RA0021"
    flows["x"].save(tmp_path / "f.yaml")
    assert FlowSpec.load(tmp_path / "f.yaml").reference == "sap:RA0021"
    # a fetched kit adds the block names of the closest reference diagrams
    folder = tmp_path / "sap" / "RA0021" / "drawio"
    folder.mkdir(parents=True)
    (folder / "a.drawio").write_text('<mxfile><diagram id="d" name="n"><mxGraphModel><root><mxCell id="1" value="SAP Cloud Connector" vertex="1"/></root></mxGraphModel></diagram></mxfile>', encoding="utf-8")
    kit = _stub({"flows": []})
    plan_flows(sap_brief, [request], kit)
    assert "# How SAP Architecture Center names the blocks" in kit.user and "- Application to Application Integration: SAP Cloud Connector" in kit.user
    # a brief about another system sees no catalogue and no reference property
    plain = _stub(answer)
    plan_flows(Brief(subject="Payroll", about="Workday sends files to the payroll provider."), [request], plain)
    assert "reference" not in plain.schema["properties"]["flows"]["items"]["properties"] and "reference architectures" not in plain.user
    assert clean_flow(FlowSpec(nodes=[FlowNode(id="a", label="A")], reference="sap:RA9999")).reference == ""
    assert clean_flow(FlowSpec(nodes=[FlowNode(id="a", label="A")], reference=" sap:RA0022 ")).reference == "sap:RA0022"


def test_plan_flows_lists_tagged_material_first_and_sends_pictures_only_when_given():
    from sdgen.material import new_material
    from sdgen.plan import FlowRequest

    request = FlowRequest(section="level_2", title="L2", purpose="")
    told = Brief(
        subject="Lockbox",
        about="Bank files to SAP S/4HANA.",
        material=[
            new_material("text", title="General note", text="Two lanes."),
            new_material("image", title="Sketch", tags=["level_2"], file="a.png", text="Bank -> SFTP -> CI -> S/4"),
            new_material("text", title="Elsewhere", tags=["other"], text="ignored"),
        ],
    )
    llm = _stub({"flows": []})
    plan_flows(told, [request], llm)
    block = llm.user.split("# Reference material", 1)[1].split("# Brief", 1)[0]
    assert block.index("## Sketch") < block.index("## General note") and "Elsewhere" not in llm.user
    assert "attached pictures" in SYSTEM_PROMPT.lower()

    class Seeing:
        name = "fake"

        def complete(self, system, user):
            return ""

        def complete_json(self, system, user, schema, name="result", images=None):
            self.images = images
            return {"flows": []}

    seeing = Seeing()
    plan_flows(told, [request], seeing, images=[(b"png", "image/png")])
    assert seeing.images == [(b"png", "image/png")]
