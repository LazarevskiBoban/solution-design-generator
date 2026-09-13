from pptx import Presentation
from pptx.util import Inches, Pt

from sdgen.fill.table import fill_table


def _table_frame(rows=4, cols=3, footer="Total"):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = slide.shapes.add_table(rows, cols, Inches(1), Inches(1), Inches(8), Inches(0.4 * rows))
    table = frame.table
    for c, name in enumerate(["Ref", "Item", "Notes"][:cols]):
        table.cell(0, c).text = name
    if rows > 1:
        table.cell(1, 0).text = "1"
        table.cell(1, 1).text = "first"
        table.cell(1, 1).text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    if footer:
        table.cell(rows - 1, 0).text = footer
    return prs, frame


def _texts(frame):
    return [[cell.text for cell in row.cells] for row in frame.table.rows]


def test_resize_columns_adds_and_removes_columns_keeping_the_total_width():
    from sdgen.fill.table import resize_columns

    prs, frame = _table_frame(rows=3, cols=3, footer=None)
    total = sum(c.width for c in frame.table.columns)
    resize_columns(frame, 5)
    assert len(frame.table.columns) == 5 and sum(c.width for c in frame.table.columns) == total
    assert all(len(row.cells) == 5 for row in frame.table.rows)
    assert _texts(frame)[0] == ["Ref", "Item", "Notes", "", ""] and _texts(frame)[1] == ["1", "first", "", "", ""]
    resize_columns(frame, 2)
    assert len(frame.table.columns) == 2 and sum(c.width for c in frame.table.columns) == total
    assert _texts(frame) == [["Ref", "Item"], ["1", "first"], ["", ""]]


def test_resize_columns_refuses_merged_cells():
    import pytest

    from sdgen.fill.table import resize_columns

    prs, frame = _table_frame(rows=2, cols=3, footer=None)
    frame.table.cell(0, 0).merge(frame.table.cell(0, 1))
    with pytest.raises(ValueError):
        resize_columns(frame, 4)
    assert len(frame.table.columns) == 3


def test_dict_rows_clone_template_row_and_keep_footer():
    prs, frame = _table_frame()
    rows = [
        {"Ref": "A", "Item": "alpha", "Notes": "n1"},
        {"ref": "B", "item": "beta"},
        {"Ref": "C", "Item": "gamma", "Extra": "ignored"},
    ]
    fill_table(frame, rows, header_rows=1, keep_last_row_if="Total")
    assert _texts(frame) == [
        ["Ref", "Item", "Notes"],
        ["A", "alpha", "n1"],
        ["B", "beta", ""],
        ["C", "gamma", ""],
        ["Total", "", ""],
    ]
    sizes = {frame.table.cell(r, 1).text_frame.paragraphs[0].runs[0].font.size for r in (1, 2, 3)}
    assert sizes == {Pt(9)}
    assert frame.height == sum(row.height for row in frame.table.rows)


def test_list_rows_and_explicit_columns():
    prs, frame = _table_frame(footer=None)
    fill_table(frame, [["9", "nine", "x", "overflow"]])
    assert _texts(frame) == [["Ref", "Item", "Notes"], ["9", "nine", "x"]]

    fill_table(frame, [{"identifier": "7", "label": "seven"}], columns=["Identifier", "Label", "Comment"])
    assert _texts(frame) == [["Ref", "Item", "Notes"], ["7", "seven", ""]]


def test_empty_rows_leave_one_blank_row(tmp_path):
    prs, frame = _table_frame()
    fill_table(frame, [], keep_last_row_if="Total")
    assert _texts(frame) == [["Ref", "Item", "Notes"], ["", "", ""], ["Total", "", ""]]
    out = tmp_path / "t.pptx"
    prs.save(out)
    assert len(Presentation(str(out)).slides[0].shapes[0].table.rows) == 3


def test_header_only_table_uses_header_as_template():
    prs, frame = _table_frame(rows=1, footer=None)
    fill_table(frame, [["a", "b", "c"], ["d", "e", "f"]])
    assert _texts(frame) == [["Ref", "Item", "Notes"], ["a", "b", "c"], ["d", "e", "f"]]


def test_row_heights_follow_wrapped_text_and_settle_the_rows():
    from sdgen.fill.table import append_rows, row_heights

    prs, frame = _table_frame(footer=None)
    template_h = frame.table.rows[1].height
    long = "This note wraps over several lines because the column is narrow and the sentence keeps going on and on without a stop."
    estimated = row_heights(frame, rows=[["1", "a", "n"], ["2", "b", long]])
    assert estimated[0] <= template_h * 1.1 and estimated[1] > template_h * 2
    fill_table(frame, [["1", "a", "n"], ["2", "b", long]])
    heights = [row.height for row in frame.table.rows]
    assert heights[1] <= template_h * 1.1 and heights[2] > template_h * 2 and frame.height == sum(heights)
    append_rows(frame, [["3", "c", "n"]])
    assert _texts(frame)[-1] == ["3", "c", "n"] and len(frame.table.rows) == 4


def test_marker_text_does_not_settle_the_template_row():
    prs, frame = _table_frame(footer=None)
    template_h = frame.table.rows[1].height
    fill_table(frame, [["{{a_very_long_marker_that_would_wrap_in_a_narrow_column}}"]], settle=False)
    assert frame.table.rows[1].height == template_h and frame.height == sum(row.height for row in frame.table.rows)
